# -*- coding: utf-8 -*-
import xlrd
from tablib import Dataset
from django.conf import settings
from django.shortcuts import render
from contabilidad.forms import *
from contabilidad.models import *
from almacenes.models import *
from activos_f.models import *
from django.template import RequestContext
from django.contrib.auth.models import User
from django.http import HttpResponseRedirect, Http404, HttpResponse
from django.http import JsonResponse
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db.models import Q
from dal import autocomplete
from . import models
from django.core import serializers
import datetime
from django.db.models import F
from django.db.models import Sum

from django.views.generic import ListView, CreateView, UpdateView, FormView, TemplateView, DeleteView
from django.core.urlresolvers import reverse_lazy

from easy_pdf.views import PDFTemplateView
from openpyxl import Workbook
from openpyxl.writer.excel import save_virtual_workbook

class crearAsiento(CreateView):
	model_pk = nivel_cinco
	model = encabezadoas
	form_class = crearEncabezadoForm
	template_name = 'contabilidad/nuevoasiento.html'
	success_url = 'contabilidad:listaasiento'
	def get_context_data(self, **kwargs):
		context = super(crearAsiento, self).get_context_data(**kwargs)
		return context
			
	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		selec = self.request.GET.getlist('envia')
		try:
			aux1 = []
			aux2 = []
			aux4 = []
			aux5 = []
			for i in selec:
				ingreso = encingresoitem.objects.get(id=i)
				for j in ingreso.ingresoitem_set.all():
					aux4.append(j.item.categoria.num_cuenta)
					aux5.append(j.sub_total)
					aux1.append(j.item.categoria.num_cuenta)
					aux1.append(0)
					aux2.append(j.sub_total)
					aux2.append(0)
			aux3 = zip(aux1, aux2)
			aux6 = zip(aux4, aux5)
			total_ingreso_objects = ingreso.ingresoitem_set.all().count()
			crearAsientoFormSetIng = inlineformset_factory(encabezadoas, asiento, form=crearAsientoForm , extra=total_ingreso_objects)
			if ingreso.iva == True:
				crear_asiento_form_set = crearAsientoFormSetIng(initial=[{'nivel_cinco': i, 'debe_b': j} for i, j in aux3])
			else:
				crear_asiento_form_set = crearAsientoFormSetIng(initial=[{'nivel_cinco': i, 'debe_b': j} for i, j in aux6])
		except Exception as e:
			crear_asiento_form_set = crearAsientoFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

	def post(self, request, *args, **kwargs):
		self.object=None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		selec = self.request.GET.getlist('envia')
		crear_asiento_form_set = crearAsientoFormSet(request.POST)
		if form.is_valid() and crear_asiento_form_set.is_valid():
			form = form.save(commit=False)
			aux = 0
			aux1 = 0
			aux2 = 0
			e = self.model.objects.filter(tipo_trans='E').exclude(anulado=True, cerrado=True).count()
			i = self.model.objects.filter(tipo_trans='I').exclude(anulado=True, cerrado=True).count()
			t = self.model.objects.filter(tipo_trans='T').exclude(anulado=True, cerrado=True).count()
			if (e == None) & (form.tipo_trans=='E'):
				form.numero_comp = 1
			elif (e!= None) & (form.tipo_trans=='E'):
				form.numero_comp = e + 1
			if (i == None) & (form.tipo_trans=='I'):
				form.numero_comp = 1
			elif (i!= None) & (form.tipo_trans=='I'):
				form.numero_comp = i + 1
			if (t == None) & (form.tipo_trans=='T'):
				form.numero_comp = 1
			elif (t!= None) & (form.tipo_trans=='T'):
				form.numero_comp = t + 1
			form.save()
			crear_asiento_form_set.instance = form
			formset = crear_asiento_form_set.save(commit=False)
			for f in formset:
				if (f.debe_b != None) & (f.debe_b != 0):
					f.haber_b = 0.00
					f.haber_d = 0.00
					f.auxiliar = 1

				elif (f.haber_b != None) & (f.haber_b != 0):
					f.debe_b = 0.00
					f.debe_d = 0.00
					f.auxiliar = 2

			crear_asiento_form_set.save()
			try:
				for i in selec:
					ingreso = encingresoitem.objects.get(id=i)
					ingreso.estado = True
					ingreso.save()
			except Exception as e:
				pass
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set, error='El haber no debe quedar vacío, verifique los datos.'))

class crearAsiento1(crearAsiento):
	model_pk = nivel_cinco
	model = encabezadoas
	form_class = crearEncabezadoForm
	template_name = 'contabilidad/nuevoasiento.html'
	success_url = 'contabilidad:listaasiento'
	def get_context_data(self, **kwargs):
		context = super(crearAsiento1, self).get_context_data(**kwargs)
		try:
			pk1 = self.kwargs['pk']
		except Exception as e:
			pass
		return context
			
	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		try:
			pk1 = self.kwargs['pk']
			aux1 = []
			aux2 = []
			aux4 = []
			aux = encsalidaitem.objects.get(id=pk1)
			for i in aux.salidaitem_set.all():
				aux1.append(i.encsalidaitem.cliente.num_cuenta)
				aux1.append(i.item.categoria.num_cuenta)
				aux2.append(0)
				aux2.append(i.sub_total)
				aux4.append(i.sub_total)
				aux4.append(0)
			aux3 = zip(aux1, aux2, aux4)
			total_salida_objects = aux.salida_item_set.all().count()
			crearAsientoFormSetSal = inlineformset_factory(encabezadoas, asiento, form=crearAsientoForm , extra=total_salida_objects)
			crear_asiento_form_set = crearAsientoFormSetSal(initial=[{'nivel_cinco': i, 'debe_b': k, 'haber_b': j} for i, j, k in aux3])
		except Exception as e:
			crear_asiento_form_set = crearAsientoFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

	def post(self, request, *args, **kwargs):
		self.object=None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_asiento_form_set = crearAsientoFormSet(request.POST)
		if form.is_valid() and crear_asiento_form_set.is_valid():
			form = form.save(commit=False)
			aux = 0
			aux1 = 0
			aux2 = 0
			e = self.model.objects.filter(tipo_trans='E').exclude(anulado=True, cerrado=True).count()
			i = self.model.objects.filter(tipo_trans='I').exclude(anulado=True, cerrado=True).count()
			t = self.model.objects.filter(tipo_trans='T').exclude(anulado=True, cerrado=True).count()
			if (e == None) & (form.tipo_trans=='E'):
				form.numero_comp = 1
			elif (e!= None) & (form.tipo_trans=='E'):
				form.numero_comp = e + 1
			if (i == None) & (form.tipo_trans=='I'):
				form.numero_comp = 1
			elif (i!= None) & (form.tipo_trans=='I'):
				form.numero_comp = i + 1
			if (t == None) & (form.tipo_trans=='T'):
				form.numero_comp = 1
			elif (t!= None) & (form.tipo_trans=='T'):
				form.numero_comp = t + 1
			form.save()
			crear_asiento_form_set.instance = form
			formset = crear_asiento_form_set.save(commit=False)
			for f in formset:
				if (f.debe_b != None) & (f.debe_b != 0):
					f.haber_b = 0.00
					f.haber_d = 0.00
					f.auxiliar = 1

				elif (f.haber_b != None) & (f.haber_b != 0):
					f.debe_b = 0.00
					f.debe_d = 0.00
					f.auxiliar = 2

			crear_asiento_form_set.save()
			try:
				pk1 = self.kwargs['pk']
				sal = encsalidaitem.objects.get(id=pk1)
				sal.estado = True
				sal.save()
			except Exception as e:
				pass
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

class crearAsiento2(CreateView):
	model_pk = nivel_cinco
	model = encabezadoas
	form_class = crearEncabezadoForm
	template_name = 'contabilidad/nuevoasiento.html'
	success_url = 'contabilidad:listaasiento'
	def get_context_data(self, **kwargs):
		context = super(crearAsiento2, self).get_context_data(**kwargs)
		return context
			
	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		selec = self.request.GET.getlist('envia')
		campo = self.request.GET.get('campo')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		fecha = f.date()
		y = f.year
		y_ant = f.year - 1
		f1 = str(y) + "-09-30"
		f2 = str(y_ant) + "-09-30"
		fecha1 = f.strftime("%d/%m/%Y")
		f3 = datetime.datetime.strptime(f1, "%Y-%m-%d")
		f4 = datetime.datetime.strptime(f2, "%Y-%m-%d")
		fecha2 = f4.strftime("%d/%m/%Y")
		fecha_ant1 = f3.date()
		fecha_ant2 = f4.date()
		if (f.month != 10) & (f.month != 11) & (f.month != 12):
			f5 = fecha_ant2
		else:
			f5 = fecha_ant1
		try:
			aux1 = []
			aux2 = []
			for i in selec:
				activo = Alta_baja_activo.objects.get(id=i)
				if activo.vinculado == True:
					aux1.append(activo.tipo_bien.rubro.num_cuenta)
					aux1.append(activo.tipo_bien.rubro.num_cuenta)
					aux2.append(activo.deprec_gestion_tot(fecha, f5, '1'))
					aux2.append(activo.deprec_acum_tot(fecha, f5, '1'))
				else:
					aux1.append(activo.tipo_bien.rubro.num_cuenta)
					aux1.append(activo.tipo_bien.rubro.num_cuenta)
					aux2.append(activo.deprec_gestion_tot(fecha, f5, 0))
					aux2.append(activo.deprec_acum_tot(fecha, f5, 0))
			aux3 = zip(aux1, aux2)
			crear_asiento_form_set = crearAsientoFormSet(initial=[{'nivel_cinco': i, 'debe_b': j} for i, j in aux3])
		except Exception as e:
			crear_asiento_form_set = crearAsientoFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

	def post(self, request, *args, **kwargs):
		self.object=None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		selec = self.request.GET.getlist('envia')
		crear_asiento_form_set = crearAsientoFormSet(request.POST)
		if form.is_valid() and crear_asiento_form_set.is_valid():
			form = form.save(commit=False)
			aux = 0
			aux1 = 0
			aux2 = 0
			e = self.model.objects.filter(tipo_trans='E').exclude(anulado=True, cerrado=True).count()
			i = self.model.objects.filter(tipo_trans='I').exclude(anulado=True, cerrado=True).count()
			t = self.model.objects.filter(tipo_trans='T').exclude(anulado=True, cerrado=True).count()
			if (e == None) & (form.tipo_trans=='E'):
				form.numero_comp = 1
			elif (e!= None) & (form.tipo_trans=='E'):
				form.numero_comp = e + 1
			if (i == None) & (form.tipo_trans=='I'):
				form.numero_comp = 1
			elif (i!= None) & (form.tipo_trans=='I'):
				form.numero_comp = i + 1
			if (t == None) & (form.tipo_trans=='T'):
				form.numero_comp = 1
			elif (t!= None) & (form.tipo_trans=='T'):
				form.numero_comp = t + 1
			form.save()
			crear_asiento_form_set.instance = form
			formset = crear_asiento_form_set.save(commit=False)
			for f in formset:
				if (f.debe_b != None) & (f.debe_b != 0):
					f.haber_b = 0.00
					f.haber_d = 0.00
					f.auxiliar = 1

				elif (f.haber_b != None) & (f.haber_b != 0):
					f.debe_b = 0.00
					f.debe_d = 0.00
					f.auxiliar = 2

			crear_asiento_form_set.save()
			try:
				for i in selec:
					activo = Alta_baja_activo.objects.get(id=i)
					activo.estado = True
					activo.save()
			except Exception as e:
				pass
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

class crearAsiento3(CreateView):
	model_pk = nivel_cinco
	model = encabezadoas
	form_class = crearEncabezadoForm
	template_name = 'contabilidad/nuevoasiento.html'
	success_url = 'contabilidad:listaasiento'
	def get_context_data(self, **kwargs):
		context = super(crearAsiento3, self).get_context_data(**kwargs)
		return context
			
	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		selec = self.request.GET.getlist('envia')
		campo = self.request.GET.get('campo')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		fecha = f.date()
		y = f.year
		y_ant = f.year - 1
		f1 = str(y) + "-09-30"
		f2 = str(y_ant) + "-09-30"
		fecha1 = f.strftime("%d/%m/%Y")
		f3 = datetime.datetime.strptime(f1, "%Y-%m-%d")
		f4 = datetime.datetime.strptime(f2, "%Y-%m-%d")
		fecha2 = f4.strftime("%d/%m/%Y")
		fecha_ant1 = f3.date()
		fecha_ant2 = f4.date()
		if (f.month != 10) & (f.month != 11) & (f.month != 12):
			f5 = fecha_ant2
		else:
			f5 = fecha_ant1
		try:
			aux1 = []
			aux2 = []
			for i in selec:
				capital = Capitalizacion.objects.get(id=i)
				if capital.alta_activo.vinculado == True:
					aux1.append(capital.item.categoria.num_cuenta)
					aux1.append(capital.item.categoria.num_cuenta)
					aux2.append(capital.deprec_gestion_cap_tot(fecha, f5, '1'))
					aux2.append(capital.deprec_acum_cap_tot(fecha, f5, '1'))
				else:
					aux1.append(capital.item.categoria.num_cuenta)
					aux1.append(capital.item.categoria.num_cuenta)
					aux2.append(capital.deprec_gestion_cap_tot(fecha, f5, '1'))
					aux2.append(capital.deprec_acum_cap_tot(fecha, f5, '1'))
			aux3 = zip(aux1, aux2)
			crear_asiento_form_set = crearAsientoFormSet(initial=[{'nivel_cinco': i, 'debe_b': j} for i, j in aux3])
		except Exception as e:
			crear_asiento_form_set = crearAsientoFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

	def post(self, request, *args, **kwargs):
		self.object=None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		selec = self.request.GET.getlist('envia')
		crear_asiento_form_set = crearAsientoFormSet(request.POST)
		if form.is_valid() and crear_asiento_form_set.is_valid():
			form = form.save(commit=False)
			aux = 0
			aux1 = 0
			aux2 = 0
			e = self.model.objects.filter(tipo_trans='E').exclude(anulado=True, cerrado=True).count()
			i = self.model.objects.filter(tipo_trans='I').exclude(anulado=True, cerrado=True).count()
			t = self.model.objects.filter(tipo_trans='T').exclude(anulado=True, cerrado=True).count()
			if (e == None) & (form.tipo_trans=='E'):
				form.numero_comp = 1
			elif (e!= None) & (form.tipo_trans=='E'):
				form.numero_comp = e + 1
			if (i == None) & (form.tipo_trans=='I'):
				form.numero_comp = 1
			elif (i!= None) & (form.tipo_trans=='I'):
				form.numero_comp = i + 1
			if (t == None) & (form.tipo_trans=='T'):
				form.numero_comp = 1
			elif (t!= None) & (form.tipo_trans=='T'):
				form.numero_comp = t + 1
			form.save()
			crear_asiento_form_set.instance = form
			formset = crear_asiento_form_set.save(commit=False)
			for f in formset:
				if (f.debe_b != None) & (f.debe_b != 0):
					f.haber_b = 0.00
					f.haber_d = 0.00
					f.auxiliar = 1

				elif (f.haber_b != None) & (f.haber_b != 0):
					f.debe_b = 0.00
					f.debe_d = 0.00
					f.auxiliar = 2
					
			crear_asiento_form_set.save()
			try:
				for i in selec:
					capital = Capitalizacion.objects.get(id=i)
					capital.estado = True
					capital.save()
			except Exception as e:
				pass
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

class updateAsiento(UpdateView):
	model_pk = nivel_cinco
	model = encabezadoas
	form_class = crearEncabezadoForm
	template_name = 'contabilidad/nuevoasiento.html'
	success_url = 'contabilidad:listaasiento'

	def get_context_data(self, **kwargs):
		context = super(updateAsiento, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'form' and 'car' in context:
			context['form'] = self.form_class()
			context['car'] = encabezadoas.objects.get(id=pk1)
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().encabezadoas.pk})

	def get(self, request, *args, **kwargs):
		self.object = self.get_object()
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_asiento_form_set = updateAsientoFormSet(instance=self.object)
		return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set))

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		crear_asiento_form_set = updateAsientoFormSet(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid() and crear_asiento_form_set.is_valid():
			self.object = form.save()
			crear_asiento_form_set.instance = self.object
			formset = crear_asiento_form_set.save(commit=False)
			for obj in crear_asiento_form_set.deleted_objects:
				obj.delete()

			for f in formset:
				if (f.debe_b != None) & (f.debe_b != 0):
					f.haber_b = 0.00
					f.haber_d = 0.00
					f.auxiliar = 1

				elif (f.haber_b != None) & (f.haber_b != 0):
					f.debe_b = 0.00
					f.debe_d = 0.00
					f.auxiliar = 2

			crear_asiento_form_set.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_asiento_form_set=crear_asiento_form_set, error=crear_asiento_form_set.non_form_errors()))

class listaAsiento(CreateView, ListView):
	model = encabezadoas
	form_class = searchForm
	template_name = 'contabilidad/listaasiento.html'
	def get_context_data(self, **kwargs):
		context = super(listaAsiento, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context

	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))

	def get_queryset(self):
		object_list = self.model.objects.filter(anulado=False)
		return object_list

class listaAsientoAnulado(listaAsiento):
	model = encabezadoas
	form_class = searchForm
	template_name = 'contabilidad/listaasiento.html'

	def get_queryset(self):
		object_list = self.model.objects.filter(anulado=True)
		return object_list

class listaAsientoForeign(CreateView, ListView):
	model = asiento
	form_class = searchForm
	template_name = 'contabilidad/listaasientof.html'
	def get_context_data(self, *args, **kwargs):
		context = super(listaAsientoForeign, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['pk1'] = pk1
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		pk1 = self.kwargs['pk']
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(encabezadoas = pk1).filter(Q(nivel_cinco__descripcion__icontains = campo) | Q(glosa_indiv__icontains = campo) | Q(banco__icontains = campo) | Q(numero_cheque__icontains = campo) | Q(nivel_cinco__codigo__icontains = campo))
		else:
			object_list = self.model.objects.filter(encabezadoas = pk1).order_by('id')
		return object_list

class borrarCuenta(DeleteView):
	model = asiento
	success_url = 'contabilidad:listaasientof'
	def get_success_url(self):
		pk1 = self.kwargs['pk']
		dat = self.model.objects.get(id=pk1).encabezadoas
		return reverse_lazy(self.success_url , kwargs = {'pk': dat.pk})

class listaAsientoForeignT(CreateView, ListView):
	model = asiento
	form_class = searchForm
	template_name = 'contabilidad/listaasientoft.html'
	def get_context_data(self, *args, **kwargs):
		context = super(listaAsientoForeignT, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))

class asientoAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return nivel_cinco.objects.none()
		qs = nivel_cinco.objects.all().order_by('codigo')
		if self.q:
			qs = qs.filter(Q(codigo__istartswith=self.q) | Q(descripcion__icontains=self.q))
		return qs

class nivel_unoAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return nivel_uno.objects.none()
		qs = nivel_uno.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(Q(id=self.q) | Q(descripcion__icontains=self.q))
		return qs

class nivel_dosAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return nivel_dos.objects.none()
		qs = nivel_dos.objects.all().order_by('codigo')
		nivel_uno = self.forwarded.get('nivel_uno', None)
		if nivel_uno:
			qs = qs.filter(nivel_uno=nivel_uno).order_by('codigo')
		if self.q:
			qs = qs.filter(Q(codigo__istartswith=self.q) | Q(descripcion__icontains=self.q))
		return qs

class nivel_tresAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return nivel_tres.objects.none()
		qs = nivel_tres.objects.all().order_by('codigo')
		nivel_dos = self.forwarded.get('nivel_dos', None)
		if nivel_dos:
			qs = qs.filter(nivel_dos=nivel_dos).order_by('codigo')
		if self.q:
			qs = qs.filter(Q(codigo__istartswith=self.q) | Q(descripcion__icontains=self.q))
		return qs

class nivel_cuatroAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return nivel_cuatro.objects.none()
		qs = nivel_cuatro.objects.all().order_by('codigo')
		nivel_tres = self.forwarded.get('nivel_tres', None)
		if nivel_tres:
			qs = qs.filter(nivel_tres=nivel_tres)
		if self.q:
			qs = qs.filter(Q(codigo__istartswith=self.q) | Q(descripcion__icontains=self.q))
		return qs

def cuenta(request):
	data = serializers.serialize("json", nivel_cinco.objects.filter(Q(codigo='1102080001') | Q(codigo='5201300001') | Q(codigo='4101010001') | Q(codigo='2101200001') | Q(codigo='2101210001')).order_by('codigo'))
	return HttpResponse(data, content_type="application/json")


class crearCuenta1(CreateView):
	form_class = crearPlanCuenta1Form
	template_name = 'contabilidad/nuevacuenta1.html'
	success_url = reverse_lazy('contabilidad:listacuenta1')

class updateCuenta1(UpdateView):
	model = nivel_uno
	form_class = crearPlanCuenta1Form
	template_name = 'contabilidad/nuevacuenta1.html'
	success_url = reverse_lazy('contabilidad:listacuenta1')

class listaCuenta1(CreateView, ListView):
	model = nivel_uno
	form_class = searchForm
	template_name = 'contabilidad/listacuenta1.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta1, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(codigo = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearCuenta2(CreateView):
	model = nivel_dos
	model_pk = nivel_uno
	form_class = crearPlanCuenta2Form
	template_name = 'contabilidad/nuevacuenta2.html'
	success_url = 'contabilidad:listacuenta2'

	def get_context_data(self, **kwargs):
		context = super(crearCuenta2, self).get_context_data(**kwargs)
		if 'form' and 'niv' in context:
			context['niv'] = nivel_uno.objects.get(id=pk1)
		return context

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			cont = self.model.objects.filter(nivel_uno=form.nivel_uno.id).count()
			if cont == 0:
				cont = 1
				aux = str(form.nivel_uno.id) + str(cont)
				form.codigo = aux
			else:
				aux = str(form.nivel_uno.id) + str(cont+1)
				form.codigo = aux
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateCuenta2(UpdateView):
	model = nivel_dos
	form_class = crearPlanCuenta2Form
	template_name = 'contabilidad/nuevacuenta2.html'
	success_url = reverse_lazy('contabilidad:listacuenta2')


class listaCuenta2(CreateView, ListView):
	model = nivel_dos
	form_class = searchForm
	template_name = 'contabilidad/listacuenta2.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta2, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(codigo = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearCuenta3(CreateView):
	model = nivel_tres
	model_pk = nivel_dos
	form_class = crearPlanCuenta3Form
	template_name = 'contabilidad/nuevacuenta3.html'
	success_url = 'contabilidad:listacuenta3'

	def get_context_data(self, **kwargs):
		context = super(crearCuenta3, self).get_context_data(**kwargs)

		return context
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			cont = self.model.objects.filter(nivel_dos=form.nivel_dos.id).count()
			if cont == 0:
				cont = 1
				aux = form.nivel_dos.codigo + str(cont).zfill(2)
				form.codigo = aux
			else:
				aux = form.nivel_dos.codigo + str(cont+1).zfill(2)
				form.codigo = aux
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateCuenta3(UpdateView):
	model = nivel_tres
	form_class = crearPlanCuenta3Form
	template_name = 'contabilidad/nuevacuenta3.html'
	success_url = reverse_lazy('contabilidad:listacuenta3')

class listaCuenta3(CreateView, ListView):
	model = nivel_tres
	form_class = searchForm
	template_name = 'contabilidad/listacuenta3.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta3, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(codigo = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearCuenta4(CreateView):
	model = nivel_cuatro
	model_pk = nivel_tres
	form_class = crearPlanCuenta4Form
	template_name = 'contabilidad/nuevacuenta4.html'
	success_url = 'contabilidad:listacuenta4'

	def get_context_data(self, **kwargs):
		context = super(crearCuenta4, self).get_context_data(**kwargs)

		return context
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			cont = self.model.objects.filter(nivel_tres=form.nivel_tres.id).count()
			if cont == 0:
				cont = 1
				aux = form.nivel_tres.codigo + str(cont).zfill(2)
				form.codigo = aux
			else:
				aux = form.nivel_tres.codigo + str(cont+1).zfill(2)
				form.codigo = aux
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateCuenta4(UpdateView):
	model = nivel_cuatro
	form_class = updatePlanCuenta4Form
	template_name = 'contabilidad/nuevacuenta4.html'
	success_url = reverse_lazy('contabilidad:listacuenta4')

class listaCuenta4(CreateView, ListView):
	model = nivel_cuatro
	form_class = searchForm
	template_name = 'contabilidad/listacuenta4.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta4, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(codigo = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearCuenta5(CreateView):
	model = nivel_cinco
	model_pk = nivel_cuatro
	form_class = crearPlanCuenta5Form
	template_name = 'contabilidad/nuevacuenta5.html'
	success_url = 'contabilidad:listacuenta5'

	def get_context_data(self, **kwargs):
		context = super(crearCuenta5, self).get_context_data(**kwargs)

		return context
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			cont = self.model.objects.filter(nivel_cuatro=form.nivel_cuatro.id).count()
			if cont == 0:
				cont = 1
				aux = form.nivel_cuatro.codigo + str(cont).zfill(4)
				form.codigo = aux
			else:
				aux = form.nivel_cuatro.codigo + str(cont+1).zfill(4)
				form.codigo = aux
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateCuenta5(UpdateView):
	model = nivel_cinco
	form_class = updatePlanCuenta5Form
	template_name = 'contabilidad/nuevacuenta5.html'
	success_url = reverse_lazy('contabilidad:listacuenta5')

class listaCuenta5(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/listacuenta5.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(codigo = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaPlanCuenta(CreateView, ListView):
	model = nivel_uno
	form_class = searchForm
	template_name = 'contabilidad/listaplancuentas.html'
	def get_context_data(self, **kwargs):
		context = super(listaPlanCuenta, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(codigo = campo)
		else:
			object_list = self.model.objects.all()
		return object_list

class listaCuenta2T(CreateView, ListView):
	model = nivel_dos
	form_class = searchForm
	template_name = 'contabilidad/listacuenta2t.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta2T, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta3T(CreateView, ListView):
	model = nivel_tres
	form_class = searchForm
	template_name = 'contabilidad/listacuenta3t.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta3T, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta4T(CreateView, ListView):
	model = nivel_cuatro
	form_class = searchForm
	template_name = 'contabilidad/listacuenta4t.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta4T, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta5TT(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/listacuenta5tt.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5TT, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta5T(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/listacuenta5t.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5T, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta5TB(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/cuentas_balance.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5TB, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta5TE(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/estado_ev.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5TE, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta5ER(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/estado_resultados.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5ER, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCuenta5SS(CreateView, ListView):
	model = nivel_cinco
	form_class = searchForm
	template_name = 'contabilidad/suma_saldo.html'
	def get_context_data(self, **kwargs):
		context = super(listaCuenta5SS, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

#Formulario para sacar reporte de retenciones
class listaAsientoRetencion(CreateView, ListView):
	model = encabezadoas
	form_class = searchForm
	template_name = 'contabilidad/retencion.html'
	def get_context_data(self, **kwargs):
		context = super(listaAsientoRetencion, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(codigo__istartswith = campo) | Q(descripcion__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearFactura(CreateView):
	model_pk = asiento
	form_class = crearFacturaForm
	template_name = 'contabilidad/nuevafactura.html'
	success_url = 'contabilidad:listaasientof'
	success_url2 = 'contabilidad:listafactura'

	def get_context_data(self, **kwargs):
		pk1 = self.kwargs['pk']
		try:
			self.asi = self.model_pk.objects.get(pk=pk1)
		except Exception as e:
			self.asi = None
		context = super(crearFactura, self).get_context_data(**kwargs)
		if 'form' in context:
			context['pk1'] = pk1
		return context

	def get_form_kwargs(self):
		kwargs = super(crearFactura, self).get_form_kwargs()
		kwargs.update({'asi':self.asi})
		return kwargs

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			pk1 = self.kwargs['pk']
			monto = form.cleaned_data['importe_fac']
			form = form.save(commit=False)
			factura = asiento.objects.get(id=pk1)
			factura.estado = True
			factura.save()
			form.asiento = self.model_pk.objects.get(id=pk1)
			form.save()
			dat = self.model_pk.objects.get(id=pk1).encabezadoas
			if(monto >= 50000):
				return HttpResponseRedirect(reverse_lazy(self.success_url2, kwargs={'pk': pk1}))
			else:
				return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs={'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateFactura(UpdateView):
	model = Factura
	form_class = crearFacturaForm
	template_name = 'contabilidad/nuevafactura.html'
	success_url = 'contabilidad:listafactura'

	def get_context_data(self, **kwargs):
		context = super(updateFactura, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'car' not in context:
			context['pk1'] = pk1
			context['car'] = self.model.objects.get(id=pk1)
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().asiento.pk})
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid():
			dat = self.model.objects.get(id=pk1).asiento
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs = {'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class listaFactura(CreateView, ListView):
	model = Factura
	form_class = searchForm
	template_name = 'contabilidad/listafactura.html'
	def get_context_data(self, **kwargs):
		context = super(listaFactura, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
			context['pk1'] = pk1
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		pk1 = self.kwargs['pk']
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.filter(asiento = pk1).order_by('id')
		return object_list

class crearFacturaV(CreateView):
	model_pk = asiento
	form_class = crearFacturaVForm
	template_name = 'contabilidad/nuevafactura_v.html'
	success_url = 'contabilidad:listaasientof'

	def get_context_data(self, **kwargs):
		pk1 = self.kwargs['pk']
		try:
			self.asi = self.model_pk.objects.get(pk=pk1)
		except Exception as e:
			self.asi = None
		context = super(crearFacturaV, self).get_context_data(**kwargs)
		if 'form' in context:
			context['pk1'] = pk1
		return context

	def get_form_kwargs(self):
		kwargs = super(crearFacturaV, self).get_form_kwargs()
		kwargs.update({'asi':self.asi})
		return kwargs

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			pk1 = self.kwargs['pk']
			monto = form.cleaned_data['importe_fac']
			form = form.save(commit=False)
			factura = asiento.objects.get(id=pk1)
			factura.estado = True
			factura.save()
			form.asiento = self.model_pk.objects.get(id=pk1)
			form.save()
			dat = self.model_pk.objects.get(id=pk1).encabezadoas
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs={'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateFacturaV(UpdateView):
	model = Factura_V
	form_class = crearFacturaVForm
	template_name = 'contabilidad/nuevafactura_v.html'
	success_url = 'contabilidad:listafactura_v'

	def get_context_data(self, **kwargs):
		context = super(updateFacturaV, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'car' not in context:
			context['pk1'] = pk1
			context['car'] = self.model.objects.get(id=pk1)
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().asiento.pk})
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid():
			dat = self.model.objects.get(id=pk1).asiento
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs = {'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class listaFacturaV(CreateView, ListView):
	model = Factura_V
	form_class = searchForm
	template_name = 'contabilidad/listafactura_v.html'
	def get_context_data(self, **kwargs):
		context = super(listaFacturaV, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
			context['pk1'] = pk1
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		pk1 = self.kwargs['pk']
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.filter(asiento = pk1).order_by('id')
		return object_list

class crearFacturaBol(CreateView):
	model_pk = asiento
	form_class = crearFacturaBolForm
	template_name = 'contabilidad/nuevafacturabol.html'
	success_url = 'contabilidad:listaasientof'

	def get_context_data(self, **kwargs):
		pk1 = self.kwargs['pk']
		try:
			self.asi = self.model_pk.objects.get(pk=pk1)
		except Exception as e:
			self.asi = None
		context = super(crearFacturaBol, self).get_context_data(**kwargs)
		if 'form' in context:
			context['pk1'] = pk1
		return context

	def get_form_kwargs(self):
		kwargs = super(crearFacturaBol, self).get_form_kwargs()
		kwargs.update({'asi':self.asi})
		return kwargs

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			pk1 = self.kwargs['pk']
			monto = form.cleaned_data['importe_fac']
			form = form.save(commit=False)
			factura = asiento.objects.get(id=pk1)
			factura.estado = True
			factura.save()
			form.asiento = self.model_pk.objects.get(id=pk1)
			form.save()
			dat = self.model_pk.objects.get(id=pk1).encabezadoas
			if(monto >= 50000):
				return HttpResponseRedirect(reverse_lazy(self.success_url2, kwargs={'pk': pk1}))
			else:
				return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs={'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateFacturaBol(UpdateView):
	model = Factura
	form_class = crearFacturaBolForm
	template_name = 'contabilidad/nuevafacturabol.html'
	success_url = 'contabilidad:listafacturabol'

	def get_context_data(self, **kwargs):
		context = super(updateFacturaBol, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'car' not in context:
			context['pk1'] = pk1
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().asiento.pk})
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		print(pk1)
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid():
			dat = self.model.objects.get(id=pk1).asiento
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs = {'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class listaFacturaBol(CreateView, ListView):
	model = Factura
	form_class = searchForm
	template_name = 'contabilidad/listafacturabol.html'
	def get_context_data(self, **kwargs):
		context = super(listaFacturaBol, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
			context['pk1'] = pk1
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		pk1 = self.kwargs['pk']
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.filter(asiento = pk1).order_by('id')
		return object_list

class crearFacturaImp(CreateView):
	model_pk = asiento
	form_class = crearFacturaImpForm
	template_name = 'contabilidad/nuevafacturaimp.html'
	success_url = 'contabilidad:listaasientof'

	def get_context_data(self, **kwargs):
		pk1 = self.kwargs['pk']
		try:
			self.asi = self.model_pk.objects.get(pk=pk1)
		except Exception as e:
			self.asi = None
		context = super(crearFacturaImp, self).get_context_data(**kwargs)
		if 'form' in context:
			context['pk1'] = pk1
		return context

	def get_form_kwargs(self):
		kwargs = super(crearFacturaImp, self).get_form_kwargs()
		kwargs.update({'asi':self.asi})
		return kwargs

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		#print(form.errors)
		if form.is_valid():
			pk1 = self.kwargs['pk']
			monto = form.cleaned_data['importe_fac']
			form = form.save(commit=False)
			factura = asiento.objects.get(id=pk1)
			factura.estado = True
			factura.save()
			form.asiento = self.model_pk.objects.get(id=pk1)
			form.save()
			dat = self.model_pk.objects.get(id=pk1).encabezadoas
			if(monto >= 50000):
				return HttpResponseRedirect(reverse_lazy(self.success_url2, kwargs={'pk': pk1}))
			else:
				return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs={'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateFacturaImp(UpdateView):
	model = Factura
	form_class = crearFacturaImpForm
	template_name = 'contabilidad/nuevafacturaimp.html'
	success_url = 'contabilidad:listafacturaimp'

	def get_context_data(self, **kwargs):
		context = super(updateFacturaImp, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'car' not in context:
			context['pk1'] = pk1
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().asiento.pk})
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		print(pk1)
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid():
			dat = self.model.objects.get(id=pk1).asiento
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs = {'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class listaFacturaImp(CreateView, ListView):
	model = Factura
	form_class = searchForm
	template_name = 'contabilidad/listafacturaimp.html'
	def get_context_data(self, **kwargs):
		context = super(listaFacturaImp, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
			context['pk1'] = pk1
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		pk1 = self.kwargs['pk']
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.filter(asiento = pk1).order_by('id')
		return object_list

class listaFacturaT(CreateView, ListView):
	model = Factura
	form_class = searchForm
	template_name = 'contabilidad/listafacturat.html'
	def get_context_data(self, **kwargs):
		context = super(listaFacturaT, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(asiento__encabezadoas__numero_comp=campo) | Q(nro_fact__icontains=campo) | Q(fecha_fac__icontains=campo) | Q(nit_prov__icontains=campo) | Q(nro_aut__icontains=campo) | Q(cod_cont__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('asiento__encabezadoas__numero_comp')
		return object_list

class listaFacturaVT(CreateView, ListView):
	model = Factura_V
	form_class = searchForm
	template_name = 'contabilidad/listafacturavt.html'
	def get_context_data(self, **kwargs):
		context = super(listaFacturaVT, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(Q(asiento__encabezadoas__numero_comp=campo) | Q(nro_fact__icontains=campo) | Q(fecha_fac__icontains=campo) | Q(nit_prov__icontains=campo) | Q(nro_aut__icontains=campo) | Q(cod_cont__icontains=campo))
		else:
			object_list = self.model.objects.all().order_by('asiento__encabezadoas__numero_comp')
		return object_list

class crearCotizacion_tc(CreateView):
	form_class = crearCotizacionForm
	template_name = 'contabilidad/nuevacotizacion_tc.html'
	success_url = reverse_lazy('contabilidad:listacotizacion_tc')

class updateCotizacion_tc(UpdateView):
	model = Cotizacion
	form_class = crearCotizacionForm
	template_name = 'contabilidad/nuevacotizacion_tc.html'
	success_url = reverse_lazy('contabilidad:listacotizacion_tc')

class listaCotizacion_tc(CreateView, ListView):
	model = Cotizacion
	form_class = searchForm
	template_name = 'contabilidad/listacotizacion_tc.html'
	def get_context_data(self, **kwargs):
		context = super(listaCotizacion_tc, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaCotizacionM(CreateView, ListView):
	model = Cotizacion
	form_class = searchForm
	template_name = 'contabilidad/listacotizacion_m.html'
	def get_context_data(self, **kwargs):
		context = super(listaCotizacionM, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class uploadCotizacionTC(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacotizacion_tc')
	template_name = 'contabilidad/cotizacion_tcupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			for i in range(0, sheet.nrows):
				f = sheet.cell_value(i, 0)
				f1 = datetime.datetime(*xlrd.xldate_as_tuple(f, book.datemode))
				f2 = f1.strftime("%Y-%m-%d")
				aux.append(f2)
			t_c = sheet.col_values(1, 0)
			u = sheet.col_values(2, 0)
			for fecha, tipo_c, ufv in zip(aux, t_c, u):
				cotizacion, created = Cotizacion.objects.get_or_create(fecha=fecha, tipo_c=tipo_c, ufv=ufv)
				cotizacion.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class uploadCotizacionUFV(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacotizacion_ufv')
	template_name = 'contabilidad/cotizacion_ufvupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			for i in range(0, sheet.nrows):
				f = sheet.cell_value(i, 0)
				f1 = datetime.datetime(*xlrd.xldate_as_tuple(f, book.datemode))
				f2 = f1.strftime("%Y-%m-%d")
				aux.append(f2)
			t_c = sheet.col_values(1, 0)

			for fecha, tipo_c in zip(aux, t_c):
				cotizacion, created = Cotizacion_ufv.objects.get_or_create(fecha=fecha, ufv=tipo_c)
				cotizacion.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class uploadNiveluno(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacuenta1')
	template_name = 'contabilidad/nivel_unoupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			descr = sheet.col_values(0, 0)
			for descripcion in descr:
				cuenta, created = nivel_uno.objects.get_or_create(descripcion=descripcion)
				cuenta.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class uploadNiveldos(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacuenta2')
	template_name = 'contabilidad/nivel_dosupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			aux1 = []
			uno = sheet.col_values(0, 0)
			for i in uno:
				aux.append(int(i))
			cod = sheet.col_values(1, 0)
			for j in cod:
				aux1.append(int(j))
			descr = sheet.col_values(2, 0)
			for nivel_uno, codigo, descripcion in zip(aux, aux1, descr):
				cuenta, created = nivel_dos.objects.get_or_create(nivel_uno_id=nivel_uno, codigo=codigo, descripcion=descripcion)
				cuenta.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class uploadNiveltres(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacuenta3')
	template_name = 'contabilidad/nivel_tresupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			aux1 = []
			aux2 = []
			uno = sheet.col_values(0, 0)
			dos = sheet.col_values(1, 0)
			for i in uno:
				aux.append(int(i))
			for j in dos:
				aux1.append(int(j))
			cod = sheet.col_values(2, 0)
			for k in cod:
				aux2.append(int(k))
			descr = sheet.col_values(3, 0)
			for nivel_uno, nivel_dos, codigo, descripcion in zip(aux, aux1, aux2, descr):
				cuenta, created = nivel_tres.objects.get_or_create(nivel_uno_id=nivel_uno, nivel_dos_id=nivel_dos, codigo=codigo, descripcion=descripcion)
				cuenta.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class uploadNivelcuatro(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacuenta4')
	template_name = 'contabilidad/nivel_cuatroupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			aux1 = []
			aux2 = []
			aux3 = []
			uno = sheet.col_values(0, 0)
			dos = sheet.col_values(1, 0)
			tres = sheet.col_values(2, 0)
			for i in uno:
				aux.append(int(i))
			for j in dos:
				aux1.append(int(j))
			for k in dos:
				aux2.append(int(k))
			cod = sheet.col_values(3, 0)
			for l in cod:
				aux3.append(int(l))
			descr = sheet.col_values(4, 0)
			for nivel_uno, nivel_dos, nivel_tres, codigo, descripcion in zip(aux, aux1, aux2, aux3, descr):
				cuenta, created = nivel_cuatro.objects.get_or_create(nivel_uno_id=nivel_uno, nivel_dos_id=nivel_dos, nivel_tres_id=nivel_tres, codigo=codigo, descripcion=descripcion)
				cuenta.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class uploadNivelcinco(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacuenta5')
	template_name = 'contabilidad/nivel_cincoupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			aux1 = []
			aux2 = []
			aux3 = []
			aux4 = []
			uno = sheet.col_values(0, 0)
			dos = sheet.col_values(1, 0)
			tres = sheet.col_values(2, 0)
			cuatro = sheet.col_values(3, 0)
			for i in uno:
				aux.append(int(i))
			for j in dos:
				aux1.append(int(j))
			for k in dos:
				aux2.append(int(k))
			for l in dos:
				aux3.append(int(l))
			cod = sheet.col_values(4, 0)
			for m in cod:
				aux4.append(int(m))
			descr = sheet.col_values(5, 0)
			for nivel_uno, nivel_dos, nivel_tres, nivel_cuatro, codigo, descripcion in zip(aux, aux1, aux2, aux3, aux4, descr):
				cuenta, created = nivel_cinco.objects.get_or_create(nivel_uno_id=nivel_uno, nivel_dos_id=nivel_dos, nivel_tres_id=nivel_tres, nivel_cuatro_id=nivel_cuatro, codigo=codigo, descripcion=descripcion)
				cuenta.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class crearProyecto(CreateView):
	form_class = crearProyectoForm
	template_name = 'contabilidad/nuevoproyecto.html'
	success_url = reverse_lazy('contabilidad:listaproyecto')

class updateProyecto(UpdateView):
	model = Proyecto
	form_class = crearProyectoForm
	template_name = 'contabilidad/nuevoproyecto.html'
	success_url = reverse_lazy('contabilidad:listaproyecto')

class listaProyecto(CreateView, ListView):
	model = Proyecto
	form_class = searchForm
	template_name = 'contabilidad/listaproyecto.html'
	def get_context_data(self, **kwargs):
		context = super(listaProyecto, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class listaProveedorMov(CreateView, ListView):
	model = encabezadoas
	form_class = searchForm
	template_name = 'contabilidad/lista_movimiento_proov.html'
	def get_context_data(self, **kwargs):
		context = super(listaProveedorMov, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(comprobante_esc__proovedor__nombre_proov__icontains = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearBancarizacion(CreateView):
	model_pk = Factura
	form_class = crearBancarizacionForm
	template_name = 'contabilidad/nuevabancarizacion.html'
	success_url = 'contabilidad:listaasientof'

	def get_context_data(self, **kwargs):
		pk1 = self.kwargs['pk']
		try:
			self.fac = self.model_pk.objects.get(pk=pk1)
		except Exception as e:
			self.fac = None
		context = super(crearBancarizacion, self).get_context_data(**kwargs)
		if 'form' in context:
			context['pk1'] = pk1
		return context

	def get_form_kwargs(self):
		kwargs = super(crearBancarizacion, self).get_form_kwargs()
		kwargs.update({'fac':self.fac})
		return kwargs

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			pk1 = self.kwargs['pk']
			form = form.save(commit=False)
			form.factura = self.model_pk.objects.get(id=pk1)
			form.save()
			dat = self.model_pk.objects.get(id=pk1).asiento.encabezadoas
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs={'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateBancarizacion(UpdateView):
	model = bancarizacion
	form_class = crearBancarizacionForm
	template_name = 'contabilidad/nuevabancarizacion.html'
	success_url = 'contabilidad:listabancarizacion'

	def get_context_data(self, **kwargs):
		context = super(updateBancarizacion, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'car' not in context:
			context['pk1'] = pk1
			context['car'] = self.model.objects.get(id=pk1)
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().factura.pk})
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid():
			dat = self.model.objects.get(id=pk1).factura
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url, kwargs = {'pk': dat.pk}))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class listaBancarizacion(CreateView, ListView):
	model = bancarizacion
	form_class = searchForm
	template_name = 'contabilidad/listabancarizacion.html'
	def get_context_data(self, **kwargs):
		context = super(listaBancarizacion, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' and 'car' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(self.request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				campo = form.cleaned_data['search']
		if(campo):
			object_list = self.model.objects.filter(id = campo)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class crearGestion(CreateView):
	form_class = crearGestion
	template_name = 'contabilidad/nuevagestion.html'
	success_url = reverse_lazy('contabilidad:listagestion')

class listaGestion(ListView):
	model = gestion
	template_name = 'contabilidad/listagestion.html'

class updateCierreGestion(UpdateView):
	model = gestion
	form_class = crearCierreGestionForm
	template_name = 'contabilidad/nuevocierregestion.html'
	success_url = 'contabilidad:listagestion'

	def get_context_data(self, **kwargs):
		context = super(updateCierreGestion, self).get_context_data(**kwargs)
		pk1=self.kwargs['pk']
		cerrado = self.model.objects.get(id=pk1)
		if 'form' and 'cerrado' not in context :
			context['form'] = self.form_class()
			context['cerrado'] = cerrado
		return context

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid():
			form.save(commit=False)
			if form.cleaned_data['cerrado'] == True:
				year = self.model.objects.get(id=pk1).year
				asientos = encabezadoas.objects.filter(fecha__year=year).exclude(Q(fecha__month=10) | Q(fecha__month=11) | Q(fecha__month=12)).update(cerrado=True)
				ingresos = encingresoitem.objects.filter(fecha__year=year).exclude(Q(fecha__month=10) | Q(fecha__month=11) | Q(fecha__month=12)).update(cerrado=True)
				salidas = encsalidaitem.objects.filter(fecha__year=year).exclude(Q(fecha__month=10) | Q(fecha__month=11) | Q(fecha__month=12)).update(cerrado=True)
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class renumerarComprobantes(FormView):
	form_class = renumerarComprobantesForm
	success_url = 'contabilidad:listaasiento'
	template_name = 'contabilidad/renum_comp.html'

	def get_context_data(self, **kwargs):
		context = super(renumerarComprobantes, self).get_context_data(**kwargs)
		if 'form' not in context:
			context['form'] = self.form_class()
		return context

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST)
		if form.is_valid():
			ingreso = encabezadoas.objects.filter(tipo_trans='I').exclude(anulado=True)
			egreso = encabezadoas.objects.filter(tipo_trans='E').exclude(anulado=True)
			traspaso = encabezadoas.objects.filter(tipo_trans='T').exclude(anulado=True)
			cont_ing = 1
			cont_egr = 1
			cont_trasp = 1

			for i in ingreso:
				i.numero_comp=cont_ing
				i.save()
				cont_ing += 1

			for e in egreso:
				e.numero_comp=cont_egr
				e.save()
				cont_egr += 1

			for t in traspaso:
				t.numero_comp=cont_trasp
				t.save()
				cont_trasp += 1

			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			self.render_to_response(self.get_context_data(form=form))


# Reportes
class listaLibroMayor(CreateView, ListView):
	model = asiento
	form_class = searchFormLib
	template_name = 'contabilidad/listalibrom.html'
	def get_context_data(self, **kwargs):
		context = super(listaLibroMayor, self).get_context_data(**kwargs)
		parametro = self.request.GET.copy()
		if parametro.has_key('page'):
			del parametro['page']
		context['parametro'] = parametro
		if 'form' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list, form=form))
	def get_queryset(self):
		campo = self.request.GET.get('search')
		campo1 = self.request.GET.get('search1')
		campo2 = self.request.GET.get('search2')
		if(campo2):
			ini = campo2[0]
			aux2 = []
			aux3 = []
			aux4 = 0
			aux5 = 0
			aux7 = 0
			aux8 = 0
			saldo = []
			saldo_aux = []
			b = 0
			niv_cinco = nivel_cinco.objects.get(Q(codigo=campo2) | Q(descripcion=campo2))
			if (niv_cinco.nivel_uno.id == 2) or (niv_cinco.nivel_uno.id == 4):
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(campo, campo1)) & Q(Q(nivel_cinco__codigo=campo2)) | Q(nivel_cinco__descripcion=campo2)).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.haber_b - i.debe_b
					aux1 = i.haber_d - i.debe_d
					aux2.append(aux)
					aux3.append(aux1)
					aux4 = aux4 + i.debe_b
					aux5 = aux5 + i.haber_b
					aux7 = aux7 + i.debe_d
					aux8 = aux8 + i.haber_d

				for f in asiento.objects.filter(Q(encabezadoas__fecha__range=(campo, campo1)) & Q(Q(nivel_cinco__codigo=campo2)) | Q(nivel_cinco__descripcion=campo2)).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = len(saldo)
					if aux == 0:
						saldo.append(aux2[b])
						saldo_aux.append(aux3[b])
						b += 1
					elif b <= aux:
						saldo.append(saldo[b-1] + aux2[b])
						saldo_aux.append(saldo_aux[b-1] + aux3[b])
						b += 1
			else:
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(campo, campo1)) & Q(Q(nivel_cinco__codigo=campo2)) | Q(nivel_cinco__descripcion=campo2)).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.debe_b - i.haber_b
					aux1 = i.debe_d - i.haber_d
					aux2.append(aux)
					aux3.append(aux1)
					aux4 = aux4 + i.debe_b
					aux5 = aux5 + i.haber_b
					aux7 = aux7 + i.debe_d
					aux8 = aux8 + i.haber_d
				for f in asiento.objects.filter(Q(encabezadoas__fecha__range=(campo, campo1)) & Q(Q(nivel_cinco__codigo=campo2)) | Q(nivel_cinco__descripcion=campo2)).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = len(saldo)
					if aux == 0:
						saldo.append(aux2[b])
						saldo_aux.append(aux3[b])
						b += 1
					elif b <= aux:
						saldo.append(saldo[b-1] + aux2[b])
						saldo_aux.append(saldo_aux[b-1] + aux3[b])
						b += 1
			lista = zip(asiento.objects.filter(Q(encabezadoas__fecha__range=(campo, campo1)) & Q(Q(nivel_cinco__codigo=campo2)) | Q(nivel_cinco__descripcion=campo2)).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'), saldo, saldo_aux)
			object_list = lista
		else:
			lista = zip(self.model.objects.all().exclude(encabezadoas__anulado=True).order_by('id'), " ", " ")
			object_list = lista
		return object_list

class generarComprobante(PDFTemplateView):
	template_name = "contabilidad/comprobante.html"
	
	def get_context_data(self, **kwargs):
		context = super(generarComprobante, self).get_context_data(**kwargs)
		pk = self.kwargs.get('pk')
		try:
			context['comp'] = encabezadoas.objects.get(id=pk)
			context['user'] = self.request.user
		except:
			raise Http404
		return context

class libroMayorP(PDFTemplateView):
	template_name = "contabilidad/libro_mayor_p.html"
	def get_context_data(self, **kwargs):
		context = super(libroMayorP, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo2 = self.request.GET.get('campo2')
		campo3 = self.request.GET.get('campo3')
		campo4 = self.request.GET.get('campo4')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo2, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		ini = campo3[0]
		aux2 = []
		aux3 = []
		aux4 = 0
		aux5 = 0
		aux7 = 0
		aux8 = 0
		saldo = []
		saldo_aux = []
		b = 0
		niv_cinco = nivel_cinco.objects.get(Q(codigo=campo3) | Q(descripcion=campo3))
		if (niv_cinco.nivel_uno.id == 2) or (niv_cinco.nivel_uno.id == 4):
			if campo4 == 'Bs':
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.haber_b - i.debe_b
					aux1 = i.haber_d - i.debe_d
					aux2.append(aux)
					aux3.append(aux1)
					aux4 = aux4 + i.debe_b
					aux5 = aux5 + i.haber_b
			elif campo4 == 'Dls':
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.haber_b - i.debe_b
					aux1 = i.haber_d - i.debe_d
					aux2.append(aux)
					aux3.append(aux1)
					aux7 = aux7 + i.debe_d
					aux8 = aux8 + i.haber_d
			elif campo4 == 'BsDl':
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.haber_b - i.debe_b
					aux1 = i.haber_d - i.debe_d
					aux2.append(aux)
					aux3.append(aux1)
					aux4 = aux4 + i.debe_b
					aux5 = aux5 + i.haber_b
					aux7 = aux7 + i.debe_d
					aux8 = aux8 + i.haber_d
			for f in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
				aux = len(saldo)
				if aux == 0:
					saldo.append(aux2[b])
					saldo_aux.append(aux3[b])
					b += 1
				elif b <= aux:
					saldo.append(saldo[b-1] + aux2[b])
					saldo_aux.append(saldo_aux[b-1] + aux3[b])
					b += 1
		else:
			if campo4 == 'Bs':
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.debe_b - i.haber_b
					aux1 = i.debe_d - i.haber_d
					aux2.append(aux)
					aux3.append(aux1)
					aux4 = aux4 + i.debe_b
					aux5 = aux5 + i.haber_b
			elif campo4 == 'Dls':
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.debe_b - i.haber_b
					aux1 = i.debe_d - i.haber_d
					aux2.append(aux)
					aux3.append(aux1)
					aux7 = aux7 + i.debe_d
					aux8 = aux8 + i.haber_d
			elif campo4 == 'BsDl':
				for i in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
					aux = 0
					aux1 = 0
					aux = i.debe_b - i.haber_b
					aux1 = i.debe_d - i.haber_d
					aux2.append(aux)
					aux3.append(aux1)
					aux4 = aux4 + i.debe_b
					aux5 = aux5 + i.haber_b
					aux7 = aux7 + i.debe_d
					aux8 = aux8 + i.haber_d
			for f in asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'):
				aux = len(saldo)
				if aux == 0:
					saldo.append(aux2[b])
					saldo_aux.append(aux3[b])
					b += 1
				elif b <= aux:
					saldo.append(saldo[b-1] + aux2[b])
					saldo_aux.append(saldo_aux[b-1] + aux3[b])
					b += 1
		context.update({'aux4': aux4})
		context.update({'aux5': aux5})
		context.update({'aux7': aux7})
		context.update({'aux8': aux8})
		context.update({'campo': fecha1})
		context.update({'campo2': fecha2})
		context.update({'campo3': campo3})
		context.update({'campo4': campo4})
		lista = zip(asiento.objects.filter(Q(encabezadoas__fecha__range=(f1, f2)) & Q(Q(nivel_cinco__codigo=campo3) | Q(nivel_cinco__descripcion=campo3))).exclude(encabezadoas__anulado=True).order_by('encabezadoas__fecha'), saldo, saldo_aux)
		context['asi'] = lista
		return context

class balanceGen(PDFTemplateView):
	template_name = "contabilidad/balance_gen.html"
	def get_context_data(self, **kwargs):
		context = super(balanceGen, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo2 = self.request.GET.get('campo2')
		campo3 = self.request.GET.get('campo3')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo2, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		context.update({'campo': fecha1})
		context.update({'campo2': fecha2})
		context.update({'campo3': campo3})
		res = nivel_uno.objects.filter(Q(id__lte=3) & (Q(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__fecha__range=(f1, f2))) & Q(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__anulado=False)).distinct()
		context['niv'] = res
		context['f1'] = f1
		context['f2'] = f2

		return context

class estadoResultado(PDFTemplateView):
	template_name = "contabilidad/estado_r.html"
	def get_context_data(self, **kwargs):
		context = super(estadoResultado, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo2 = self.request.GET.get('campo2')
		campo3 = self.request.GET.get('campo3')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo2, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		aux = 0
		aux1 = 0
		aux2 = 0
		aux3 = 0
		context.update({'campo': fecha1})
		context.update({'campo2': fecha2})
		context.update({'campo3': campo3})
		res = nivel_uno.objects.filter(Q(id__gte=4) & (Q(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__fecha__range=(f1, f2))) & Q(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__anulado=False)).distinct()
		context['niv'] = res
		for i in nivel_uno.objects.filter(Q(id__gte=4) & (Q(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__fecha__range=(f1, f2))) & Q(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__anulado=False)).distinct():
			for j in i.nivel_dos_set.all():
				for k in j.nivel_tres_set.all():
					for l in k.nivel_cuatro_set.all():
						for m in l.nivel_cinco_set.all():
							for n in m.asiento_set.all():
								if(n.debe_b==None):
									n.debe_b=0
								elif(n.haber_b==None):
									n.haber_b=0
								elif(n.debe_d==None):
									n.debe_d=0
								elif(n.haber_d==None):
									n.haber_d=0
								aux = aux + n.debe_b
								aux1 = aux1 + n.haber_b
								aux2 = aux2 + n.debe_d
								aux3 = aux3 + n.haber_d
		utilidad_b = 0
		utilidad_d = 0
		utilidad_b = aux - aux1
		utilidad_d = aux2 - aux3
		
		context['f1'] = f1
		context['f2'] = f2
		context.update({'utilidad_b': utilidad_b})
		context.update({'utilidad_d': utilidad_d})
		return context

class sumasSaldos(PDFTemplateView):
	template_name = "contabilidad/sumas_saldos.html"
	def get_context_data(self, **kwargs):
		context = super(sumasSaldos, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo2 = self.request.GET.get('campo2')
		campo3 = self.request.GET.get('campo3')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo2, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		aux = 0
		aux1 = 0
		aux2 = 0
		aux3 = 0
		context.update({'campo': fecha1})
		context.update({'campo2': fecha2})
		context.update({'campo3': campo3})
		context['niv'] = nivel_uno.objects.filter(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__fecha__range=(f1, f2)).distinct()
		for i in nivel_uno.objects.filter(nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__fecha__range=(f1, f2)).distinct():
			for j in i.nivel_dos_set.all():
				for k in j.nivel_tres_set.all():
					for l in k.nivel_cuatro_set.all():
						for m in l.nivel_cinco_set.all():
							for n in m.asiento_set.all():
								if(n.debe_b==None):
									n.debe_b=0
								elif(n.haber_b==None):
									n.haber_b=0
								elif(n.debe_d==None):
									n.debe_d=0
								elif(n.haber_d==None):
									n.haber_d=0
								aux = aux + n.debe_b
								aux1 = aux1 + n.haber_b
								aux2 = aux2 + n.debe_d
								aux3 = aux3 + n.haber_d
		utilidad_b = 0
		utilidad_d = 0
		utilidad_b = aux - aux1
		utilidad_d = aux2 - aux3
		context.update({'utilidad_b': utilidad_b})
		context.update({'utilidad_d': utilidad_d})
		return context

class FacturasExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = request.GET.get('campo')
		campo2 = request.GET.get('campo2')
		res = Factura.objects.filter(Q(fecha_fac__month=campo) & Q(fecha_fac__year=campo2)).exclude(asiento__encabezadoas__anulado=True)
		wb = Workbook()
		ws = wb.active
	
		ws['A2'] = 'Especificación'
		ws['B2'] = 'N°'
		ws['C2'] = 'Fecha de la Factura o DUI'
		ws['D2'] = 'NIT Proveedor'
		ws['E2'] = 'Nombre y Apellidos/Razón Social'
		ws['F2'] = 'N° de la Factura'
		ws['G2'] = 'N° de DUI'
		ws['H2'] = 'N° de Autorización'
		ws['I2'] = 'Importe Total de la Compra'
		ws['J2'] = 'Importe no Sujeto a Crédito Fiscal'
		ws['K2'] = 'Subtotal'
		ws['L2'] = 'Descuentos Bonificaciones y Rebajas Obtenidas'
		ws['M2'] = 'Importe Base para Crédito Fiscal'
		ws['N2'] = 'Crédito Fiscal'
		ws['O2'] = 'Código de Control'
		ws['P2'] = 'Tipo de Compra'
		ws['Q2'] = 'Número de Comprobante'
		ws['R2'] = 'Fecha Comprobante'


		cont = 3
		num = 1
	
		for facturas in res:
			ws.cell(row = cont, column = 1).value = 1
			ws.cell(row = cont, column = 2).value = num
			ws.cell(row = cont, column = 3).value = facturas.fecha_f()
			ws.cell(row = cont, column = 4).value = facturas.nit_prov
			ws.cell(row = cont, column = 5).value = facturas.nomb_razon
			ws.cell(row = cont, column = 6).value = facturas.nro_fact
			ws.cell(row = cont, column = 7).value = facturas.nro_dui
			ws.cell(row = cont, column = 8).value = facturas.nro_aut
			ws.cell(row = cont, column = 9).value = facturas.importe_fac
			ws.cell(row = cont, column = 10).value = facturas.importe_credit
			ws.cell(row = cont, column = 11).value = facturas.subtot
			ws.cell(row = cont, column = 12).value = facturas.descuento
			ws.cell(row = cont, column = 13).value = facturas.base_credito
			ws.cell(row = cont, column = 14).value = facturas.credit_fisc
			ws.cell(row = cont, column = 15).value = facturas.codigo()
			ws.cell(row = cont, column = 16).value = facturas.tipo_fac
			ws.cell(row = cont, column = 17).value = facturas.asiento.encabezadoas.numero_comp
			ws.cell(row = cont, column = 18).value = facturas.asiento.encabezadoas.fecha
			num+=1
			cont+=1
	
		nombre_archivo = "Facturas_Compras.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class FacturasVExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = request.GET.get('campo')
		campo2 = request.GET.get('campo2')
		res = Factura_V.objects.filter(Q(fecha_fac__month=campo) & Q(fecha_fac__year=campo2)).exclude(asiento__encabezadoas__anulado=True)
		wb = Workbook()
		ws = wb.active
	
		ws['A2'] = 'N°'
		ws['B2'] = 'Fecha de la Factura'
		ws['C2'] = 'N° de la Factura'
		ws['D2'] = 'N° de Autorización'
		ws['E2'] = 'Estado'
		ws['F2'] = 'NIT/CI Cliente'
		ws['G2'] = 'Nombre y Apellidos/Razón Social'
		ws['H2'] = 'Importe Total de Ventas'
		ws['I2'] = 'Importe ICE/IEHD/Tasas'
		ws['J2'] = 'Exportaciones y Operaciones exentas'
		ws['K2'] = 'Ventas gravadas a tasa cero'
		ws['L2'] = 'Subtotal'
		ws['M2'] = 'Descuentos Bonificaciones y Rebajas Obtenidas'
		ws['N2'] = 'Importe Base para Débito Fiscal'
		ws['O2'] = 'Débito Fiscal'
		ws['P2'] = 'Código de Control'


		cont = 3
		num = 1
	
		for facturas in res:
			ws.cell(row = cont, column = 1).value = num
			ws.cell(row = cont, column = 2).value = facturas.fecha_f()
			ws.cell(row = cont, column = 3).value = facturas.nro_fact
			ws.cell(row = cont, column = 4).value = facturas.nro_aut
			ws.cell(row = cont, column = 5).value = facturas.estado_p
			ws.cell(row = cont, column = 6).value = facturas.nit_cliente
			ws.cell(row = cont, column = 7).value = facturas.nomb_razon
			ws.cell(row = cont, column = 8).value = facturas.importe_fac
			ws.cell(row = cont, column = 9).value = facturas.importe_ice
			ws.cell(row = cont, column = 10).value = facturas.export_oper
			ws.cell(row = cont, column = 11).value = facturas.ventas_grav
			ws.cell(row = cont, column = 12).value = facturas.subtot
			ws.cell(row = cont, column = 13).value = facturas.descuento
			ws.cell(row = cont, column = 14).value = facturas.base_debito
			ws.cell(row = cont, column = 15).value = facturas.debito_fisc
			ws.cell(row = cont, column = 16).value = facturas.codigo()

			num+=1
			cont+=1
	
		nombre_archivo = "Facturas_Ventas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RetencionExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = request.GET.get('campo')
		campo2 = request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo2, "%Y-%m-%d")
		res = encabezadoas.objects.filter(Q(fecha__range=(f1,f2)) & Q(Q(asientos__opciones_p='R') | Q(asientos__opciones_p='A'))).exclude(anulado=True)
		wb = Workbook()
		ws = wb.active
	
		ws['A2'] = 'Nº'
		ws['B2'] = 'Nº COMP'
		ws['C2'] = 'FECHA'
		ws['D2'] = 'TIPO'
		ws['E2'] = 'DETALLE'
		ws['F2'] = 'IMPORTE'
		ws['G2'] = 'IMPORTE TOTAL'
		ws['H2'] = 'SERVICIO'
		ws['I2'] = 'COMPRAS'
		ws['J2'] = 'IT 3%'


		cont = 3
		num = 1
		for i in res:
			aux = 0
			aux1 = 0
			aux2 = 0
			aux3 = 0
			result = 0
			ws.cell(row = cont, column = 1).value = num
			ws.cell(row = cont, column = 2).value = i.numero_comp
			ws.cell(row = cont, column = 3).value = i.fecha_fi()
			for j in i.asientos.all():
				if((j.opciones_p == 'R') | (j.opciones_p == 'A')):
					ws.cell(row = cont, column = 4).value = j.opciones_s
					ws.cell(row = cont, column = 5).value = j.refer
					ws.cell(row = cont, column = 7).value = j.debe_b
					aux = j.debe_b
				elif(j.nivel_cinco.codigo == '2101110005'):
					ws.cell(row = cont, column = 8).value = j.haber_b
					aux1 = j.haber_b
				elif(j.nivel_cinco.codigo == '2101110007'):
					ws.cell(row = cont, column = 9).value = j.haber_b
					aux2 = j.haber_b
				elif(j.nivel_cinco.codigo == '2101110006'):
					ws.cell(row = cont, column = 10).value = j.haber_b
					aux3 = j.haber_b
				result = aux - aux1 - aux2 - aux3
				ws.cell(row = cont, column = 6).value = result
			
			num+=1
			cont+=1
	
		nombre_archivo = "Retenciones.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class estadoEvolucionP(PDFTemplateView):
	template_name = "contabilidad/evolucion_p.html"
	def get_context_data(self, **kwargs):
		context = super(estadoEvolucionP, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		year_act = f.year
		year = f.year - 1
		res = 0
		fecha = f.strftime("%d/%m/%Y")
		context['campo'] = fecha
		context['asi'] = asiento.objects.filter(encabezadoas__fecha__lte=f)
		
		if (f.month!=10) | (f.month!=11) | (f.month!=12):
			c1_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year)+"-09-30") & Q(nivel_cinco__codigo="3101010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c2_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year)+"-09-30") & Q(nivel_cinco__codigo="3101030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c3_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year)+"-09-30") & Q(nivel_cinco__codigo="3102010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c4_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year)+"-09-30") & Q(nivel_cinco__codigo="3102030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c5_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year)+"-09-30") & Q(nivel_cinco__codigo="3104010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
		else:
			c1_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year_act)+"-09-30") & Q(nivel_cinco__codigo="3101010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('debe_b') - F('haber_b')))
			c2_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year_act)+"-09-30") & Q(nivel_cinco__codigo="3101030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('debe_b') - F('haber_b')))
			c3_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year_act)+"-09-30") & Q(nivel_cinco__codigo="3102010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('debe_b') - F('haber_b')))
			c4_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year_act)+"-09-30") & Q(nivel_cinco__codigo="3102030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('debe_b') - F('haber_b')))
			c5_a = asiento.objects.filter(Q(encabezadoas__fecha__lte=str(year_act)+"-09-30") & Q(nivel_cinco__codigo="3104010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('debe_b') - F('haber_b')))
		
		
		if c1_a['total']:
			aux = c1_a['total'] #capital social anterior gestión
		else:
			aux = 0
		if c2_a['total']:
			aux1 = c2_a['total'] #ajuste de capital anterior gestión
		else:
			aux1 = 0
		if c3_a['total']:
			aux2 = c3_a['total'] #reserva legal anterior gestión
		else:
			aux2 = 0
		if c4_a['total']:
			aux3 = c4_a['total'] #ajuste reservas patrimoniales anterior gestión
		else:
			aux3 = 0
		if c5_a['total']:
			aux4 = c5_a['total'] #resultados acumulados anterior gestión
		else:
			aux4 = 0

		patrimonio_neto_a = aux + aux1 + aux2 + aux3 + aux4
		
		if (f.month!=10) | (f.month!=11) | (f.month!=12):
			c1 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year)+"-10-01", f)) & Q(nivel_cinco__codigo="3101010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c2 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year)+"-10-01", f)) & Q(nivel_cinco__codigo="3101030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c3 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year)+"-10-01", f)) & Q(nivel_cinco__codigo="3102010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c4 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year)+"-10-01", f)) & Q(nivel_cinco__codigo="3102030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c5 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year)+"-10-01", f)) & Q(nivel_cinco__codigo="3104010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c6 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year)+"-10-01", f)) & Q(nivel_cinco__codigo="3104020001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
		else:
			c1 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year_act)+"-10-01", f)) & Q(nivel_cinco__codigo="3101010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c2 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year_act)+"-10-01", f)) & Q(nivel_cinco__codigo="3101030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c3 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year_act)+"-10-01", f)) & Q(nivel_cinco__codigo="3102010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c4 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year_act)+"-10-01", f)) & Q(nivel_cinco__codigo="3102030001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c5 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year_act)+"-10-01", f)) & Q(nivel_cinco__codigo="3104010001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			c6 = asiento.objects.filter(Q(encabezadoas__fecha__range=(str(year_act)+"-10-01", f)) & Q(nivel_cinco__codigo="3104020001")).exclude(encabezadoas__anulado=True).aggregate(total=Sum(F('haber_b') - F('debe_b')))
			
		if c1['total']:
			aux5 = c1['total']  #capital social
		else:
			aux5 = 0
		if c2['total']:
			aux6 = c2['total']	#ajuste de capital	
		else:
			aux6 = 0
		if c3['total']:
			aux7 = c3['total']	#reserva legal	
		else:
			aux7 = 0
		if c4['total']:
			aux8 = c4['total']	#ajuste reservas patrimoniales	
		else:
			aux8 = 0
		if c5['total']:
			aux9 = c5['total']	#resultados acumulados	
		else:
			aux9 = 0
		if c6['total']:
			aux10 = c6['total']	#resultado del ejercicio
		else:
			aux10 = 0

		capital_social = aux + aux5
		ajuste_capital = aux1 + aux6
		reserva_legal = aux2 + aux7
		ajuste_reserv = aux3 + aux8
		resultados_acum = aux4 + aux5 + aux7 + aux9 + aux10
		patrimonio_total = patrimonio_neto_a + aux6 + aux8 + aux9 + aux10

		context['patri_n'] = patrimonio_neto_a
		context['cap_soc_a'] = aux
		context['ajuste_cap_a'] = aux1
		context['res_leg_a'] = aux2
		context['ajust_res_a'] = aux3
		context['res_acum_a'] = aux4
		context['cap_soc'] = aux5
		context['ajuste_cap'] = aux6
		context['res_leg'] = aux7
		context['ajust_res'] = aux8
		context['res_acum'] = aux9
		context['resultados_ejer'] = aux10
		context['cap_soc_t'] = capital_social
		context['ajuste_cap_t'] = ajuste_capital
		context['res_leg_t'] = reserva_legal
		context['ajust_res_t'] = ajuste_reserv
		context['res_acum_t'] = resultados_acum
		context['patri_t'] = patrimonio_total

		return context

class libroDiario(PDFTemplateView):
	template_name = "contabilidad/lib_diario.html"
	def get_context_data(self, **kwargs):
		context = super(libroDiario, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		res = asiento.objects.filter(encabezadoas__fecha__range=(f1, f2)).exclude(encabezadoas__anulado=True).order_by('encabezadoas__numero_comp')
		context['lib'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2

		return context

class ReporteCostoServicioTemplate(FormView):
	form_class = repCostoServicioForm
	template_name = 'contabilidad/costo_servicio.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteCostoServicioTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteCostoServicio(PDFTemplateView):
	template_name = "contabilidad/rep_costo_servicio.html"
	def get_context_data(self, **kwargs):
		context = super(ReporteCostoServicio, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")

		res = nivel_cinco.objects.filter(Q(nivel_cuatro__codigo='510101') | Q(nivel_cuatro__codigo='510102') | Q(nivel_cuatro__codigo='510103'), asiento__encabezadoas__fecha__range=(f1, f2), nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__anulado=False).distinct().order_by('codigo')

		res1 = nivel_cinco.objects.filter(Q(nivel_cuatro__codigo='510101') | Q(nivel_cuatro__codigo='510102') | Q(nivel_cuatro__codigo='510103'), asiento__encabezadoas__fecha__range=(f1, f2), nivel_dos__nivel_tres__nivel_cuatro__nivel_cinco__asiento__encabezadoas__anulado=False).aggregate(total=Sum('asiento__debe_b'))

		context['res'] = res
		context['total'] = res1['total']
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2

		return context

class RepPlanCuentas(PDFTemplateView):
	template_name = "contabilidad/rep_plan_cuenta.html"
	def get_context_data(self, **kwargs):
		context = super(RepPlanCuentas, self).get_context_data(**kwargs)
		res = nivel_uno.objects.all().order_by('id')
		context['res'] = res

		return context
