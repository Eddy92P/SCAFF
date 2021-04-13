# -*- coding: utf-8 -*-
from __future__ import unicode_literals

from django.shortcuts import render
from activos_f.forms import *
from activos_f.models import *
from usuarios.models import *
from almacenes.models import *

from django.views.generic import ListView, CreateView, UpdateView, FormView, TemplateView
from django.core.urlresolvers import reverse_lazy
from django.http import HttpResponseRedirect, Http404, HttpResponse

from easy_pdf.views import PDFTemplateView
from django.db.models import Q

import datetime
from itertools import chain
from operator import attrgetter
from django.db.models import F
from django.db.models import Sum
from django.db.models import Count, Avg
from django.db.models import Subquery, OuterRef

from itertools import chain
from operator import attrgetter, methodcaller

from openpyxl import Workbook
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

class rubroAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return Rubro.objects.none()
		qs = Rubro.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs

class tipo_bienAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return Tipo_bien.objects.none()
		qs = Tipo_bien.objects.all().order_by('id')
		rubro = self.forwarded.get('rubro', None)
		if rubro:
			qs = qs.filter(rubro=rubro)
		if self.q:
			qs = qs.filter(tipo__istartswith=self.q)
		return qs

class crearRubro(CreateView):
	form_class = crearRubroForm
	template_name = 'activos_f/nuevorubro.html'
	success_url = reverse_lazy('activos_f:listarubro')

class updateRubro(UpdateView):
	model = Rubro
	form_class = crearRubroForm
	template_name = 'activos_f/nuevorubro.html'
	success_url = reverse_lazy('activos_f:listarubro')

class listaRubro(ListView):
	model = Rubro
	template_name = 'activos_f/listarubro.html'

class crearTipo_bien(CreateView):
	model = Tipo_bien
	form_class = crearTipo_bienForm
	template_name = 'activos_f/nuevotipo_bien.html'
	success_url = 'activos_f:listatipo_bien'

	def get_context_data(self, **kwargs):
		context = super(crearTipo_bien, self).get_context_data(**kwargs)
		return context
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			r = self.model.objects.filter(rubro=form.rubro.id).count()
			aux = form.rubro.prefijo + '-' + str(r+1).zfill(4)
			form.codigo = aux
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateTipo_bien(UpdateView):
	model = Tipo_bien
	form_class = crearTipo_bienForm
	template_name = 'activos_f/nuevotipo_bien.html'
	success_url = reverse_lazy('activos_f:listatipo_bien')

class listaTipo_bien(ListView):
	model = Tipo_bien
	template_name = 'activos_f/listatipo_bien.html'

class crearUbicacion(CreateView):
	form_class = crearUbicacionForm
	template_name = 'activos_f/nuevaubicacion.html'
	success_url = reverse_lazy('activos_f:listaubicacion')

class updateUbicacion(UpdateView):
	model = Ubicacion_gen
	form_class = crearUbicacionForm
	template_name = 'activos_f/nuevaubicacion.html'
	success_url = reverse_lazy('activos_f:listaubicacion')

class listaUbicacion(ListView):
	model = Ubicacion_gen
	template_name = 'activos_f/listaubicacion.html'

class crearAltaActivo(CreateView):
	model = Alta_baja_activo
	form_class = crearAltaActivoForm
	template_name = 'activos_f/nueva_altaactivo.html'
	success_url = 'activos_f:lista_altaactivo'

	def get_context_data(self, **kwargs):
		context = super(crearAltaActivo, self).get_context_data(**kwargs)
		return context

	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		try:
			pk1 = self.kwargs['pk']
			act = salidaitem.objects.get(id=pk1)
			form = self.form_class(initial={'cantidad':act.cantidad, 'tipo_bien':act.item.item_act.tipo_bien.id, 'ubicacion':act.encsalidaitem.cliente, 'valor_inic':act.sub_total})
		except Exception as e:
			pass
		crear_personal_form_set = crearPersonalFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_personal_form_set=crear_personal_form_set))

	def post(self, request, *args, **kwargs):
		self.object = None	
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_personal_form_set = crearPersonalFormSet(request.POST, request.FILES)
		if form.is_valid() and crear_personal_form_set.is_valid():
			form1 = form.save(commit=False) 
			r = Rubro.objects.get(id=form1.tipo_bien.rubro.id)
			a = self.model.objects.all().count()
			cod = str(a+1) + '-' + form1.tipo_bien.codigo
			if form1.vida_util == 0:
				form1.vida_util = r.vida_util * 12
				form1.vida_util_rest = r.vida_util * 12
			form1.codigo = cod
			crear_personal_form_set.instance = form1
			persona = crear_personal_form_set.save(commit=False)
			form1.save()
			crear_personal_form_set.save()
			try:
				pk1 = self.kwargs['pk']
				act = salidaitem.objects.get(id=pk1)
				act.encsalidaitem.estado_act = True
				act.save()
			except Exception as e:
				pass
			return  HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_personal_form_set = crear_personal_form_set))

class updateAltaActivo(UpdateView):
	model = Alta_baja_activo
	form_class = crearAltaActivoForm
	template_name = 'activos_f/nueva_altaactivo.html'
	success_url = 'activos_f:lista_altaactivo'

	def get_context_data(self, **kwargs):
		context = super(updateAltaActivo, self).get_context_data(**kwargs)
		return context

	def get(self, request, *args, **kwargs):
		self.object = self.get_object()
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_personal_form_set = updatePersonalFormSet(instance=self.object)
		return self.render_to_response(self.get_context_data(form=form, crear_personal_form_set=crear_personal_form_set))

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']	
		form = self.form_class(request.POST, request.FILES, instance=self.model.objects.get(id=pk1))
		crear_personal_form_set = updatePersonalFormSet(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid() and crear_personal_form_set.is_valid():
			form1 = form.save(commit=False) 
			crear_personal_form_set.instance = form1
			form1.save()
			crear_personal_form_set.save()
			return  HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_personal_form_set = crear_personal_form_set))

class listaAltaActivo(ListView):
	model = Alta_baja_activo
	template_name = 'activos_f/lista_altaactivo.html'

	def get_context_data(self, **kwargs):
		context = super (listaAltaActivo, self).get_context_data(**kwargs)	
		return context

	def get(self, request, *args, **kwargs):
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list))

	def get_queryset(self):
		object_list = self.model.objects.filter(baja=False).order_by('id')
		return object_list

class updateBajaActivo(UpdateView):
	model = Alta_baja_activo
	form_class = crearBajaActivoForm
	template_name = 'activos_f/nueva_bajaactivo.html'
	success_url = reverse_lazy('activos_f:lista_bajaactivo')

	def get_context_data(self, **kwargs):
		context = super(updateBajaActivo, self).get_context_data(**kwargs)
		pk1=self.kwargs['pk']
		alta = self.model.objects.get(id=pk1)
		if 'form' and 'alta' not in context :
			context['form'] = self.form_class()
			context['alta'] = alta
		return context

class listaBajaActivo(ListView):
	model = Alta_baja_activo
	template_name = 'activos_f/lista_bajaactivo.html'

	def get_context_data(self, **kwargs):
		context = super (listaBajaActivo, self).get_context_data(**kwargs)	
		return context

	def get(self, request, *args, **kwargs):
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list))

	def get_queryset(self):
		object_list = self.model.objects.filter(baja=True).order_by('id')
		return object_list

class crearCapitalizacion(CreateView):
	model = Capitalizacion
	form_class = crearCapitalizacionForm
	template_name = 'activos_f/nueva_capitalizacion.html'
	success_url = 'activos_f:lista_capitalizacion'

	def get_context_data(self, **kwargs):
		context = super(crearCapitalizacion, self).get_context_data(**kwargs)
		return context

	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		try:
			pk1 = self.kwargs['pk']
			capit = salidaitem.objects.get(id=pk1)
			form = self.form_class(initial={'item':capit.item.id, 'valor_inic':capit.sub_total})
		except Exception as e:
			pass
		return self.render_to_response(self.get_context_data(form=form))

	def post(self, request, *args, **kwargs):
		self.object = None	
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		if form.is_valid():
			form1 = form.save(commit=False) 
			form1.save()
			try:
				pk1 = self.kwargs['pk']
				capit = salidaitem.objects.get(id=pk1)
				capit.estado_cap = True
				capit.save()
			except Exception as e:
				pass
			return  HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateCapitalizacion(UpdateView):
	model = Capitalizacion
	form_class = crearCapitalizacionForm
	template_name = 'activos_f/nueva_capitalizacion.html'
	success_url = reverse_lazy('activos_f:lista_capitalizacion')

class listaCapitalizacion(ListView):
	model = Capitalizacion
	template_name = 'activos_f/lista_capitalizacion.html'

#Reportes
class ReporteEspecBienTemplate(FormView):
	form_class = reporteEspecBienForm
	template_name = 'activos_f/temp_espec_bien.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteEspecBienTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteEspecBien(PDFTemplateView):
	template_name = "activos_f/reporte_espec_bien.html"
	def get_context_data(self, **kwargs):
		context = super(ReporteEspecBien, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		fecha = f.date()
		y = f.year
		y_ant = f.year - 1
		f1 = str(y) + "-09-30"
		f2 = str(y_ant) + "-09-30"
		fecha1 = f.strftime("%d/%m/%Y")
		f3 = datetime.datetime.strptime(f1, "%Y-%m-%d")
		f4 = datetime.datetime.strptime(f2, "%Y-%m-%d")
		fecha2 = f4.strftime("%d/%m/%Y")
		fecha_ant1 = f3.date()
		fecha_ant2 = f4.date()

		aux = []
		aux1 = []

		valor_inic = 0
		valor_actual = 0
		actualizacion = 0
		deprec_gestion = 0
		deprec_acum_inicial = 0
		valor_act_deprec_acum = 0
		actualiz_deprec_acum = 0
		deprec_acum = 0
		saldo_activo = 0

		valor_inic_cap = 0
		valor_actual_cap = 0
		actualizacion_cap = 0
		deprec_gestion_cap = 0
		deprec_acum_inicial_cap = 0
		valor_act_deprec_acum_cap = 0
		actualiz_deprec_acum_cap = 0
		deprec_acum_cap = 0
		saldo_activo_cap = 0

		if (f.month != 10) & (f.month != 11) & (f.month != 12):
			context['f2'] = fecha_ant2
			f5 = fecha_ant2
		else:
			context['f2'] = fecha_ant1
			f5 = fecha_ant1

		if (campo1 == 'T') & (campo2 == '1'):
			res = Alta_baja_activo.objects.filter(fecha_alta__lte=f, vinculado=True, baja=False).order_by('tipo_bien__rubro', 'fecha_alta')
		elif (campo1 == 'T') & (campo2 == '2'):
			res = Alta_baja_activo.objects.filter(fecha_alta__lte=f, vinculado=False, baja=False).order_by('tipo_bien__rubro', 'fecha_alta')
		elif (campo1 != 'T') & (campo2 == '1'):
			res = Alta_baja_activo.objects.filter(fecha_alta__lte=f, tipo_bien__rubro=campo1, vinculado=True, baja=False).order_by('fecha_alta')
		elif (campo1 != 'T') & (campo2 == '2'):
			res = Alta_baja_activo.objects.filter(fecha_alta__lte=f, tipo_bien__rubro=campo1, vinculado=False, baja=False).order_by('fecha_alta')

		context['res'] = res
		context['fecha'] = fecha1
		context['fecha2'] = fecha2
		context['f'] = fecha
		context['campo1'] = campo1
		context['campo2'] = campo2


		for i in res:
			valor_inic = valor_inic + i.valor_inicial(f5)
			valor_actual = valor_actual + i.valor_actual(fecha, f5)
			actualizacion = actualizacion + i.actualizacion(fecha, f5)
			deprec_gestion = deprec_gestion + i.deprec_gestion(fecha, f5)
			deprec_acum_inicial = deprec_acum_inicial + i.deprec_acum_inicial(fecha, f5)
			valor_act_deprec_acum = valor_act_deprec_acum + i.valor_act_deprec_acum(fecha, f5)
			actualiz_deprec_acum = actualiz_deprec_acum + i.actualiz_deprec_acum(fecha, f5)
			deprec_acum = deprec_acum + i.deprec_acum(fecha, f5)
			saldo_activo = saldo_activo + i.saldo_activo(fecha, f5)

			if i.capitalizacion_set.filter(fecha_cap__lte=f):
				for j in i.capitalizacion_set.filter(fecha_cap__lte=f):
					valor_inic_cap = valor_inic_cap + j.valor_inicial_cap(f5)
					valor_actual_cap = valor_actual_cap + j.valor_actual_cap(fecha, f5)
					actualizacion_cap = actualizacion_cap + j.actualizacion_cap(fecha, f5)
					deprec_gestion_cap = deprec_gestion_cap + j.deprec_gestion_cap(fecha, f5)
					deprec_acum_inicial_cap = deprec_acum_inicial_cap + j.deprec_acum_inicial_cap(fecha, f5)
					valor_act_deprec_acum_cap = valor_act_deprec_acum_cap + j.valor_act_deprec_acum_cap(fecha, f5)
					actualiz_deprec_acum_cap = actualiz_deprec_acum_cap + j.actualiz_deprec_acum_cap(fecha, f5)
					deprec_acum_cap = deprec_acum_cap + j.deprec_acum_cap(fecha, f5)
					saldo_activo_cap = saldo_activo_cap + j.saldo_activo_cap(fecha, f5)

		total_valor_inic = valor_inic + valor_inic_cap
		total_valor_actual = valor_actual + valor_actual_cap
		total_actualizacion = actualizacion + actualizacion_cap
		total_deprec_gestion = deprec_gestion + deprec_gestion_cap
		total_deprec_acum_inicial = deprec_acum_inicial + deprec_acum_inicial_cap
		total_valor_act_deprec_acum = valor_act_deprec_acum + valor_act_deprec_acum_cap
		total_actualiz_deprec_acum = actualiz_deprec_acum + actualiz_deprec_acum_cap
		total_deprec_acum = deprec_acum + deprec_acum_cap
		total_saldo_activo = saldo_activo + saldo_activo_cap

		context['total_valor_inic'] = total_valor_inic
		context['total_valor_actual'] = total_valor_actual
		context['total_actualizacion'] = total_actualizacion
		context['total_deprec_gestion'] = total_deprec_gestion
		context['total_deprec_acum_inicial'] = total_deprec_acum_inicial
		context['total_valor_act_deprec_acum'] = total_valor_act_deprec_acum
		context['total_actualiz_deprec_acum'] = total_actualiz_deprec_acum
		context['total_deprec_acum'] = total_deprec_acum
		context['total_saldo_activo'] = total_saldo_activo

		return context

class ReporteAltaBajaTemplate(FormView):
	form_class = reporteAltaBajaForm
	template_name = 'activos_f/temp_alta_baja.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteAltaBajaTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteAltaBaja(PDFTemplateView):
	template_name = "activos_f/reporte_alta_baja.html"
	def get_context_data(self, **kwargs):
		context = super(ReporteAltaBaja, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		campo3 = self.request.GET.get('campo3')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f1 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha = f.strftime("%d/%m/%Y")
		fecha1 = f1.strftime("%d/%m/%Y")

		if (campo2 == 'T') & (campo3 == '1'):
			res = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=False).order_by('tipo_bien__rubro', 'fecha_alta')
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=False).aggregate(total=Sum('valor_inic'))
		elif (campo2 == 'T') & (campo3 == '2'):
			res = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=True).order_by('tipo_bien__rubro', 'fecha_baja')
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=True).aggregate(total=Sum('valor_inic'))
		elif (campo2 != 'T') & (campo3 == '1'):
			res = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=False, tipo_bien__rubro=campo2).order_by('tipo_bien__rubro', 'fecha_alta')
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=False, tipo_bien__rubro=campo2).aggregate(total=Sum('valor_inic'))
		else:
			res = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=True, tipo_bien__rubro=campo2).order_by('tipo_bien__rubro', 'fecha_baja')
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=True, tipo_bien__rubro=campo2).aggregate(total=Sum('valor_inic'))

		context['fecha'] = fecha
		context['fecha1'] = fecha1
		context['f'] = f
		context['f1'] = f1
		context['campo2'] = campo2
		context['campo3'] = campo3
		context['total_val_inic'] = res1['total']
		context['res'] = res

		return context

class ReporteInventarioTemplate(FormView):
	form_class = reporteInventarioForm
	template_name = 'activos_f/temp_inventario.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteInventarioTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteInventario(PDFTemplateView):
	template_name = "activos_f/reporte_inventario.html"
	def get_context_data(self, **kwargs):
		context = super(ReporteInventario, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f1 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha = f.strftime("%d/%m/%Y")
		fecha1 = f1.strftime("%d/%m/%Y")

		res = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1)).order_by('fecha_alta')

		context['fecha'] = fecha
		context['fecha1'] = fecha1
		context['res'] = res

		return context

class ReporteEspecBienExcelTemplate(FormView):
	form_class = reporteEspecBienForm
	template_name = 'activos_f/temp_espec_bien_excel.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteEspecBienExcelTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteEspecBienExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		fecha = f.date()
		y = f.year
		y_ant = f.year - 1
		f1 = str(y) + "-09-30"
		f2 = str(y_ant) + "-09-30"
		fecha1 = f.strftime("%d/%m/%Y")
		f3 = datetime.datetime.strptime(f1, "%Y-%m-%d")
		f4 = datetime.datetime.strptime(f2, "%Y-%m-%d")
		fecha2 = f4.strftime("%d/%m/%Y")
		fecha_ant1 = f3.date()
		fecha_ant2 = f4.date()

		aux = []
		aux1 = []

		valor_inic = 0
		valor_actual = 0
		actualizacion = 0
		deprec_gestion = 0
		deprec_acum_inicial = 0
		valor_act_deprec_acum = 0
		actualiz_deprec_acum = 0
		deprec_acum = 0
		saldo_activo = 0

		valor_inic_cap = 0
		valor_actual_cap = 0
		actualizacion_cap = 0
		deprec_gestion_cap = 0
		deprec_acum_inicial_cap = 0
		valor_act_deprec_acum_cap = 0
		actualiz_deprec_acum_cap = 0
		deprec_acum_cap = 0
		saldo_activo_cap = 0

		wb = Workbook()
		ws = wb.active
		
		if campo2 == '1':
			ws['A1'] = 'REPORTE ESPECIFICO DE BIENES PRACTICADO AL ' + str(fecha2) + ' (VINCULADO AL PROCESO DE PRODUCCION)'
		else:
			ws['A1'] = 'REPORTE ESPECIFICO DE BIENES PRACTICADO AL ' + str(fecha2) + ' (NO VINCULADO AL PROCESO DE PRODUCCION)'
		ws['A2'] = 'Código'
		ws['B2'] = 'Descripción'
		ws['C2'] = 'Fecha compra/revaluo'
		ws['D2'] = 'Valor Inicial al ' + fecha2
		ws['E2'] = 'Factor'
		ws['F2'] = 'Valor actual ' + fecha1
		ws['G2'] = 'Actualización'
		ws['H2'] = 'Depreciación %'
		ws['I2'] = 'Depreciación de la gestión'
		ws['J2'] = 'Depreciación acumulada'
		ws['N2'] = 'Saldo al ' + fecha1
		ws['J3'] = 'Saldo al ' + fecha2
		ws['K3'] = 'Valor actual'
		ws['L3'] = 'Actualización'
		ws['M3'] = 'Saldo al ' + fecha1

		ws.column_dimensions['A'].width = 15
		ws.column_dimensions['B'].width = 25
		ws.column_dimensions['C'].width = 25
		ws.column_dimensions['D'].width = 25
		ws.column_dimensions['E'].width = 15
		ws.column_dimensions['F'].width = 25
		ws.column_dimensions['G'].width = 15
		ws.column_dimensions['H'].width = 15
		ws.column_dimensions['I'].width = 25
		ws.column_dimensions['J'].width = 25
		ws.column_dimensions['K'].width = 15
		ws.column_dimensions['L'].width = 20
		ws.column_dimensions['M'].width = 25
		ws.column_dimensions['N'].width = 20

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['K2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['L2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['M2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['N2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['K3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['L3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['M3'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)
		ws['J2'].font = Font(bold=True)
		ws['K2'].font = Font(bold=True)
		ws['L2'].font = Font(bold=True)
		ws['M2'].font = Font(bold=True)
		ws['N2'].font = Font(bold=True)
		ws['J3'].font = Font(bold=True)
		ws['K3'].font = Font(bold=True)
		ws['L3'].font = Font(bold=True)
		ws['M3'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:N1')
		ws.merge_cells('J2:M2')
		ws.merge_cells('A2:A3')
		ws.merge_cells('B2:B3')
		ws.merge_cells('C2:C3')
		ws.merge_cells('D2:D3')
		ws.merge_cells('E2:E3')
		ws.merge_cells('F2:F3')
		ws.merge_cells('G2:G3')
		ws.merge_cells('H2:H3')
		ws.merge_cells('I2:I3')
		
		cont = 4
		iteration = 1

		if (f.month != 10) & (f.month != 11) & (f.month != 12):
			f5 = fecha_ant2
		else:
			f5 = fecha_ant1

		if (campo1 == 'T') & (campo2 == '1'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, tipo_bien__alta_baja_activo__vinculado=True, tipo_bien__alta_baja_activo__baja=False).order_by('id').distinct()
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, tipo_bien__alta_baja_activo__vinculado=True, tipo_bien__alta_baja_activo__baja=False).distinct().count()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__lte=f, vinculado=True, baja=False).order_by('tipo_bien__rubro', 'fecha_alta')
		elif (campo1 == 'T') & (campo2 == '2'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, tipo_bien__alta_baja_activo__vinculado=False, tipo_bien__alta_baja_activo__baja=False).order_by('id').distinct()
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, tipo_bien__alta_baja_activo__vinculado=False, tipo_bien__alta_baja_activo__baja=False).distinct().count
			res1 = Alta_baja_activo.objects.filter(fecha_alta__lte=f, vinculado=False, baja=False).order_by('tipo_bien__rubro', 'fecha_alta')
		elif (campo1 != 'T') & (campo2 == '1'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, id=campo1, tipo_bien__alta_baja_activo__vinculado=True, tipo_bien__alta_baja_activo__baja=False).order_by('id').distinct()
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, id=campo1, tipo_bien__alta_baja_activo__vinculado=True, tipo_bien__alta_baja_activo__baja=False).distinct().count()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__lte=f, tipo_bien__rubro=campo1, vinculado=True, baja=False).order_by('fecha_alta')
		elif (campo1 != 'T') & (campo2 == '2'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, id=campo1, tipo_bien__alta_baja_activo__vinculado=False, tipo_bien__alta_baja_activo__baja=False).order_by('id').distinct()
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__lte=f, id=campo1, tipo_bien__alta_baja_activo__vinculado=False, tipo_bien__alta_baja_activo__baja=False).distinct().count()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__lte=f, tipo_bien__rubro=campo1, vinculado=False, baja=False).order_by('fecha_alta')

		for i in res1:
			valor_inic = valor_inic + i.valor_inicial(f5)
			valor_actual = valor_actual + i.valor_actual(fecha, f5)
			actualizacion = actualizacion + i.actualizacion(fecha, f5)
			deprec_gestion = deprec_gestion + i.deprec_gestion(fecha, f5)
			deprec_acum_inicial = deprec_acum_inicial + i.deprec_acum_inicial(fecha, f5)
			valor_act_deprec_acum = valor_act_deprec_acum + i.valor_act_deprec_acum(fecha, f5)
			actualiz_deprec_acum = actualiz_deprec_acum + i.actualiz_deprec_acum(fecha, f5)
			deprec_acum = deprec_acum + i.deprec_acum(fecha, f5)
			saldo_activo = saldo_activo + i.saldo_activo(fecha, f5)

			if i.capitalizacion_set.filter(fecha_cap__lte=f):
				for j in i.capitalizacion_set.filter(fecha_cap__lte=f):
					valor_inic_cap = valor_inic_cap + j.valor_inicial_cap(f5)
					valor_actual_cap = valor_actual_cap + j.valor_actual_cap(fecha, f5)
					actualizacion_cap = actualizacion_cap + j.actualizacion_cap(fecha, f5)
					deprec_gestion_cap = deprec_gestion_cap + j.deprec_gestion_cap(fecha, f5)
					deprec_acum_inicial_cap = deprec_acum_inicial_cap + j.deprec_acum_inicial_cap(fecha, f5)
					valor_act_deprec_acum_cap = valor_act_deprec_acum_cap + j.valor_act_deprec_acum_cap(fecha, f5)
					actualiz_deprec_acum_cap = actualiz_deprec_acum_cap + j.actualiz_deprec_acum_cap(fecha, f5)
					deprec_acum_cap = deprec_acum_cap + j.deprec_acum_cap(fecha, f5)
					saldo_activo_cap = saldo_activo_cap + j.saldo_activo_cap(fecha, f5)

		total_valor_inic = valor_inic + valor_inic_cap
		total_valor_actual = valor_actual + valor_actual_cap
		total_actualizacion = actualizacion + actualizacion_cap
		total_deprec_gestion = deprec_gestion + deprec_gestion_cap
		total_deprec_acum_inicial = deprec_acum_inicial + deprec_acum_inicial_cap
		total_valor_act_deprec_acum = valor_act_deprec_acum + valor_act_deprec_acum_cap
		total_actualiz_deprec_acum = actualiz_deprec_acum + actualiz_deprec_acum_cap
		total_deprec_acum = deprec_acum + deprec_acum_cap
		total_saldo_activo = saldo_activo + saldo_activo_cap

		for rubros in res:
			if iteration < num_queries:
				ws.cell(row = cont, column = 1).value = rubros.nombre
				cont += 1
				if (campo1 == 'T') & (campo2 == '1'):
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f, alta_baja_activo__vinculado=True, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=True, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				elif (campo1 == 'T') & (campo2 == '2'):
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f, alta_baja_activo__vinculado=False, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=False, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				elif (campo1 != 'T') & (campo2 == '1'):
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f,alta_baja_activo__vinculado=True, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=True, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				else:
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f,alta_baja_activo__vinculado=False, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=False, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + rubros.nombre
				ws.cell(row = cont, column = 4).value = rubros.val_inic_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 6).value = rubros.valor_actual_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 7).value = rubros.actualizacion_tot_excel(fecha, f5, campo2)
				
				ws.cell(row = cont, column = 9).value = rubros.deprec_gestion_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 10).value = rubros.deprec_acum_inicial_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 11).value = rubros.valor_act_deprec_acum_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 12).value = rubros.actualiz_deprec_acum_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 13).value = rubros.deprec_acum_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 14).value = rubros.saldo_activo_tot_excel(fecha, f5, campo2)
				cont += 1
				iteration += 1
			else:
				ws.cell(row = cont, column = 1).value = rubros.nombre
				cont += 1
				if (campo1 == 'T') & (campo2 == '1'):
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f, alta_baja_activo__vinculado=True, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=True, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				elif (campo1 == 'T') & (campo2 == '2'):
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f, alta_baja_activo__vinculado=False, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=False, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				elif (campo1 != 'T') & (campo2 == '1'):
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f,alta_baja_activo__vinculado=True, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=True, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				else:
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__lte=f, alta_baja_activo__vinculado=False, alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__lte=f, vinculado=False, baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							ws.cell(row = cont, column = 3).value = activo.fecha_alt()
							ws.cell(row = cont, column = 4).value = activo.valor_inicial_r(f5)
							ws.cell(row = cont, column = 5).value = activo.factor(fecha, f5)
							ws.cell(row = cont, column = 6).value = activo.valor_actual(fecha, f5)
							ws.cell(row = cont, column = 7).value = activo.actualizacion(fecha, f5)
							ws.cell(row = cont, column = 8).value = activo.depreciacion()
							ws.cell(row = cont, column = 9).value = activo.deprec_gestion(fecha, f5)
							ws.cell(row = cont, column = 10).value = activo.deprec_acum_inicial_r(fecha, f5)
							ws.cell(row = cont, column = 11).value = activo.valor_act_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 12).value = activo.actualiz_deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 13).value = activo.deprec_acum(fecha, f5)
							ws.cell(row = cont, column = 14).value = activo.saldo_activo(fecha, f5)
							cont += 1
							if activo.capitalizacion_set.exists():
								for capitalizacion in activo.cap(fecha):
									ws.cell(row = cont, column = 2).value = capitalizacion.tipo_bien.tipo
									ws.cell(row = cont, column = 3).value = capitalizacion.fecha_f()
									ws.cell(row = cont, column = 4).value = capitalizacion.valor_inicial_cap_r(f5)
									ws.cell(row = cont, column = 5).value = capitalizacion.factor_cap(fecha, f5)
									ws.cell(row = cont, column = 6).value = capitalizacion.valor_actual_cap(fecha, f5)
									ws.cell(row = cont, column = 7).value = capitalizacion.actualizacion_cap(fecha, f5)
									ws.cell(row = cont, column = 8).value = capitalizacion.depreciacion_cap_r
									ws.cell(row = cont, column = 9).value = capitalizacion.deprec_gestion_cap(fecha, f5)
									ws.cell(row = cont, column = 10).value = capitalizacion.deprec_acum_inicial_cap_r(fecha, f5)
									ws.cell(row = cont, column = 11).value = capitalizacion.valor_act_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 12).value = capitalizacion.actualiz_deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 13).value = capitalizacion.deprec_acum_cap(fecha, f5)
									ws.cell(row = cont, column = 14).value = capitalizacion.saldo_activo_cap(fecha, f5)
									cont += 1
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + rubros.nombre
				ws.cell(row = cont, column = 4).value = rubros.val_inic_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 6).value = rubros.valor_actual_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 7).value = rubros.actualizacion_tot_excel(fecha, f5, campo2)
				
				ws.cell(row = cont, column = 9).value = rubros.deprec_gestion_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 10).value = rubros.deprec_acum_inicial_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 11).value = rubros.valor_act_deprec_acum_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 12).value = rubros.actualiz_deprec_acum_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 13).value = rubros.deprec_acum_tot_excel(fecha, f5, campo2)
				ws.cell(row = cont, column = 14).value = rubros.saldo_activo_tot_excel(fecha, f5, campo2)
				cont += 1
				if campo1 == 'T':
					ws.cell(row = cont, column = 1).value = "TOTALES"
					ws.cell(row = cont, column = 4).value = total_valor_inic

					ws.cell(row = cont, column = 6).value = total_valor_actual
					ws.cell(row = cont, column = 7).value = total_actualizacion

					ws.cell(row = cont, column = 9).value = total_deprec_gestion
					ws.cell(row = cont, column = 10).value = total_valor_act_deprec_acum
					ws.cell(row = cont, column = 11).value = total_deprec_acum_inicial
					ws.cell(row = cont, column = 12).value = total_actualiz_deprec_acum
					ws.cell(row = cont, column = 13).value = total_deprec_acum
					ws.cell(row = cont, column = 14).value = total_saldo_activo

		nombre_archivo = "Reporte Especifico de Bienes.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class ReporteAltaBajaTemplateExcel(FormView):
	form_class = reporteAltaBajaForm
	template_name = 'activos_f/temp_alta_baja_excel.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteAltaBajaTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteAltaBajaExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		campo3 = self.request.GET.get('campo3')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f1 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha = f.strftime("%d/%m/%Y")
		fecha1 = f1.strftime("%d/%m/%Y")

		if (campo2 == 'T') & (campo3 == '1'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=False).order_by('id').distinct()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=False).aggregate(total=Sum('valor_inic'))
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=False).distinct().count()
			

		elif (campo2 == 'T') & (campo3 == '2'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=True).order_by('id').distinct()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=True).aggregate(total=Sum('valor_inic'))
			num_queries = Rubro.objects.filter(tipo_bien_alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=True).distinct().count()

		elif (campo2 != 'T') & (campo3 == '1'):
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=False, id=campo2).order_by('id').distinct()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=False, tipo_bien__rubro=campo2).aggregate(total=Sum('valor_inic'))
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=False, id=campo2).distinct().count()

		else:
			res = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=True, id=campo2).order_by('id').distinct()
			res1 = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1), baja=True, tipo_bien__rubro=campo2).aggregate(total=Sum('valor_inic'))
			num_queries = Rubro.objects.filter(tipo_bien__alta_baja_activo__fecha_alta__range=(f, f1), tipo_bien__alta_baja_activo__baja=True, id=campo2).distinct().count()

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'REPORTE DE ALTAS DESDE EL ' + str(fecha) + ' al ' + str(fecha1)
		ws['A2'] = 'Código'
		ws['B2'] = 'Descripción'
		ws['C2'] = 'Imagen del Activo'
		ws['D2'] = 'Fecha compra o revaluo'
		ws['E2'] = 'Vinculado'
		ws['F2'] = 'Valor Bs.'
		ws['G2'] = 'Depreciación %'
		ws['H2'] = 'Vida útil(mes)'
		ws['I2'] = 'Vida útil restante(mes)'

		ws.column_dimensions['A'].width = 13
		ws.column_dimensions['B'].width = 25
		ws.column_dimensions['C'].width = 40
		ws.column_dimensions['D'].width = 25
		ws.column_dimensions['E'].width = 10
		ws.column_dimensions['F'].width = 15
		ws.column_dimensions['G'].width = 15
		ws.column_dimensions['H'].width = 15
		ws.column_dimensions['I'].width = 25

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:I1')

		cont = 3
		iteration = 1

		for rubros in res:
			if iteration < num_queries:
				ws.cell(row = cont, column = 1).value = rubros.nombre.upper()
				cont += 1
				if campo3 == '1':
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__range=(f, f1), alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__range=(f, f1), baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							if activo.img_activo:
								ws.cell(row = cont, column = 3).value = activo.img_activo.url
							else:
								ws.cell(row = cont, column = 3).value = "Sin Imagen"
							ws.cell(row = cont, column = 4).value = activo.fecha_alt()
							if activo.vinculado == True:
								ws.cell(row = cont, column = 5).value = "S"
							else:
								ws.cell(row = cont, column = 5).value = "N"
							ws.cell(row = cont, column = 6).value = activo.valor_inic
							ws.cell(row = cont, column = 7).value = activo.depreciacion()
							ws.cell(row = cont, column = 8).value = activo.vida_util
							ws.cell(row = cont, column = 9).value = activo.vida_util_rest
							cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + rubros.nombre
					ws.cell(row = cont, column = 6).value = rubros.val_ini_excel(f, f1)
					cont += 1
					iteration += 1

				else:
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__range=(f, f1), alta_baja_activo__baja=True):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__range=(f, f1), baja=True):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							if activo.img_activo:
								ws.cell(row = cont, column = 3).value = activo.img_activo.url
							else:
								ws.cell(row = cont, column = 3).value = "Sin Imagen"
							ws.cell(row = cont, column = 4).value = activo.fecha_alt()
							if activo.vinculado == True:
								ws.cell(row = cont, column = 5).value = "S"
							else:
								ws.cell(row = cont, column = 5).value = "N"
							ws.cell(row = cont, column = 6).value = activo.valor_inic
							ws.cell(row = cont, column = 7).value = activo.depreciacion()
							ws.cell(row = cont, column = 8).value = activo.vida_util
							ws.cell(row = cont, column = 9).value = activo.vida_util_rest
							cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + rubros.nombre
					ws.cell(row = cont, column = 6).value = rubros.val_ini_excel(f, f1)
					cont += 1
					iteration += 1
			else:
				ws.cell(row = cont, column = 1).value = rubros.nombre.upper()
				cont += 1
				if campo3 == '1':
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__range=(f, f1), alta_baja_activo__baja=False):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__range=(f, f1), baja=False):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							if activo.img_activo:
								ws.cell(row = cont, column = 3).value = activo.img_activo.url
							else:
								ws.cell(row = cont, column = 3).value = "Sin Imagen"
							ws.cell(row = cont, column = 4).value = activo.fecha_alt()
							if activo.vinculado == True:
								ws.cell(row = cont, column = 5).value = "S"
							else:
								ws.cell(row = cont, column = 5).value = "N"
							ws.cell(row = cont, column = 6).value = activo.valor_inic
							ws.cell(row = cont, column = 7).value = activo.depreciacion()
							ws.cell(row = cont, column = 8).value = activo.vida_util
							ws.cell(row = cont, column = 9).value = activo.vida_util_rest
							cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + rubros.nombre
					ws.cell(row = cont, column = 6).value = rubros.val_ini_excel(f, f1)
					cont += 1

				else:
					for t_bien in rubros.tipo_bien_set.filter(alta_baja_activo__fecha_alta__range=(f, f1), alta_baja_activo__baja=True):
						for activo in t_bien.alta_baja_activo_set.filter(fecha_alta__range=(f, f1), baja=True):
							ws.cell(row = cont, column = 1).value = activo.codigo
							ws.cell(row = cont, column = 2).value = activo.tipo_bien.tipo
							if activo.img_activo:
								ws.cell(row = cont, column = 3).value = activo.img_activo.url
							else:
								ws.cell(row = cont, column = 3).value = "Sin Imagen"
							ws.cell(row = cont, column = 4).value = activo.fecha_alt()
							if activo.vinculado == True:
								ws.cell(row = cont, column = 5).value = "S"
							else:
								ws.cell(row = cont, column = 5).value = "N"
							ws.cell(row = cont, column = 6).value = activo.valor_inic
							ws.cell(row = cont, column = 7).value = activo.depreciacion()
							ws.cell(row = cont, column = 8).value = activo.vida_util
							ws.cell(row = cont, column = 9).value = activo.vida_util_rest
							cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + rubros.nombre
					ws.cell(row = cont, column = 6).value = rubros.val_ini_excel(f, f1)
					cont += 1
					if campo2 == 'T':
						ws.cell(row = cont, column = 1).value = "TOTALES"
						ws.cell(row = cont, column = 6).value = res1['total']


		nombre_archivo = "Reporte de Altas y Bajas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class ReporteInventarioTemplateExcel(FormView):
	form_class = reporteInventarioForm
	template_name = 'activos_f/temp_inventario_excel.html'
	def get_context_data(self, **kwargs):
		context = super(ReporteInventarioTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class ReporteInventarioExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f1 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha = f.strftime("%d/%m/%Y")
		fecha1 = f1.strftime("%d/%m/%Y")

		res = Alta_baja_activo.objects.filter(fecha_alta__range=(f, f1)).order_by('fecha_alta')

		wb = Workbook()
		ws = wb.active
		
		ws['A2'] = 'INVENTARIO DE ACTIVOS FIJOS DESDE EL ' + str(fecha) + ' ' + str(fecha1)
		ws['A2'] = 'Código'
		ws['B2'] = 'Rubro'
		ws['C2'] = 'Tipo de bien'
		ws['D2'] = 'Marca'
		ws['E2'] = 'Modelo'
		ws['F2'] = 'Características'
		ws['G2'] = 'Color'
		ws['H2'] = 'Serie'
		ws['I2'] = 'Estado'
		ws['J2'] = 'Responsable'
		ws['K2'] = 'Cargo'

		ws.column_dimensions['A'].width = 13
		ws.column_dimensions['B'].width = 25
		ws.column_dimensions['C'].width = 25
		ws.column_dimensions['D'].width = 25
		ws.column_dimensions['E'].width = 25
		ws.column_dimensions['F'].width = 25
		ws.column_dimensions['G'].width = 20
		ws.column_dimensions['H'].width = 20
		ws.column_dimensions['I'].width = 15
		ws.column_dimensions['J'].width = 35
		ws.column_dimensions['K'].width = 25

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['K2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)
		ws['J2'].font = Font(bold=True)
		ws['K2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:K1')
		
		cont = 3
		iteration = 1

		for activos in res:
			ws.cell(row = cont, column = 1).value = activos.codigo
			ws.cell(row = cont, column = 2).value = activos.tipo_bien.rubro.nombre
			ws.cell(row = cont, column = 3).value = activos.tipo_bien.tipo
			ws.cell(row = cont, column = 4).value = activos.tipo_bien.marca
			ws.cell(row = cont, column = 5).value = activos.tipo_bien.modelo
			ws.cell(row = cont, column = 6).value = activos.tipo_bien.caracter
			ws.cell(row = cont, column = 7).value = activos.tipo_bien.color
			ws.cell(row = cont, column = 8).value = activos.tipo_bien.serie
			ws.cell(row = cont, column = 9).value = activos.estado_bien
			for personals in activos.personal_aut_act_set.all():
				ws.cell(row = cont, column = 10).value = personals.personal.nombre + ' ' + personals.personal.ap_paterno
				ws.cell(row = cont, column = 11).value = personals.personal.cargo
				break

		nombre_archivo = "Reporte de Inventario.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response
