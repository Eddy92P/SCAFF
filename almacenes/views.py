# -*- coding: utf-8 -*-
from django.shortcuts import render
from almacenes.forms import *
from almacenes.models import *
from activos_f.models import *
from usuarios.models import *
from contabilidad.views import *
# Create your views here.

from django.views.generic import ListView, CreateView, UpdateView, FormView, TemplateView
from django.core.urlresolvers import reverse_lazy
from django.http import HttpResponseRedirect, Http404, HttpResponse

from easy_pdf.views import PDFTemplateView
from django.db.models import Q

import datetime
from itertools import chain
from operator import attrgetter
from django.db.models import F
from django.db.models import Sum
from django.db.models import Count, Avg
from django.db.models import Subquery, OuterRef

from itertools import chain
from operator import attrgetter, methodcaller

from openpyxl import Workbook
from decimal import Decimal
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
"""
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, cm
from django.http import HttpResponse

from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
from reportlab.platypus.paragraph import Paragraph
from reportlab.platypus import Table,TableStyle
from reportlab.lib import colors
"""

class crearProveedor(CreateView):
	form_class = crearProveedorForm
	template_name = 'almacen/nuevoproveedor.html'
	success_url = reverse_lazy('almacenes:listaproveedor')

class updateProveedor(UpdateView):
	model = proveedor
	form_class = crearProveedorForm
	template_name = 'almacen/nuevoproveedor.html'
	success_url = reverse_lazy('almacenes:listaproveedor')

class listaProveedor(ListView):
	model = proveedor
	template_name = 'almacen/listaproveedor.html'

class crearCliente(CreateView):
	form_class = crearClienteForm
	template_name = 'almacen/nuevocliente.html'
	success_url = reverse_lazy('almacenes:listacliente')

class updateCliente(UpdateView):
	model = cliente
	form_class = crearClienteForm
	template_name = 'almacen/nuevocliente.html'
	success_url = reverse_lazy('almacenes:listacliente')

class listaCliente(CreateView,ListView):
	model = cliente
	form_class = searchForm
	template_name = 'almacen/listacliente.html'
	def get_context_data(self, **kwargs):
		context = super (listaCliente, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

class crearTipoEntrada(CreateView):
	form_class = crearTipoEntradaForm
	template_name = 'almacen/nuevotipo_entrada.html'
	success_url = reverse_lazy('almacenes:listatipo_entrada')

class updateTipoEntrada(UpdateView):
	model = tipo_entrada
	form_class = crearTipoEntradaForm
	template_name = 'almacen/nuevotipo_entrada.html'
	success_url = reverse_lazy('almacenes:listatipo_entrada')

class listaTipoEntrada(CreateView,ListView):
	model = tipo_entrada
	form_class = searchForm
	template_name = 'almacen/listatipo_entrada.html'
	def get_context_data(self, **kwargs):
		context = super (listaTipoEntrada, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

class crearTipoSalida(CreateView):
	form_class = crearTipoSalidaForm
	template_name = 'almacen/nuevotipo_salida.html'
	success_url = reverse_lazy('almacenes:listatipo_salida')

class updateTipoSalida(UpdateView):
	model = tipo_salida
	form_class = crearTipoSalidaForm
	template_name = 'almacen/nuevotipo_salida.html'
	success_url = reverse_lazy('almacenes:listatipo_salida')

class listaTipoSalida(CreateView,ListView):
	model = tipo_salida
	form_class = searchForm
	template_name = 'almacen/listatipo_salida.html'
	def get_context_data(self, **kwargs):
		context = super (listaTipoSalida, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

class crearUnidad(CreateView):
	form_class = crearUnidadForm
	template_name = 'almacen/nuevaunidad.html'
	success_url = reverse_lazy('almacenes:listaunidad')

class updateUnidad(UpdateView):
	model = unidad
	form_class = crearUnidadForm
	template_name = 'almacen/nuevaunidad.html'
	success_url = reverse_lazy('almacenes:listaunidad')

class listaUnidad(CreateView,ListView):
	model = unidad
	form_class = searchForm
	template_name = 'almacen/listaunidad.html'
	def get_context_data(self, **kwargs):
		context = super (listaUnidad, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

#Autocompletado encadenado
"""class almacenAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return almacen.objects.none()
		qs = almacen.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs"""

class categoriaAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return categoria.objects.none()
		qs = categoria.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs

class subcategoriaAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return subcategoria.objects.none()
		qs = subcategoria.objects.all().order_by('id')
		categoria = self.forwarded.get('categoria', None)
		if categoria:
			qs = qs.filter(categoria=categoria)
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs

class itemAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return item.objects.none()
		qs = item.objects.all().order_by('codigo')
		if self.q:
			qs = qs.filter(Q(nombre__istartswith=self.q) | Q(codigo__icontains=self.q))
		return qs

class itemActAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return item.objects.none()
		qs = item.objects.filter(activo_f=True).order_by('codigo')
		if self.q:
			qs = qs.filter(Q(nombre__istartswith=self.q) | Q(codigo__icontains=self.q))
		return qs

class personalAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return personal.objects.none()
		qs = personal.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs

class proveedorAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return proveedor.objects.none()
		qs = proveedor.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs

class clienteAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return cliente.objects.none()
		qs = cliente.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(nombre__istartswith=self.q)
		return qs

class ubicacionAutocomplete(autocomplete.Select2QuerySetView):
	def get_queryset(self):
		if not self.request.user.is_authenticated():
			return Ubicacion_gen.objects.none()
		qs = Ubicacion_gen.objects.all().order_by('id')
		if self.q:
			qs = qs.filter(ubicacion_g__istartswith=self.q)
		return qs

"""class crearAlmacen(CreateView):
	form_class = crearAlmacenForm
	template_name = 'almacen/crearalmacen.html'
	success_url = reverse_lazy('almacenes:listaalmacenes')

class updateAlmacen(UpdateView):
	model = almacen
	form_class = crearAlmacenForm
	template_name = 'almacen/crearalmacen.html'
	success_url = reverse_lazy('almacenes:listaalmacenes')

class listaAlmacenes(CreateView,ListView):
	model = almacen
	form_class = searchForm
	template_name = 'almacen/listaalmacenes.html'
	paginate_by = 10
	def get_context_data(self, **kwargs):
		context = super (listaAlmacenes, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list"""

class crearCategoria(CreateView):
	form_class = crearCategoriaForm
	template_name = 'almacen/nuevacategoria.html'
	success_url = reverse_lazy('almacenes:listacategoria')

class updateCategoria(UpdateView):
	model = categoria
	form_class = crearCategoriaForm
	template_name = 'almacen/nuevacategoria.html'
	success_url = reverse_lazy('almacenes:listacategoria')

class listaCategoria(CreateView,ListView):
	model = categoria
	form_class = searchForm
	template_name = 'almacen/listacategoria.html'
	def get_context_data(self, **kwargs):
		context = super (listaCategoria, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

class crearSubCategoria(CreateView):
	model = subcategoria
	form_class = crearSubCategoriaForm
	template_name = 'almacen/nuevasubcategoria.html'
	success_url = 'almacenes:listasubcategoria'

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			cant = self.model.objects.filter(categoria__id=form.categoria.id).count()
			form.codigo = str(cant + 1)
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateSubCategoria(UpdateView):
	model = subcategoria
	form_class = crearSubCategoriaForm
	template_name = 'almacen/nuevasubcategoria.html'
	success_url = reverse_lazy('almacenes:listasubcategoria')

class listaSubCategoria(CreateView,ListView):
	model = subcategoria
	form_class = searchForm
	template_name = 'almacen/listasubcategoria.html'
	def get_context_data(self, **kwargs):
		context = super (listaSubCategoria, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

class crearItem(CreateView):
	model = item
	form_class = crearItemForm
	template_name = 'almacen/nuevoitem.html'
	success_url = 'almacenes:listaitem'

	def get_context_data(self, **kwargs):
		context = super(crearItem, self).get_context_data(**kwargs)
		return context
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			aux = str(form.subcategoria.categoria.id) + '-' + form.subcategoria.codigo + '-' + form.codigo.zfill(4)
			form.codigo = aux
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateItem(UpdateView):
	model = item
	form_class = crearItemForm
	template_name = 'almacen/nuevoitem.html'
	success_url = reverse_lazy('almacenes:listaitem')

class listaItem(CreateView,ListView):
	model = item
	form_class = searchForm
	template_name = 'almacen/listaitem.html'
	def get_context_data(self, **kwargs):
		context = super(listaItem, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		object_list = self.model.objects.all().order_by('id')
		return object_list

class crearItemAct(CreateView):
	model = item_act
	form_class = crearItemActForm
	template_name = 'almacen/nuevoitemact.html'
	success_url = 'almacenes:listaitemact'

	def get_context_data(self, **kwargs):
		context = super(crearItemAct, self).get_context_data(**kwargs)
		return context
	
	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class(request.POST)
		if form.is_valid():
			form = form.save(commit=False)
			form.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form))

class updateItemAct(UpdateView):
	model = item_act
	form_class = crearItemActForm
	template_name = 'almacen/nuevoitemact.html'
	success_url = reverse_lazy('almacenes:listaitemact')

class listaItemAct(ListView):
	model = item_act
	template_name = 'almacen/listaitemact.html'

class listaAltaAct(ListView):
	model = salidaitem
	template_name = 'almacen/listasalidaitemact.html'

	def get_context_data(self, **kwargs):
		context = super (listaAltaAct, self).get_context_data(**kwargs)	
		return context

	def get(self, request, *args, **kwargs):
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list))

	def get_queryset(self):
		object_list = self.model.objects.filter(item__activo_f=True).order_by('id')
		return object_list

class listaSalidaItemf(ListView):
	model = salidaitem
	template_name = 'almacen/listasalidaitemf.html'

# lazys complejos

class ingresoItem(CreateView):
	model_pk = item
	form_class = crearEncabezadoIngresoForm
	template_name = 'almacen/ingresoitem.html'
	success_url = 'almacenes:listaingreso'

	def get_context_data(self, **kwargs):
		context = super(ingresoItem, self).get_context_data(**kwargs)
		return context

	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_ingreso_form_set = crearIngresoItemFormSet()
		crear_personal_form_set = crearPersonalFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_ingreso_form_set=crear_ingreso_form_set, crear_personal_form_set=crear_personal_form_set))

	def post(self, request, *args, **kwargs):
		self.object = None	
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_ingreso_form_set = crearIngresoItemFormSet(request.POST)
		crear_personal_form_set = crearPersonalFormSet(request.POST)
		if form.is_valid() and crear_ingreso_form_set.is_valid() and crear_personal_form_set.is_valid():
			form1 = form.save(commit=False) 
			crear_personal_form_set.instance = form1
			persona = crear_personal_form_set.save(commit=False)
			crear_ingreso_form_set.instance = form1
			ingreso = crear_ingreso_form_set.save(commit=False)
			for formset in ingreso:
				formset.costo_unit_actualizado = formset.costo_unit
				formset.sub_total_actualizado = formset.sub_total
				a=itemalmacen.objects.get_or_create(item = formset.item)
				if a[0].cantidad >= formset.item.maximo:
					return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='El almacen se encuentra lleno, realice salidas del item ' + a[0].item.nombre))
				elif formset.cantidad >= formset.item.maximo:
					return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el maximo permitido de '+ str(formset.item.maximo) + ' del item ' + a[0].item.nombre + ', por favor verifique los datos'))
				else:
					a[0].cantidad = a[0].cantidad + formset.cantidad
					if a[0].cantidad >= formset.item.maximo:
						return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el maximo permitido de '+ str(formset.item.maximo) + ' del item ' + a[0].item.nombre + ', por favor verifique los datos'))
					elif a[0].cantidad < formset.item.minimo:
						return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad esta por debajo del minimo permitido de '+ str(formset.item.minimo) + ' del item ' + a[0].item.nombre + ', por favor verifique los datos'))
					else:
						a[0].precio = 0
						a[0].save()
					form1.save()
					crear_personal_form_set.save()
					crear_ingreso_form_set.save()
			return  HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form,  crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set))

class listaIngresoItem(ListView):
	model = encingresoitem
	form_class = searchForm
	template_name = 'almacen/listaingresoitem.html'

	def get_context_data(self, **kwargs):
		context = super (listaIngresoItem, self).get_context_data(**kwargs)
		if 'form' and 'item' not in context:
			context['form'] = self.form_class()
		return context

class updateIngresoItem(UpdateView):
	model = encingresoitem
	form_class = crearEncabezadoIngresoForm
	template_name = 'almacen/ingresoitem.html'
	success_url = 'almacenes:listaingreso'

	def get_context_data(self, **kwargs):
		context = super(updateIngresoItem, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'form' and 'car' in context:
			context['form'] = self.form_class()
			context['car'] = encingresoitem.objects.get(id=pk1)
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().encingresoitem.pk})

	def get(self, request, *args, **kwargs):
		self.object = self.get_object()
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_ingreso_form_set = updateIngresoItemFormSet(instance=self.object)
		crear_personal_form_set = updatePersonalFormSet(instance=self.object)
		return self.render_to_response(self.get_context_data(form=form, crear_ingreso_form_set=crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set))

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		crear_ingreso_form_set = updateIngresoItemFormSet(request.POST, instance=self.model.objects.get(id=pk1))
		crear_personal_form_set = updatePersonalFormSet(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid() and crear_ingreso_form_set.is_valid() and crear_personal_form_set.is_valid():
			aux = 0
			form1 = form.save(commit=False) 
			crear_personal_form_set.instance = form1
			persona = crear_personal_form_set.save(commit=False)
			crear_ingreso_form_set.instance = form1
			ingreso = crear_ingreso_form_set.save(commit=False)

			for obj in crear_ingreso_form_set.deleted_objects:
				if (obj.cantidad != 0) | (obj.cantidad != None):
					item_alm = itemalmacen.objects.get(item = obj.item)
					item_alm.cantidad -= obj.cantidad
					item_alm.save()
				obj.delete()

			for formset in ingreso:
				a = itemalmacen.objects.get_or_create(item = formset.item)
				i = ingresoitem.objects.filter(item = formset.item, encingresoitem=pk1, id=formset.id).order_by('id').aggregate(total=Sum('cantidad'))
				if a[0].cantidad >= formset.item.maximo:
					return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='El almacen se encuentra lleno, realice salidas del item ' + a.item.nombre))
				elif formset.cantidad >= formset.item.maximo:
					return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el maximo permitido de '+ str(formset.item.maximo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
				else:
					if i['total'] == None:
						a[0].cantidad = a[0].cantidad + formset.cantidad
						if a[0].cantidad >= formset.item.maximo:
							return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el maximo permitido de '+ str(formset.item.maximo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
						elif a[0].cantidad < formset.item.minimo:
							return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad esta por debajo del minimo permitido de '+ str(formset.item.minimo) + ' del item ' + a[0].item.nombre + ', por favor verifique los datos'))
						else:
							a[0].save()
					elif(formset.cantidad > i['total']):
						aux = formset.cantidad - i['total']
						a[0].cantidad = a[0].cantidad + aux
						if a[0].cantidad >= formset.item.maximo:
							return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el maximo permitido de '+ str(formset.item.maximo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
						else:
							a[0].save()
					elif(formset.cantidad < i['total']):
						aux = i['total'] - formset.cantidad
						a[0].cantidad = a[0].cantidad - aux
						if a[0].cantidad >= formset.item.maximo:
							return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el maximo permitido de '+ str(formset.item.maximo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
						elif a[0].cantidad < formset.item.minimo:
							return self.render_to_response(self.get_context_data(form = form, crear_ingreso_form_set = crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad esta por debajo del minimo permitido de '+ str(formset.item.minimo) + ' del item ' + a[0].item.nombre + ', por favor verifique los datos'))
						else:
							a[0].save()
					form1.save()
					crear_personal_form_set.save()
					formset.save()
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_ingreso_form_set=crear_ingreso_form_set, crear_personal_form_set = crear_personal_form_set, error_formset=crear_ingreso_form_set.non_form_errors()))

class salidaItem(CreateView):
	model_pk = item
	form_class = crearEncabezadoSalidaForm
	second_form_class = crearSalidaItemForm
	template_name = 'almacen/salidaitem.html'
	success_url = 'almacenes:listasalida'

	def get_context_data(self, **kwargs):
		context = super (salidaItem, self).get_context_data(**kwargs)
		if 'form' not in context:
			context['form'] = self.form_class()
		return context

	def get(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_personal_form_set = crearPersonalSFormSet()
		crear_salida_form_set = crearSalidaItemFormSet()
		return self.render_to_response(self.get_context_data(form=form, crear_salida_form_set=crear_salida_form_set, crear_personal_form_set=crear_personal_form_set))
	
	def post(self, request, *args, **kwargs):
		self.object = None
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_personal_form_set = crearPersonalSFormSet(request.POST)
		crear_salida_form_set = crearSalidaItemFormSet(request.POST)
		if form.is_valid() and crear_salida_form_set.is_valid() and crear_personal_form_set.is_valid():
			from django.db.models import Avg
			form1 = form.save(commit=False)
			crear_salida_form_set.instance = form1
			crear_personal_form_set.instance = form1
			persona = crear_personal_form_set.save(commit=False)
			salida = crear_salida_form_set.save(commit=False)
			for formset in salida:
				try:
					a = itemalmacen.objects.get(item = formset.item)
				except Exception as e:
					return self.render_to_response(self.get_context_data(form1 = form1, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='El item ' + str(formset.item)+ ' no se ha ingresado aún'))
				if a.cantidad < formset.cantidad:
					return self.render_to_response(self.get_context_data(form1 = form1, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el stock actual de '+ str(a.cantidad) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
				elif a.cantidad < formset.item.minimo:
					return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el minimo permitido de '+ str(formset.item.minimo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
				else:
					try:
						aux = ingresoitem.objects.filter(item = formset.item, encingresoitem__fecha__lte=form1.fecha, encingresoitem__cerrado=False).aggregate(total=Sum('cantidad'))
						aux1 = salidaitem.objects.filter(item = formset.item, encsalidaitem__fecha__lte=form1.fecha, encsalidaitem__cerrado=False).aggregate(total=Sum('cantidad'))
						aux2 = ingresoitem.objects.filter(item = formset.item, encingresoitem__fecha__lte=form1.fecha, encingresoitem__cerrado=False).aggregate(total=Sum('sub_total'))

						if aux1['total'] == None:
							aux1['total'] = 0
						res = aux['total'] - aux1['total']
						res1 = aux2['total']/Decimal(res)
						formset.costo_unit = round(res1, 2)
						formset.sub_total = Decimal(formset.cantidad) * res1
						a.cantidad = a.cantidad - formset.cantidad
						
						form1.save()
						crear_salida_form_set.save()
						crear_personal_form_set.save()
						form.save_m2m()
						a.save()
					except Exception as e:
						raise ValueError('La cantidad de salida excede el saldo existente a la fecha especificada, o la fecha de salida es anterior a la fecha de ingreso, verifique los datos')

			return  HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set))

class listaSalidaItem(CreateView,ListView):
	model = encsalidaitem
	form_class = searchForm
	template_name = 'almacen/listasalidaitem.html'

	def get_context_data(self, **kwargs):
		context = super (listaSalidaItem, self).get_context_data(**kwargs)
		if 'form' and 'item' not in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
				#kwargs['id']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.all().order_by('id')
		return object_list

class updateSalidaItem(UpdateView):
	model = encsalidaitem
	form_class = crearEncabezadoSalidaForm
	template_name = 'almacen/salidaitem.html'
	success_url = 'almacenes:listasalida'

	def get_context_data(self, **kwargs):
		context = super(updateSalidaItem, self).get_context_data(**kwargs)
		pk1 = self.kwargs['pk']
		if 'form' and 'car' in context:
			context['form'] = self.form_class()
			context['car'] = encsalidaitem.objects.get(id=pk1)
		return context

	def get_success_url(self):
		return reverse_lazy(self.success_url, kwargs = {'pk': self.get_object().encsalidaitem.pk})

	def get(self, request, *args, **kwargs):
		self.object = self.get_object()
		form_class = self.get_form_class()
		form = self.get_form(form_class)
		crear_salida_form_set = crearSalidaItemFormSet(instance=self.object)
		crear_personal_form_set = updatePersonalSFormSet(instance=self.object)
		return self.render_to_response(self.get_context_data(form=form, crear_salida_form_set=crear_salida_form_set, crear_personal_form_set=crear_personal_form_set))

	def post(self, request, *args, **kwargs):
		self.object = self.get_object
		pk1=self.kwargs['pk']
		form = self.form_class(request.POST, instance=self.model.objects.get(id=pk1))
		crear_salida_form_set = crearSalidaItemFormSet(request.POST, instance=self.model.objects.get(id=pk1))
		crear_personal_form_set = updatePersonalSFormSet(request.POST, instance=self.model.objects.get(id=pk1))
		if form.is_valid() and crear_salida_form_set.is_valid():
			form1 = form.save(commit=False)
			crear_salida_form_set.instance = form1
			persona = crear_personal_form_set.save(commit=False)
			salida = crear_salida_form_set.save(commit=False)

			for obj in crear_salida_form_set.deleted_objects:
				if (obj.cantidad != 0) | (obj.cantidad != None):
					item_alm = itemalmacen.objects.get(item = obj.item)
					item_alm.cantidad += obj.cantidad
					item_alm.save()
				obj.delete()

			for formset in salida:
				a = itemalmacen.objects.get(item = formset.item)
				if a.cantidad < formset.cantidad:
					return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el stock actual de '+ str(a.cantidad) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
				elif a.cantidad < formset.item.minimo:
					return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el minimo permitido de '+ str(formset.item.minimo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
				else:
					try:
						aux = ingresoitem.objects.filter(item = formset.item, encingresoitem__fecha__lte=form1.fecha, encingresoitem__cerrado=False).aggregate(total=Sum('cantidad'))
						aux1 = salidaitem.objects.filter(item = formset.item, encsalidaitem__fecha__lte=form1.fecha, encsalidaitem__cerrado=False).aggregate(total=Sum('cantidad'))
						aux2 = ingresoitem.objects.filter(item = formset.item, encingresoitem__fecha__lte=form1.fecha, encingresoitem__cerrado=False).aggregate(total=Sum('sub_total'))
						
						if aux1['total'] == None:
							aux1['total'] = 0
						res = aux['total'] - aux1['total']
						res1 = aux2['total']/Decimal(res)
						formset.costo_unit = round(res1, 2)
						formset.sub_total = Decimal(formset.cantidad) * res1
						s = salidaitem.objects.filter(item = formset.item, encsalidaitem=pk1, id=formset.id).aggregate(total=Sum('cantidad'))
						if s['total'] == None:
							a.cantidad = a.cantidad - formset.cantidad
							if a.cantidad < formset.item.minimo:
								return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el mínimo permitido de '+ str(formset.item.minimo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
							else:
								a.save()
						elif(formset.cantidad > s['total']):
							aux = formset.cantidad - s['total']
							a.cantidad = a.cantidad - aux
							if a.cantidad < formset.item.minimo:
								return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el mínimo permitido de '+ str(formset.item.minimo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
							else:
								a.save()
						elif(formset.cantidad < s['total']):
							aux = s['total'] - formset.cantidad
							a.cantidad = a.cantidad + aux
							if a.cantidad < formset.item.minimo:
								return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad excede el mínimo permitido de '+ str(formset.item.minimo) + ' del item ' + a.item.nombre + ', por favor verifique los datos'))
							else:
								a.save()
						formset.save()
						crear_personal_form_set.save()
						form.save_m2m()
					except Exception as e:
						return self.render_to_response(self.get_context_data(form = form, crear_salida_form_set = crear_salida_form_set, crear_personal_form_set = crear_personal_form_set, error='La cantidad de salida de uno de los items excede el saldo existente a la fecha especificada, o la fecha de salida es anterior a la fecha de ingreso, verifique los datos'))
			form1.save()
					
			return HttpResponseRedirect(reverse_lazy(self.success_url))
		else:
			return self.render_to_response(self.get_context_data(form=form, crear_salida_form_set=crear_salida_form_set, error_formset=crear_salida_form_set.non_form_errors()))

class listaItemT(CreateView,ListView):
	model = item
	form_class = searchForm
	template_name = 'almacen/listaitemt.html'
	def get_context_data(self, **kwargs):
		context = super (listaItemT, self).get_context_data(**kwargs)	
		if 'form' in context:
			context['form'] = self.form_class()
		return context
	def get(self, request, *args, **kwargs):
		self.object = self.get_object
		form = self.form_class()
		if request.GET:
			form = self.form_class(request.GET)
		self.object_list = self.get_queryset()
		return self.render_to_response(self.get_context_data(object_list=self.object_list , form=form))
	def get_queryset(self):
		id = None
		if self.request.method == "GET":
			form = self.form_class(self.request.GET)
			if form.is_valid():
				id = form.cleaned_data['search']
		if (id):
			#object_list = self.model.objects.filter(name__icontains = id)
			object_list = self.model.objects.filter(id = id)
		else:
			object_list = self.model.objects.filter().order_by('id')
		return object_list

class generarComprobanteIng(PDFTemplateView):
	template_name = "almacen/comprobante_ingresoitem.html"
	def get_context_data(self, **kwargs):
		context = super(generarComprobanteIng, self).get_context_data(**kwargs)
		pk1 = self.kwargs.get('pk')
		try:
			context['res'] = encingresoitem.objects.get(id=pk1)
			context['user'] = self.request.user
		except:
			raise Http404
		return context

class generarComprobanteSal(PDFTemplateView):
	template_name = "almacen/comprobante_salidaitem.html"
	def get_context_data(self, **kwargs):
		context = super(generarComprobanteSal, self).get_context_data(**kwargs)
		pk1 = self.kwargs.get('pk')
		try:
			context['res'] = encsalidaitem.objects.get(id=pk1)
			context['user'] = self.request.user
		except:
			raise Http404
		return context

class reporteDiarioETemplate(FormView):
	form_class = reporteDiarioEForm
	template_name = "almacen/rep_diario_e_temp.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDiarioETemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteDiarioE(PDFTemplateView):
	template_name = "almacen/rep_diario_e.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDiarioE, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		context['res'] = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).order_by('encingresoitem__fecha')
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2

		return context

class reporteDiarioSTemplate(FormView):
	form_class = reporteDiarioEForm
	template_name = "almacen/rep_diario_s_temp.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDiarioSTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteDiarioS(PDFTemplateView):
	template_name = "almacen/rep_diario_s.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDiarioS, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")	
		context['res'] = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).order_by('encsalidaitem__fecha')
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2
		return context

class reporteKardexVal(PDFTemplateView):
	template_name = "almacen/kardex_fisico_v.html"
	def get_context_data(self, **kwargs):
		context = super(reporteKardexVal, self).get_context_data(**kwargs)
		op = 0
		op1 = 0
		suma_c_e = 0
		suma_sub_e = 0
		suma_c_s = 0
		suma_sub_s = 0
		saldo = []
		saldo1 = []
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		aux1 = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), item__codigo=campo2, encingresoitem__cerrado=False)
		aux2 = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), item__codigo=campo2, encsalidaitem__cerrado=False)
		lista = list(sorted(chain(aux1, aux2), key=methodcaller('encabezado_f')))
		
		for i in lista:
			if i.encabezado().tipo() == "E":
				suma_c_e = suma_c_e + i.cantidad
				suma_sub_e = suma_sub_e + i.sub_total
				op = op + i.cantidad
				op1 = op1 + i.sub_total
				saldo.append(op)
				saldo1.append(op1)
			else:
				suma_c_s = suma_c_s + i.cantidad
				suma_sub_s = suma_sub_s + i.sub_total
				op = op - i.cantidad
				op1 = op1 - i.sub_total
				saldo.append(op)
				saldo1.append(op1)

		lista_res = zip(lista, saldo, saldo1)
		context['cant_e'] = suma_c_e
		context['sub_e'] = suma_sub_e
		context['cant_s'] = suma_c_s
		context['sub_s'] = suma_sub_s
		context['res'] = lista_res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		
		return context

class resumenKardexTemplate(FormView):
	form_class = reporteResKardexForm
	template_name = "almacen/res_kardex.html"
	def get_context_data(self, **kwargs):
		context = super(resumenKardexTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class resumenKardex(PDFTemplateView):
	template_name = "almacen/resumenKardex.html"
	model_pk = item
	def get_context_data(self, **kwargs):
		context = super(resumenKardex, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		f3 = f1 - datetime.timedelta(days=1)
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		fecha3 = f3.strftime("%d/%m/%Y")
		b = 0
		c = 0

		d = 0
		e = 0

		f = 0

		tot_saldo_a = 0
		tot_saldo_e = 0
		tot_saldo_s = 0
		tot_saldo = 0

		salida = []

		saldo_c = []
		cantidad_e = []
		cantidad_s = []
		saldo_sub = []
		sub_total_e = []
		sub_total_s = []

		saldo_c_a = []
		cantidad_e_a = []
		cantidad_s_a = []
		saldo_sub_a = []
		sub_total_e_a = []
		sub_total_s_a = []

		saldo_c_t = []
		saldo_sub_t = []

		try:
			aux_a = item.objects.filter(ingresoitem__encingresoitem__fecha__lte=f3, ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__cantidad'))
			aux1_a = item.objects.filter(salidaitem__encsalidaitem__fecha__lte=f3, salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__cantidad'))
			aux2_a = item.objects.filter(ingresoitem__encingresoitem__fecha__lte=f3, ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__sub_total'))
			aux3_a = item.objects.filter(salidaitem__encsalidaitem__fecha__lte=f3, salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__sub_total'))
				
			resp = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).order_by('codigo').distinct()
			resp_a = item.objects.filter(ingresoitem__encingresoitem__fecha__lte=f3, ingresoitem__encingresoitem__cerrado=False).order_by('codigo').distinct()
	
			aux = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__cantidad'))
			aux1 = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__cantidad'))
			aux2 = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__sub_total'))
			aux3 = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__sub_total'))

			for q in aux_a:
				cantidad_e_a.append(q.total)
	
			if aux1_a:
				for i in aux_a:
					cont = 0
					if i in aux1_a:
						cantidad_s_a.append(aux1_a[cont].total)
						cont+=1
					else:
						cantidad_s_a.append(0)
			
			else:
				for i in aux_a:
					cantidad_s_a.append(0)
	
	
			for x in aux2_a:
				sub_total_e_a.append(x.total)
	
			if aux3_a:
				cont1 = 0
				for i in aux2_a:
					if i in aux3_a:
						sub_total_s_a.append(aux3_a[cont1].total)
						cont1+=1
					else:
						sub_total_s_a.append(0)
			else:
				for i in aux2_a:
					sub_total_s_a.append(0)
			
			for r in resp:
				if r in resp_a:
					saldo_c_a.append(cantidad_e_a[d] - cantidad_s_a[d])
					saldo_sub_a.append(sub_total_e_a[d] - sub_total_s_a[d])
					d+=1
				else:
					saldo_c_a.append(0)
					saldo_sub_a.append(0)
	
			for i in saldo_sub_a:
				tot_saldo_a = tot_saldo_a + i
				
			for i in aux:
				cantidad_e.append(i.total)
	
			if aux1:
				cont2 = 0
				for i in aux:
					if i in aux1:
						cantidad_s.append(aux1[cont2].total)
						cont2+=1
					else:
						cantidad_s.append(0)
			else:
				for i in aux1:
					cantidad_s.append(0)
	
			for k in item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo'):
				saldo_c.append(cantidad_e[b] - cantidad_s[b])
				b+=1
	
			for m in aux2:
				sub_total_e.append(m.total)
	
			if aux3:
				cont3 = 0
				for i in aux2:
					if i in aux3:
						sub_total_s.append(aux3[cont3].total)
						cont3+=1
					else:
						sub_total_s.append(0)
	
			else:
				for i in aux2:
					sub_total_s.append(0)
	
			for i in sub_total_e:
				tot_saldo_e = tot_saldo_e + i
	
			for i in sub_total_s:
				tot_saldo_s = tot_saldo_s + i
	
			for o in item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo'):
				saldo_sub.append(sub_total_e[c] - sub_total_s[c])
				c+=1
	
			res = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total_cant_e=Sum('ingresoitem__cantidad'), total_sub_e=Sum('ingresoitem__sub_total'))
			res1 = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total_cant_s=Sum('salidaitem__cantidad'), total_sub_s=Sum('salidaitem__sub_total'))
	
			if res1:
				x = 0
				for i in res:
					if i in res1:
						salida.append(res1[x])
						x+=1
					else:
						salida.append(0)
			else:
				for i in res:
					salida.append(0)
	
	
			for p in item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo'):
				saldo_c_t.append(saldo_c[f] + saldo_c_a[f])
				saldo_sub_t.append(saldo_sub[f] + saldo_sub_a[f])
				f+=1
	
			for i in saldo_sub_t:
				tot_saldo = tot_saldo + i
	
			lista = zip(res, salida, saldo_c_t, saldo_sub_t, saldo_c_a, saldo_sub_a)
	
			context['res'] = lista
			context['total_a'] = tot_saldo_a
			context['total_e'] = tot_saldo_e
			context['total_s'] = tot_saldo_s
			context['total_saldo'] = tot_saldo
			context['fecha1'] = fecha1
			context['fecha2'] = fecha2
			context['fecha3'] = fecha3
			return context

		except Exception as e:
			return None

class reporteKardexTemplate(FormView):
	form_class = repGeneralForm
	template_name = 'almacen/reporte_general.html'
	def get_context_data(self, **kwargs):
		context = super(reporteKardexTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteGeneral(PDFTemplateView):
	template_name = "almacen/rep_gral.html"
	def get_context_data(self, **kwargs):
		context = super(reporteGeneral, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		f3 = f1 - datetime.timedelta(days=1)
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		fecha3 = f3.strftime("%d/%m/%Y")
		
		if campo2 == 'T':
			res = item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2)) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2)), Q(ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo')
		else:
			res = item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2)) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2)), Q(ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__cerrado=False), categoria=campo2).distinct().order_by('codigo')

		context['res'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['fecha3'] = fecha3
		context['f1'] = f1
		context['f2'] = f2
		context['f3'] = f3
		context['cat'] = campo2
		
		return context

class ResumenEySTemplate(FormView):
	form_class = repResumenEySForm
	template_name = 'almacen/resumen_eys.html'
	def get_context_data(self, **kwargs):
		context = super(ResumenEySTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteResumenEyS(PDFTemplateView):
	template_name = "almacen/reporte_resumen_eys.html"
	def get_context_data(self, **kwargs):
		context = super(reporteResumenEyS, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		
		if campo2 == 'E':
			res = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).order_by('id')
			sumas = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total_i=Sum('total_iva'), total_imp_i=Sum('importe_iva'), total=Sum('total'))
			context['total_i'] = sumas['total_i']
			context['total_imp_i'] = sumas['total_imp_i']
			context['total'] = sumas['total']
		else:
			res = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).order_by('id')
			sumas_s = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total=Sum('salidaitem__sub_total'))
			context['total_s'] = sumas_s['total']

		context['res'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2
		context['tipo_mov'] = campo2
		
		return context

class DetalleEntradasTemplate(FormView):
	form_class = repDetalleEntradasForm
	template_name = 'almacen/detalle_entradas.html'
	def get_context_data(self, **kwargs):
		context = super(DetalleEntradasTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteDetalleEntradas(PDFTemplateView):
	template_name = "almacen/reporte_detalle_entradas.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDetalleEntradas, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		
		if campo2 == 'T':
			res = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).order_by('encingresoitem__id').order_by('encingresoitem__tipo_entrada')
			sumas = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total_i=Sum('total_iva'), total_imp_i=Sum('importe_iva'), total=Sum('total'))
			context['total_i'] = sumas['total_i']
			context['total_imp_i'] = sumas['total_imp_i']
			context['total'] = sumas['total']
		else:
			res = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo')
			sumas = ingresoitem.objects.filter(item__in=res, encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total=Sum('costo_total'), total_imp_i=Sum(F('costo_total') - F('sub_total')), total_i=Sum('sub_total'))
			context['total_cant'] = sumas['total_cant']
			context['total'] = sumas['total']
			context['total_imp_i'] = sumas['total_imp_i']
			context['total_i'] = sumas['total_i']
			
			

		context['res'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2
		context['tipo_mov'] = campo2
		
		return context

class DetalleSalidasTemplate(FormView):
	form_class = repDetalleSalidasForm
	template_name = 'almacen/detalle_salidas.html'
	def get_context_data(self, **kwargs):
		context = super(DetalleSalidasTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteDetalleSalidas(PDFTemplateView):
	template_name = "almacen/reporte_detalle_salidas.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDetalleSalidas, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		
		if campo2 == 'T':
			res = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).order_by('encsalidaitem__tipo_salida')
			sumas = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total=Sum('salidaitem__sub_total'))
			context['total'] = sumas['total']
		else:
			res = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo')
			sumas = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total=Sum('sub_total'))
			context['total_cant'] = sumas['total_cant']
			context['total'] = sumas['total']

		context['res'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2
		context['tipo_mov'] = campo2
		
		return context

class DetalleClientesTemplate(FormView):
	form_class = repDetalleClienteForm
	template_name = 'almacen/detalle_clientes.html'
	def get_context_data(self, **kwargs):
		context = super(DetalleClientesTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteDetalleClientes(PDFTemplateView):
	template_name = "almacen/reporte_detalle_clientes.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDetalleClientes, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		
		if campo2 == 'T':
			res = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).distinct().order_by('encsalidaitem__id')
			sumas = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))
		else:
			res = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cliente=campo2, encsalidaitem__cerrado=False).distinct().order_by('encsalidaitem__id')
			sumas = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cliente=campo2, encsalidaitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))

		context['total_cant'] = sumas['total_cant']
		context['total_costo_u'] = sumas['total_costo_u']
		context['total'] = sumas['total']
		context['res'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2
		context['tipo_mov'] = campo2
		
		return context

class DetalleProveedoresTemplate(FormView):
	form_class = repDetalleProveedorForm
	template_name = 'almacen/detalle_proveedores.html'
	def get_context_data(self, **kwargs):
		context = super(DetalleProveedoresTemplate, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class reporteDetalleProveedores(PDFTemplateView):
	template_name = "almacen/reporte_detalle_proveedores.html"
	def get_context_data(self, **kwargs):
		context = super(reporteDetalleProveedores, self).get_context_data(**kwargs)
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		
		if campo2 == 'T':
			res = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).distinct().order_by('encingresoitem__id')
			sumas = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))
		else:
			res = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__proveedor=campo2, encingresoitem__cerrado=False).distinct().order_by('encingresoitem__id')
			sumas = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__proveedor=campo2, encingresoitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))

		context['total_cant'] = sumas['total_cant']
		context['total_costo_u'] = sumas['total_costo_u']
		context['total'] = sumas['total']
		context['res'] = res
		context['fecha1'] = fecha1
		context['fecha2'] = fecha2
		context['f1'] = f1
		context['f2'] = f2
		context['tipo_mov'] = campo2
		
		return context

class uploadCotizacionTC(FormView):
	form_class = XlsInputForm
	success_url = reverse_lazy('contabilidad:listacotizacion_tc')
	template_name = 'contabilidad/cotizacion_tcupload.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST, request.FILES)
		if form.is_valid():
			input_excel = request.FILES['input_excel']
			book = xlrd.open_workbook(file_contents=input_excel.read())
			sheet = book.sheets()[0]
			aux = []
			for i in range(0, sheet.nrows):
				f = sheet.cell_value(i, 0)
				f1 = datetime.datetime(*xlrd.xldate_as_tuple(f, book.datemode))
				f2 = f1.strftime("%Y-%m-%d")
				aux.append(f2)
			t_c = sheet.col_values(1, 0)
			u = sheet.col_values(2, 0)
			for fecha, tipo_c, ufv in zip(aux, t_c, u):
				cotizacion, created = Cotizacion.objects.get_or_create(fecha=fecha, tipo_c=tipo_c, ufv=ufv)
				cotizacion.save()
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

#Actualización de inventarios
class actualizarInventario(FormView):
	form_class = actualizarInventarioForm
	success_url = reverse_lazy('almacenes:listasalida')
	template_name = 'almacen/act_inventario.html'

	def get(self, request, *args, **kwargs):
		form = self.form_class()
		return self.render_to_response({'form':form})

	def post(self, request, *args, **kwargs):
		form = self.form_class(request.POST)
		campo = self.request.POST.get('campo')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		if form.is_valid():
			ingreso = ingresoitem.objects.filter(encingresoitem__fecha__lte=f1)
			for i in ingreso:
				fecha = i.encingresoitem.fecha
				ufv_p = Cotizacion.objects.get(fecha=fecha).ufv
				ufv_u = Cotizacion.objects.get(fecha=f1).ufv
				u = ufv_u/ufv_p
				ufv = round(u, 7)
				aux1 = i.costo_unit * Decimal(ufv)
				aux2 = i.sub_total * Decimal(ufv)
				res = round(aux1, 2)
				res1 = round(aux2, 2)
				ingresoitem.objects.filter(encingresoitem__fecha__lte=f1, id=i.id).update(costo_unit_actualizado=res, sub_total_actualizado=res1)
			return HttpResponseRedirect(self.success_url)
		else:
			return render(request, {'form':form})

class RepGeneralTemplateExcel(FormView):
	form_class = repGeneralForm
	template_name = 'almacen/temp_reporte_gral_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepGeneralTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepGeneralExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		f3 = f1 - datetime.timedelta(days=1)
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		fecha3 = f3.strftime("%d/%m/%Y")
		iteration = 1
		
		if campo2 == 'T':
			res = categoria.objects.filter(Q(item__ingresoitem__encingresoitem__fecha__range=(f1, f2), item__ingresoitem__encingresoitem__cerrado=False) | Q(item__salidaitem__encsalidaitem__fecha__range=(f1, f2), item__salidaitem__encsalidaitem__cerrado=False)).order_by('id').distinct()
			res1 = item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).order_by('codigo').distinct()
			total_queries = item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().count()
		else:
			res = categoria.objects.filter(Q(item__ingresoitem__encingresoitem__fecha__range=(f1, f2), item__ingresoitem__encingresoitem__cerrado=False) | Q(item__salidaitem__encsalidaitem__fecha__range=(f1, f2), item__salidaitem__encsalidaitem__cerrado=False), id=campo2).order_by('id').distinct()
			res1 = item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False), categoria=campo2).order_by('codigo').distinct()
			total_queries = item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2)) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2)), categoria=campo2).distinct().count()

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'REPORTE GENERAL DEL ' + str(fecha1) + ' al ' + str(fecha2)
		ws['A2'] = 'Código'
		ws['B2'] = 'Descripción'
		ws['C2'] = 'Unid.'
		ws['D2'] = 'Saldo al ' + str(fecha3)
		ws['F2'] = 'Entradas'
		ws['H2'] = 'Salidas'
		ws['J2'] = 'Saldos'
		ws['D3'] = 'Cant.'
		ws['E3'] = 'Valor'
		ws['F3'] = 'Cant.'
		ws['G3'] = 'Valor'
		ws['H3'] = 'Cant.'
		ws['I3'] = 'Valor'
		ws['J3'] = 'Cant.'
		ws['K3'] = 'Valor'

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['K2'].alignment = Alignment(horizontal="center", vertical="center")

		ws.column_dimensions['A'].width = 10
		ws.column_dimensions['B'].width = 50
		ws.column_dimensions['C'].width = 10
		ws.column_dimensions['D'].width = 10
		ws.column_dimensions['E'].width = 10
		ws.column_dimensions['F'].width = 10
		ws.column_dimensions['G'].width = 10
		ws.column_dimensions['H'].width = 10
		ws.column_dimensions['I'].width = 10
		ws.column_dimensions['J'].width = 10
		ws.column_dimensions['K'].width = 10

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)
		ws['J2'].font = Font(bold=True)
		ws['K2'].font = Font(bold=True)
		ws['D3'].font = Font(bold=True)
		ws['E3'].font = Font(bold=True)
		ws['F3'].font = Font(bold=True)
		ws['G3'].font = Font(bold=True)
		ws['H3'].font = Font(bold=True)
		ws['I3'].font = Font(bold=True)
		ws['J3'].font = Font(bold=True)
		ws['K3'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:K1')
		ws.merge_cells('D2:E2')
		ws.merge_cells('F2:G2')
		ws.merge_cells('H2:I2')
		ws.merge_cells('J2:K2')
		
		cont = 4

		for categorias in res:
			ws.cell(row = cont, column = 1).value = categorias.nombre
			cont += 1
			for subcategorias in categorias.subcategoria_set.filter(Q(item__ingresoitem__encingresoitem__fecha__range=(f1, f2)) | Q(item__salidaitem__encsalidaitem__fecha__range=(f1, f2)), Q(item__ingresoitem__encingresoitem__cerrado=False) | Q(item__salidaitem__encsalidaitem__cerrado=False)).distinct():
				ws.cell(row = cont, column = 1).value = subcategorias.nombre
				cont += 1
				for items in subcategorias.item_set.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2)) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2)), Q(ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__cerrado=False)).distinct():
					ws.cell(row = cont, column = 1).value = items.codigo
					ws.cell(row = cont, column = 2).value = items.nombre
					ws.cell(row = cont, column = 3).value = items.unidad.simbolo
					ws.cell(row = cont, column = 4).value = items.total_cant_a(f1)
					ws.cell(row = cont, column = 5).value = items.total_subt_a(f1)
					ws.cell(row = cont, column = 6).value = items.total_cant_e(f1, f2)
					ws.cell(row = cont, column = 7).value = items.total_subt_e(f1, f2)
					ws.cell(row = cont, column = 8).value = items.total_cant_s(f1, f2)
					ws.cell(row = cont, column = 9).value = items.total_subt_s(f1, f2)
					ws.cell(row = cont, column = 10).value = items.total_cant(f1, f2)
					ws.cell(row = cont, column = 11).value = items.total_subt(f1, f2)
					cont += 1
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + subcategorias.nombre.upper()
				ws.cell(row = cont, column = 5).value = subcategorias.subcategoria_a_t(f1, res1)
				ws.cell(row = cont, column = 7).value = subcategorias.total_sub_e(f1, f2)
				ws.cell(row = cont, column = 9).value = subcategorias.total_sub_s(f1, f2)
				ws.cell(row = cont, column = 11).value = subcategorias.subcategoria_t(f1, f2, res1)
				cont += 2
			ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + categorias.nombre.upper()
			ws.cell(row = cont, column = 5).value = categorias.categoria_a_t(f1, res1)
			ws.cell(row = cont, column = 7).value = categorias.total_categ_e(f1, f2, res1)
			ws.cell(row = cont, column = 9).value = categorias.total_categ_s(f1, f2, res1)
			ws.cell(row = cont, column = 11).value = categorias.categoria_t(f1, f2, res1)
			cont += 2
		for i in res1:
			ws.cell(row = cont, column = 1).value = 'TOTALES'
			ws.cell(row = cont, column = 5).value = i.subtotal_t_a(f1, res1)
			ws.cell(row = cont, column = 7).value = i.ingreso_t(f1, f2, res1)
			ws.cell(row = cont, column = 9).value = i.salida_t(f1, f2, res1)
			ws.cell(row = cont, column = 11).value = i.saldo_t(f1, f2, res1)
			break
			
		nombre_archivo = "Reporte General.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepDiarioEntTemplateExcel(FormView):
	form_class = reporteDiarioEForm
	template_name = 'almacen/temp_reporte_diarioent_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepDiarioEntTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepDiarioEntExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		res = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).order_by('fecha')

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'REPORTE DIARIO DE INGRESOS DEL ' + str(fecha1) + ' ' + str(fecha2)
		ws['A2'] = 'Fecha'
		ws['B2'] = 'Proveedor'
		ws['C2'] = 'Responsable'
		ws['D2'] = 'Código'
		ws['E2'] = 'Descripción'
		ws['F2'] = 'Cantidad'
		ws['G2'] = 'Unidad'
		ws['H2'] = 'Costo Unitario Bs.'
		ws['I2'] = 'Importe Bs.'

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I2'].alignment = Alignment(horizontal="center", vertical="center")

		ws.column_dimensions['A'].width = 15
		ws.column_dimensions['B'].width = 30
		ws.column_dimensions['C'].width = 30
		ws.column_dimensions['D'].width = 15
		ws.column_dimensions['E'].width = 45
		ws.column_dimensions['F'].width = 10
		ws.column_dimensions['G'].width = 8
		ws.column_dimensions['H'].width = 20
		ws.column_dimensions['I'].width = 15

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:I1')
		
		cont = 3

		for encingreso in res:
			for ingresos in encingreso.ingresoitem_set.all():
				ws.cell(row = cont, column = 1).value = ingresos.encingresoitem.fecha
				ws.cell(row = cont, column = 2).value = ingresos.encingresoitem.proveedor.nombre
				for personals in ingresos.encingresoitem.personal.all(): 
					ws.cell(row = cont, column = 3).value = personals.nombre + ' ' + personals.ap_paterno
					break
				ws.cell(row = cont, column = 4).value = ingresos.item.codigo
				ws.cell(row = cont, column = 5).value = ingresos.item.nombre
				ws.cell(row = cont, column = 6).value = ingresos.cantidad
				ws.cell(row = cont, column = 7).value = ingresos.item.unidad.simbolo
				ws.cell(row = cont, column = 8).value = ingresos.costo_unit
				ws.cell(row = cont, column = 9).value = ingresos.sub_total
				cont += 1
			for ingresos in encingreso.ingresoitem_set.all():
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + ingresos.item.subcategoria.nombre
				ws.cell(row = cont, column = 9).value = ingresos.total(encingreso.fecha)
				cont += 1
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + ingresos.item.categoria.nombre
				ws.cell(row = cont, column = 9).value = ingresos.total1(encingreso.fecha)
				cont += 1
				break
			
		nombre_archivo = "Reporte Diario de Entradas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepDiarioSalTemplateExcel(FormView):
	form_class = reporteDiarioEForm
	template_name = 'almacen/temp_reporte_diariosal_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepDiarioSalTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepDiarioSalExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		res = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).order_by('fecha')

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'REPORTE DIARIO DE SALIDAS DEL ' + str(fecha1) + ' al ' + str(fecha2)
		ws['A2'] = 'Fecha'
		ws['B2'] = 'Destino'
		ws['C2'] = 'Responsable'
		ws['D2'] = 'Código'
		ws['E2'] = 'Descripción'
		ws['F2'] = 'Cantidad'
		ws['G2'] = 'Unidad'
		ws['H2'] = 'Costo Unitario Bs.'
		ws['I2'] = 'Importe Bs.'

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I2'].alignment = Alignment(horizontal="center", vertical="center")

		ws.column_dimensions['A'].width = 12
		ws.column_dimensions['B'].width = 40
		ws.column_dimensions['C'].width = 30
		ws.column_dimensions['D'].width = 10
		ws.column_dimensions['E'].width = 45
		ws.column_dimensions['F'].width = 10
		ws.column_dimensions['G'].width = 8
		ws.column_dimensions['H'].width = 20
		ws.column_dimensions['I'].width = 15

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')
		
		ws.merge_cells('A1:I1')
		
		cont = 3

		for encsalida in res:
			for salidas in encsalida.salidaitem_set.all():
				ws.cell(row = cont, column = 1).value = salidas.encsalidaitem.fecha
				ws.cell(row = cont, column = 2).value = salidas.encsalidaitem.cliente.nombre
				for personals in salidas.encsalidaitem.personal.all(): 
					ws.cell(row = cont, column = 3).value = personals.nombre + ' ' + personals.ap_paterno
					break
				ws.cell(row = cont, column = 4).value = salidas.item.codigo
				ws.cell(row = cont, column = 5).value = salidas.item.nombre
				ws.cell(row = cont, column = 6).value = salidas.cantidad
				ws.cell(row = cont, column = 7).value = salidas.item.unidad.simbolo
				ws.cell(row = cont, column = 8).value = salidas.costo_unit
				ws.cell(row = cont, column = 9).value = salidas.sub_total
				cont += 1
			for salidas in encsalida.salidaitem_set.all():
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + salidas.item.subcategoria.nombre
				ws.cell(row = cont, column = 9).value = salidas.total(encsalida.fecha)
				cont += 1
				ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + salidas.item.categoria.nombre
				ws.cell(row = cont, column = 9).value = salidas.total1(encsalida.fecha)
				cont += 1
				break
			
		nombre_archivo = "Reporte Diario de Salidas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepKardexValTemplateExcel(FormView):
	form_class = repKardexValForm
	template_name = 'almacen/temp_reporte_kardexval_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepKardexValTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepKardexValExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		op = 0
		op1 = 0
		suma_c_e = 0
		suma_sub_e = 0
		suma_c_s = 0
		suma_sub_s = 0
		iteration = 1
		saldo = []
		saldo1 = []
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		aux1 = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1,f2), item__id=campo2, encingresoitem__cerrado=False)
		aux2 = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1,f2), item__id=campo2, encsalidaitem__cerrado=False)
		lista = list(sorted(chain(aux1, aux2), key=methodcaller('encabezado_f')))

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'KARDEX FISICO VALORADO DEL ' + str(fecha1) + ' ' + str(fecha2)
		for i in lista:
			ws['A2'] = 'CODIGO: ' + i.item.codigo
			ws['D2'] = 'ARTICULO: ' + i.item.nombre
			ws['G2'] = 'UNIDAD: ' + i.item.unidad.simbolo
			break

		ws['A3'] = 'Fecha'
		ws['B3'] = 'E/S'
		ws['C3'] = 'Referencia'
		ws['D3'] = 'Nº'
		ws['E3'] = 'Costo Unit.'
		ws['F3'] = 'Entradas'
		ws['H3'] = 'Salidas'
		ws['J3'] = 'Saldos'
		ws['F4'] = 'Cantidad'
		ws['G4'] = 'Valor'
		ws['H4'] = 'Cantidad'
		ws['I4'] = 'Valor'
		ws['J4'] = 'Cantidad'
		ws['K4'] = 'Valor'

		ws.merge_cells('A1:K1')
		ws.merge_cells('A2:C2')
		ws.merge_cells('D2:F2')
		ws.merge_cells('G2:K2')
		ws.merge_cells('F3:G3')
		ws.merge_cells('H3:I3')
		ws.merge_cells('J3:K3')

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J3'].alignment = Alignment(horizontal="center", vertical="center")

		ws['F4'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G4'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H4'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I4'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J4'].alignment = Alignment(horizontal="center", vertical="center")
		ws['K4'].alignment = Alignment(horizontal="center", vertical="center")

		ws.column_dimensions['A'].width = 12
		ws.column_dimensions['B'].width = 5
		ws.column_dimensions['C'].width = 20
		ws.column_dimensions['D'].width = 5
		ws.column_dimensions['E'].width = 15
		ws.column_dimensions['F'].width = 10
		ws.column_dimensions['G'].width = 10
		ws.column_dimensions['H'].width = 10
		ws.column_dimensions['I'].width = 10
		ws.column_dimensions['J'].width = 10
		ws.column_dimensions['K'].width = 10

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)
		ws['J2'].font = Font(bold=True)
		ws['K2'].font = Font(bold=True)

		ws['A3'].font = Font(bold=True)
		ws['B3'].font = Font(bold=True)
		ws['C3'].font = Font(bold=True)
		ws['D3'].font = Font(bold=True)
		ws['E3'].font = Font(bold=True)
		ws['F3'].font = Font(bold=True)
		ws['H3'].font = Font(bold=True)
		ws['J3'].font = Font(bold=True)

		ws['F4'].font = Font(bold=True)
		ws['G4'].font = Font(bold=True)
		ws['H4'].font = Font(bold=True)
		ws['I4'].font = Font(bold=True)
		ws['J4'].font = Font(bold=True)
		ws['K4'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')
		
		cont = 5
		
		for i in lista:
			if i.encabezado().tipo() == "E":
				suma_c_e = suma_c_e + i.cantidad
				suma_sub_e = suma_sub_e + i.sub_total
				op = op + i.cantidad
				op1 = op1 + i.sub_total
				saldo.append(op)
				saldo1.append(op1)
			else:
				suma_c_s = suma_c_s + i.cantidad
				suma_sub_s = suma_sub_s + i.sub_total
				op = op - i.cantidad
				op1 = op1 - i.sub_total
				saldo.append(op)
				saldo1.append(op1)

		lista_res = zip(lista, saldo, saldo1)
		num_queries = len(lista_res)

		for i, j, k in lista_res:
			if iteration < num_queries:
				ws.cell(row = cont, column = 1).value = i.encabezado().fecha_f()
				ws.cell(row = cont, column = 2).value = i.encabezado().tipo()
				ws.cell(row = cont, column = 3).value = i.encabezado().concepto
				ws.cell(row = cont, column = 4).value = i.encabezado().id
				ws.cell(row = cont, column = 5).value = i.costo_unit
				if i.encabezado().tipo() == "E":
					ws.cell(row = cont, column = 6).value = i.cantidad
					ws.cell(row = cont, column = 7).value = i.sub_total
				else:
					ws.cell(row = cont, column = 8).value = i.cantidad
					ws.cell(row = cont, column = 9).value = i.sub_total
				ws.cell(row = cont, column = 10).value = j
				ws.cell(row = cont, column = 11).value = k
				cont += 1
				iteration += 1
			else:
				ws.cell(row = cont, column = 1).value = i.encabezado().fecha_f()
				ws.cell(row = cont, column = 2).value = i.encabezado().tipo()
				ws.cell(row = cont, column = 3).value = i.encabezado().concepto
				ws.cell(row = cont, column = 4).value = i.encabezado().id
				ws.cell(row = cont, column = 5).value = i.costo_unit
				if i.encabezado().tipo() == "E":
					ws.cell(row = cont, column = 6).value = i.cantidad
					ws.cell(row = cont, column = 7).value = i.sub_total
				else:
					ws.cell(row = cont, column = 8).value = i.cantidad
					ws.cell(row = cont, column = 9).value = i.sub_total
				ws.cell(row = cont, column = 10).value = j
				ws.cell(row = cont, column = 11).value = k
				cont += 1
				ws.cell(row = cont, column = 1).value = 'TOTALES'
				ws.cell(row = cont, column = 6).value = suma_c_e
				ws.cell(row = cont, column = 7).value = suma_sub_e
				ws.cell(row = cont, column = 8).value = suma_c_s
				ws.cell(row = cont, column = 9).value = suma_sub_s
			
		nombre_archivo = "Reporte Kardex Físico Valorado.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepResKardexValTemplateExcel(FormView):
	form_class = reporteResKardexForm
	template_name = 'almacen/temp_reporte_reskardexval_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepResKardexValTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepResKardexValExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		f3 = f1 - datetime.timedelta(days=1)
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")
		fecha3 = f3.strftime("%d/%m/%Y")
		b = 0
		c = 0

		d = 0
		e = 0

		f = 0

		tot_saldo_a = 0
		tot_saldo_e = 0
		tot_saldo_s = 0
		tot_saldo = 0

		iteration = 1

		salida = []

		saldo_c = []
		cantidad_e = []
		cantidad_s = []
		saldo_sub = []
		sub_total_e = []
		sub_total_s = []

		saldo_c_a = []
		cantidad_e_a = []
		cantidad_s_a = []
		saldo_sub_a = []
		sub_total_e_a = []
		sub_total_s_a = []

		saldo_c_t = []
		saldo_sub_t = []

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'RESUMEN KARDEX FISICO VALORADO DEL ' + str(fecha1) + ' al ' + str(fecha2)
		ws['A2'] = 'Código'
		ws['B2'] = 'Descripción'
		ws['C2'] = 'Unidad'
		ws['D2'] = 'Saldo al ' + str(fecha3)
		ws['F2'] = 'Entradas'
		ws['H2'] = 'Salidas'
		ws['J2'] = 'Saldos'
		ws['D3'] = 'Cantidad'
		ws['E3'] = 'Valor'
		ws['F3'] = 'Cantidad'
		ws['G3'] = 'Valor'
		ws['H3'] = 'Cantidad'
		ws['I3'] = 'Valor'
		ws['J3'] = 'Cantidad'
		ws['K3'] = 'Valor'

		ws.merge_cells('A1:K1')
		ws.merge_cells('D2:E2')
		ws.merge_cells('F2:G2')
		ws.merge_cells('H2:I2')
		ws.merge_cells('J2:K2')

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['I3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['J3'].alignment = Alignment(horizontal="center", vertical="center")
		ws['K3'].alignment = Alignment(horizontal="center", vertical="center")

		ws.column_dimensions['A'].width = 12
		ws.column_dimensions['B'].width = 45
		ws.column_dimensions['C'].width = 8
		ws.column_dimensions['D'].width = 10
		ws.column_dimensions['E'].width = 10
		ws.column_dimensions['F'].width = 10
		ws.column_dimensions['G'].width = 10
		ws.column_dimensions['H'].width = 10
		ws.column_dimensions['I'].width = 10
		ws.column_dimensions['J'].width = 10
		ws.column_dimensions['K'].width = 10

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)
		ws['I2'].font = Font(bold=True)
		ws['J2'].font = Font(bold=True)
		ws['K2'].font = Font(bold=True)

		ws['A3'].font = Font(bold=True)
		ws['B3'].font = Font(bold=True)
		ws['C3'].font = Font(bold=True)
		ws['D3'].font = Font(bold=True)
		ws['E3'].font = Font(bold=True)
		ws['F3'].font = Font(bold=True)
		ws['G3'].font = Font(bold=True)
		ws['H3'].font = Font(bold=True)
		ws['I3'].font = Font(bold=True)
		ws['J3'].font = Font(bold=True)
		ws['K3'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')
		
		contexc = 5

		aux_a = item.objects.filter(ingresoitem__encingresoitem__fecha__lte=f3, ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__cantidad'))
		aux1_a = item.objects.filter(salidaitem__encsalidaitem__fecha__lte=f3, salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__cantidad'))
		aux2_a = item.objects.filter(ingresoitem__encingresoitem__fecha__lte=f3, ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__sub_total'))
		aux3_a = item.objects.filter(salidaitem__encsalidaitem__fecha__lte=f3, salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__sub_total'))
			
		resp = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).order_by('codigo').distinct()
		resp_a = item.objects.filter(ingresoitem__encingresoitem__fecha__lte=f3, ingresoitem__encingresoitem__cerrado=False).order_by('codigo').distinct()

		aux = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__cantidad'))
		aux1 = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__cantidad'))
		aux2 = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('ingresoitem__sub_total'))
		aux3 = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total=Sum('salidaitem__sub_total'))
		for q in aux_a:
			cantidad_e_a.append(q.total)

		if aux1_a:
			for i in aux_a:
				cont = 0
				if i in aux1_a:
					cantidad_s_a.append(aux1_a[cont].total)
					cont+=1
				else:
					cantidad_s_a.append(0)
		
		else:
			for i in aux_a:
				cantidad_s_a.append(0)


		for x in aux2_a:
			sub_total_e_a.append(x.total)

		if aux3_a:
			cont1 = 0
			for i in aux2_a:
				if i in aux3_a:
					sub_total_s_a.append(aux3_a[cont1].total)
					cont1+=1
				else:
					sub_total_s_a.append(0)
		else:
			for i in aux2_a:
				sub_total_s_a.append(0)
		
		for r in resp:
			if r in resp_a:
				saldo_c_a.append(cantidad_e_a[d] - cantidad_s_a[d])
				saldo_sub_a.append(sub_total_e_a[d] - sub_total_s_a[d])
				d+=1
			else:
				saldo_c_a.append(0)
				saldo_sub_a.append(0)

		for i in saldo_sub_a:
			tot_saldo_a = tot_saldo_a + i
			
		for i in aux:
			cantidad_e.append(i.total)

		if aux1:
			cont2 = 0
			for i in aux:
				if i in aux1:
					cantidad_s.append(aux1[cont2].total)
					cont2+=1
				else:
					cantidad_s.append(0)
		else:
			for i in aux1:
				cantidad_s.append(0)

		for k in item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo'):
			saldo_c.append(cantidad_e[b] - cantidad_s[b])
			b+=1

		for m in aux2:
			sub_total_e.append(m.total)

		if aux3:
			cont3 = 0
			for i in aux2:
				if i in aux3:
					sub_total_s.append(aux3[cont3].total)
					cont3+=1
				else:
					sub_total_s.append(0)

		else:
			for i in aux2:
				sub_total_s.append(0)

		for i in sub_total_e:
			tot_saldo_e = tot_saldo_e + i

		for i in sub_total_s:
			tot_saldo_s = tot_saldo_s + i

		for o in item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo'):
			saldo_sub.append(sub_total_e[c] - sub_total_s[c])
			c+=1

		res = item.objects.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct().order_by('codigo').annotate(total_cant_ent=Sum('ingresoitem__cantidad'), total_sub_ent=Sum('ingresoitem__sub_total'))
		res1 = item.objects.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct().order_by('codigo').annotate(total_cant_sal=Sum('salidaitem__cantidad'), total_sub_sal=Sum('salidaitem__sub_total'))

		if res1:
			x = 0
			for i in res:
				if i in res1:
					salida.append(res1[x])
					x+=1
				else:
					salida.append(0)
		else:
			for i in res:
				salida.append(0)


		for p in item.objects.filter(Q(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False) | Q(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False)).distinct().order_by('codigo'):
			saldo_c_t.append(saldo_c[f] + saldo_c_a[f])
			saldo_sub_t.append(saldo_sub[f] + saldo_sub_a[f])
			f+=1

		for i in saldo_sub_t:
			tot_saldo = tot_saldo + i

		lista = zip(res, salida, saldo_c_t, saldo_sub_t, saldo_c_a, saldo_sub_a)
		num_queries = len(lista)

		for i, j, k, l, m, n in lista:
			if iteration < num_queries:
				ws.cell(row = contexc, column = 1).value = i.codigo
				ws.cell(row = contexc, column = 2).value = i.nombre
				ws.cell(row = contexc, column = 3).value = i.unidad.simbolo
				ws.cell(row = contexc, column = 4).value = m
				ws.cell(row = contexc, column = 5).value = n
				if (i == None) | (i == 0):
					ws.cell(row = contexc, column = 6).value = 0
					ws.cell(row = contexc, column = 7).value = 0
				else:
					ws.cell(row = contexc, column = 6).value = i.total_cant_ent
					ws.cell(row = contexc, column = 7).value = i.total_sub_ent
				if (j == None) | (j == 0):
					ws.cell(row = contexc, column = 8).value = 0
					ws.cell(row = contexc, column = 9).value = 0
				else:
					ws.cell(row = contexc, column = 8).value = j.total_cant_sal
					ws.cell(row = contexc, column = 9).value = j.total_sub_sal
				ws.cell(row = contexc, column = 10).value = k
				ws.cell(row = contexc, column = 11).value = l
				contexc += 1
				iteration += 1
			else:
				ws.cell(row = contexc, column = 1).value = i.codigo
				ws.cell(row = contexc, column = 2).value = i.nombre
				ws.cell(row = contexc, column = 3).value = i.unidad.simbolo
				ws.cell(row = contexc, column = 4).value = m
				ws.cell(row = contexc, column = 5).value = n
				if (i == None) | (i == 0):
					ws.cell(row = contexc, column = 6).value = 0
					ws.cell(row = contexc, column = 7).value = 0
				else:
					ws.cell(row = contexc, column = 6).value = i.total_cant_ent
					ws.cell(row = contexc, column = 7).value = i.total_sub_ent
				if (j == None) | (j == 0):
					ws.cell(row = contexc, column = 8).value = 0
					ws.cell(row = contexc, column = 9).value = 0
				else:
					ws.cell(row = contexc, column = 8).value = j.total_cant_sal
					ws.cell(row = contexc, column = 9).value = j.total_sub_sal
				ws.cell(row = contexc, column = 10).value = k
				ws.cell(row = contexc, column = 11).value = l
				contexc += 1
				ws.cell(row = contexc, column = 1).value = 'TOTALES'
				ws.cell(row = contexc, column = 5).value = tot_saldo_a
				ws.cell(row = contexc, column = 7).value = tot_saldo_e
				ws.cell(row = contexc, column = 9).value = tot_saldo_s
				ws.cell(row = contexc, column = 11).value = tot_saldo

		nombre_archivo = "Reporte Resumen Kardex Físico Valorado.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepResumenEntSalTemplateExcel(FormView):
	form_class = repResumenEySForm
	template_name = 'almacen/temp_reporte_resumen_entsal_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepResumenEntSalTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepResumenEntSalExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")

		wb = Workbook()
		ws = wb.active
	
		ws['A2'] = 'Nº'
		ws['B2'] = 'Fecha'
		if campo2 == 'E':
			ws['A1'] = 'RESUMEN ENTRADAS AL ALMACEN DEL ' + str(fecha1) + ' al ' + str(fecha2)
			ws['C2'] = 'Referencia'
			ws['D2'] = 'Importe Factura'
			ws['E2'] = 'Importe IVA'
			ws['F2'] = 'Importe'
			ws.merge_cells('A1:F1')
		else:
			ws['A1'] = 'RESUMEN SALIDAS DEL ALMACEN DEL ' + str(fecha1) + ' al ' + str(fecha2)
			ws['C2'] = 'Referencia'
			ws['D2'] = 'Importe'
			ws.merge_cells('A1:D1')

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")

		ws.column_dimensions['A'].width = 5
		ws.column_dimensions['B'].width = 10
		ws.column_dimensions['C'].width = 40
		ws.column_dimensions['D'].width = 15
		ws.column_dimensions['E'].width = 15
		ws.column_dimensions['F'].width = 15

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')
		
		cont = 3
		iteration = 1
		
		if campo2 == 'E':
			res = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).order_by('id')
			num_queries = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).count()
			sumas = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total_i=Sum('total_iva'), total_imp_i=Sum('importe_iva'), total=Sum('total'))
		else:
			res = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).order_by('id')
			num_queries = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).count()
			sumas_s = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total=Sum('salidaitem__sub_total'))

		for encabezado in res:
			if iteration < num_queries:
				ws.cell(row = cont, column = 1).value = encabezado.id
				ws.cell(row = cont, column = 2).value = encabezado.fecha_f()
				ws.cell(row = cont, column = 3).value = encabezado.concepto
				if campo2 == 'E':
					ws.cell(row = cont, column = 4).value = encabezado.total
					ws.cell(row = cont, column = 5).value = encabezado.importe_iva
					ws.cell(row = cont, column = 6).value = encabezado.total_iva
				else:
					ws.cell(row = cont, column = 4).value = encabezado.importe(f1, f2)
				cont += 1
				iteration += 1
			else:
				ws.cell(row = cont, column = 1).value = encabezado.id
				ws.cell(row = cont, column = 2).value = encabezado.fecha_f()
				ws.cell(row = cont, column = 3).value = encabezado.concepto
				if campo2 == 'E':
					ws.cell(row = cont, column = 4).value = encabezado.total
					ws.cell(row = cont, column = 5).value = encabezado.importe_iva
					ws.cell(row = cont, column = 6).value = encabezado.total_iva
				else:
					ws.cell(row = cont, column = 4).value = encabezado.importe(f1, f2)
				cont += 1
				ws.cell(row = cont, column = 1).value = "TOTALES"
				if campo2 == 'E':
					ws.cell(row = cont, column = 4).value = sumas['total']
					ws.cell(row = cont, column = 5).value = sumas['total_imp_i']
					ws.cell(row = cont, column = 6).value = sumas['total_i']
				else:
					ws.cell(row = cont, column = 4).value = sumas_s['total']
			
		nombre_archivo = "Reporte Resumen de Entradas y Salidas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepDetalleEntTemplateExcel(FormView):
	form_class = repDetalleEntradasForm
	template_name = 'almacen/temp_reporte_detalle_ent_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepDetalleEntTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepDetalleEntExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")

		wb = Workbook()
		ws = wb.active
		
		if campo2 == 'T':
			ws['A1'] = 'DETALLE DE ENTRADAS AL ALMACEN (POR TIPO DE ENTRADA) DEL ' + str(fecha1) + ' al ' + str(fecha2)
			ws['A2'] = 'Nº'
			ws['B2'] = 'Fecha'
			ws['C2'] = 'Factura o Recibo'
			ws['D2'] = 'Proveedor'
			ws['E2'] = 'Importe'
			ws['F2'] = 'I.V.A.'
			ws['G2'] = 'Costo'

			ws.column_dimensions['A'].width = 5
			ws.column_dimensions['B'].width = 10
			ws.column_dimensions['C'].width = 15
			ws.column_dimensions['D'].width = 25
			ws.column_dimensions['E'].width = 15
			ws.column_dimensions['F'].width = 15
			ws.column_dimensions['G'].width = 15
		else:
			ws['A1'] = 'DETALLE DE ENTRADAS AL ALMACEN (POR CATEGORIA Y SUB CATEGORIA) DEL ' + str(fecha1) + ' al ' + str(fecha2)
			ws['A2'] = 'Código'
			ws['B2'] = 'Descripción'
			ws['C2'] = 'Unidad'
			ws['D2'] = 'Cantidad'
			ws['E2'] = 'Importe'
			ws['F2'] = 'I.V.A.'
			ws['G2'] = 'Costo'

			ws.column_dimensions['A'].width = 10
			ws.column_dimensions['B'].width = 45
			ws.column_dimensions['C'].width = 8
			ws.column_dimensions['D'].width = 10
			ws.column_dimensions['E'].width = 15
			ws.column_dimensions['F'].width = 15
			ws.column_dimensions['G'].width = 15

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)

		ws.merge_cells('A1:G1')

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')
		
		cont = 3
		iteration = 1
		
		if campo2 == 'T':
			res = tipo_entrada.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).order_by('id').distinct()
			num_queries = tipo_entrada.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).distinct().count()
			sumas = encingresoitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total_i=Sum('total_iva'), total_imp_i=Sum('importe_iva'), total=Sum('total'))

		else:
			res = categoria.objects.filter(item__ingresoitem__encingresoitem__fecha__range=(f1, f2), item__ingresoitem__encingresoitem__cerrado=False).order_by('id').distinct()
			num_queries = categoria.objects.filter(item__ingresoitem__encingresoitem__fecha__range=(f1, f2), item__ingresoitem__encingresoitem__cerrado=False).distinct().count()
			sumas = ingresoitem.objects.filter(item__categoria__in=res, encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total=Sum('costo_total'), total_imp_i=Sum(F('costo_total') - F('sub_total')), total_i=Sum('sub_total'))

		if campo2 == "T":
			for entrada in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = entrada.nombre.upper()
					cont += 1
					for encingreso in entrada.encingresoitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for ingreso in encingreso.ingresoitem_set.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = ingreso.encingresoitem.id
							ws.cell(row = cont, column = 2).value = ingreso.encingresoitem.fecha_f()
							ws.cell(row = cont, column = 3).value = ingreso.encingresoitem.factura
							ws.cell(row = cont, column = 4).value = ingreso.encingresoitem.proveedor.nombre
							ws.cell(row = cont, column = 5).value = ingreso.costo_total
							ws.cell(row = cont, column = 6).value = ingreso.iva()
							ws.cell(row = cont, column = 7).value = ingreso.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + entrada.nombre.upper()
					ws.cell(row = cont, column = 5).value = entrada.tipo_ent_importe_excel(f1, f2)
					ws.cell(row = cont, column = 6).value = entrada.tipo_ent_iva_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = entrada.tipo_ent_costo_excel(f1, f2)
					cont += 1
					iteration += 1
				else:
					ws.cell(row = cont, column = 1).value = entrada.nombre.upper()
					cont += 1
					for encingreso in entrada.encingresoitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for ingreso in encingreso.ingresoitem_set.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = ingreso.encingresoitem.id
							ws.cell(row = cont, column = 2).value = ingreso.encingresoitem.fecha_f()
							ws.cell(row = cont, column = 3).value = ingreso.encingresoitem.factura
							ws.cell(row = cont, column = 4).value = ingreso.encingresoitem.proveedor.nombre
							ws.cell(row = cont, column = 5).value = ingreso.costo_total
							ws.cell(row = cont, column = 6).value = ingreso.iva()
							ws.cell(row = cont, column = 7).value = ingreso.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + entrada.nombre.upper()
					ws.cell(row = cont, column = 5).value = entrada.tipo_ent_importe_excel(f1, f2)
					ws.cell(row = cont, column = 6).value = entrada.tipo_ent_iva_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = entrada.tipo_ent_costo_excel(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = "TOTALES"
					ws.cell(row = cont, column = 5).value = sumas['total']
					ws.cell(row = cont, column = 6).value = sumas['total_imp_i']
					ws.cell(row = cont, column = 7).value = sumas['total_i']

		else:
			for categorias in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = categorias.nombre.upper()
					cont += 1
					for subcategorias in categorias.subcategoria_set.filter(item__ingresoitem__encingresoitem__fecha__range=(f1, f2), item__ingresoitem__encingresoitem__cerrado=False).distinct():
						ws.cell(row = cont, column = 1).value = subcategorias.nombre.upper()
						cont += 1
						for items in subcategorias.item_set.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct():
							ws.cell(row = cont, column = 1).value = items.codigo
							ws.cell(row = cont, column = 2).value = items.nombre
							ws.cell(row = cont, column = 3).value = items.unidad.simbolo
							ws.cell(row = cont, column = 4).value = items.total_cant_e(f1, f2)
							ws.cell(row = cont, column = 5).value = items.importe_total_e(f1, f2)
							ws.cell(row = cont, column = 6).value = items.importe_iva(f1, f2)
							ws.cell(row = cont, column = 7).value = items.importe_total_iva(f1, f2)
							cont += 1
						ws.cell(row = cont, column = 1).value = "TOTAL.- " + subcategorias.nombre.upper()
						ws.cell(row = cont, column = 4).value = subcategorias.total_cant_e(f1, f2)
						ws.cell(row = cont, column = 5).value = subcategorias.importe_total(f1, f2)
						ws.cell(row = cont, column = 6).value = subcategorias.importe_iva(f1, f2)
						ws.cell(row = cont, column = 7).value = subcategorias.importe_total_iva(f1, f2)
						cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + categorias.nombre.upper()
					ws.cell(row = cont, column = 4).value = categorias.total_cant_e(f1, f2)
					ws.cell(row = cont, column = 5).value = categorias.importe_total(f1, f2)
					ws.cell(row = cont, column = 6).value = categorias.importe_iva(f1, f2)
					ws.cell(row = cont, column = 7).value = categorias.importe_total_iva(f1, f2)
					cont += 1
					iteration += 1

				else:
					ws.cell(row = cont, column = 1).value = categorias.nombre.upper()
					cont += 1
					for subcategorias in categorias.subcategoria_set.filter(item__ingresoitem__encingresoitem__fecha__range=(f1, f2), item__ingresoitem__encingresoitem__cerrado=False).distinct():
						ws.cell(row = cont, column = 1).value = subcategorias.nombre.upper()
						cont += 1
						for items in subcategorias.item_set.filter(ingresoitem__encingresoitem__fecha__range=(f1, f2), ingresoitem__encingresoitem__cerrado=False).distinct():
							ws.cell(row = cont, column = 1).value = items.codigo
							ws.cell(row = cont, column = 2).value = items.nombre
							ws.cell(row = cont, column = 3).value = items.unidad.simbolo
							ws.cell(row = cont, column = 4).value = items.total_cant_e(f1, f2)
							ws.cell(row = cont, column = 5).value = items.importe_total_e(f1, f2)
							ws.cell(row = cont, column = 6).value = items.importe_iva(f1, f2)
							ws.cell(row = cont, column = 7).value = items.importe_total_iva(f1, f2)
							cont += 1
						ws.cell(row = cont, column = 1).value = "TOTAL.- " + subcategorias.nombre.upper()
						ws.cell(row = cont, column = 4).value = subcategorias.total_cant_e(f1, f2)
						ws.cell(row = cont, column = 5).value = subcategorias.importe_total(f1, f2)
						ws.cell(row = cont, column = 6).value = subcategorias.importe_iva(f1, f2)
						ws.cell(row = cont, column = 7).value = subcategorias.importe_total_iva(f1, f2)
						cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + categorias.nombre.upper()
					ws.cell(row = cont, column = 4).value = categorias.total_cant_e(f1, f2)
					ws.cell(row = cont, column = 5).value = categorias.importe_total(f1, f2)
					ws.cell(row = cont, column = 6).value = categorias.importe_iva(f1, f2)
					ws.cell(row = cont, column = 7).value = categorias.importe_total_iva(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = "TOTALES"
					ws.cell(row = cont, column = 4).value = sumas['total_cant']
					ws.cell(row = cont, column = 5).value = sumas['total']
					ws.cell(row = cont, column = 6).value = sumas['total_imp_i']
					ws.cell(row = cont, column = 7).value = sumas['total_i']
			
		nombre_archivo = "Reporte Detalle de Entradas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepDetalleSalTemplateExcel(FormView):
	form_class = repDetalleSalidasForm
	template_name = 'almacen/temp_reporte_detalle_sal_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepDetalleSalTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepDetalleSalExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")

		wb = Workbook()
		ws = wb.active
		
		if campo2 == 'T':
			ws['A1'] = 'DETALLE DE SALIDAS AL ALMACEN (POR TIPO DE SALIDA) DEL ' + str(fecha1) + ' al ' + str(fecha2)
			ws['A2'] = 'Nº'
			ws['B2'] = 'Fecha'
			ws['C2'] = 'Cliente'
			ws['D2'] = 'Importe'

			ws.column_dimensions['A'].width = 5
			ws.column_dimensions['B'].width = 10
			ws.column_dimensions['C'].width = 50
			ws.column_dimensions['D'].width = 15
			ws.merge_cells('A1:D1')
		else:
			ws['A1'] = 'DETALLE DE SALIDAS AL ALMACEN (POR CATEGORIA Y SUB CATEGORIA) DEL ' + str(fecha1) + ' al ' + str(fecha2)
			ws['A2'] = 'Código'
			ws['B2'] = 'Descripción'
			ws['C2'] = 'Unidad'
			ws['D2'] = 'Cantidad'
			ws['E2'] = 'Importe'

			ws.column_dimensions['A'].width = 10
			ws.column_dimensions['B'].width = 45
			ws.column_dimensions['C'].width = 8
			ws.column_dimensions['D'].width = 15
			ws.column_dimensions['E'].width = 15
			ws.merge_cells('A1:E1')

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		cont = 3
		iteration = 1
		
		if campo2 == 'T':
			res = tipo_salida.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).order_by('id').distinct()
			num_queries = tipo_salida.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).distinct().count()
			sumas = encsalidaitem.objects.filter(fecha__range=(f1, f2), cerrado=False).aggregate(total=Sum('salidaitem__sub_total'))
		else:
			res = categoria.objects.filter(item__salidaitem__encsalidaitem__fecha__range=(f1, f2), item__salidaitem__encsalidaitem__cerrado=False).order_by('id').distinct()
			num_queries = categoria.objects.filter(item__salidaitem__encsalidaitem__fecha__range=(f1, f2), item__salidaitem__encsalidaitem__cerrado=False).distinct().count()
			sumas = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total=Sum('sub_total'))

		if campo2 == "T":
			for t_salida in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = t_salida.nombre.upper()
					cont += 1
					for encsalida in t_salida.encsalidaitem_set.filter(fecha__range=(f1, f2), cerrado=False).distinct():
						for salida in encsalida.salidaitem_set.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).distinct():
							ws.cell(row = cont, column = 1).value = salida.encsalidaitem.id
							ws.cell(row = cont, column = 2).value = salida.encsalidaitem.fecha_f()
							ws.cell(row = cont, column = 3).value = salida.encsalidaitem.cliente.nombre
							ws.cell(row = cont, column = 4).value = salida.encsalidaitem.importe(f1, f2)
							cont += 1
							break
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + t_salida.nombre.upper()
					ws.cell(row = cont, column = 4).value = t_salida.tipo_sal_importe_excel(f1, f2)
					cont += 1
					iteration += 1
				else:
					ws.cell(row = cont, column = 1).value = t_salida.nombre.upper()
					cont += 1
					for encsalida in t_salida.encsalidaitem_set.filter(fecha__range=(f1, f2), cerrado=False).distinct():
						for salida in encsalida.salidaitem_set.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).distinct():
							ws.cell(row = cont, column = 1).value = salida.encsalidaitem.id
							ws.cell(row = cont, column = 2).value = salida.encsalidaitem.fecha_f()
							ws.cell(row = cont, column = 3).value = salida.encsalidaitem.cliente.nombre
							ws.cell(row = cont, column = 4).value = salida.encsalidaitem.importe(f1, f2)
							cont += 1
							break
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + t_salida.nombre.upper()
					ws.cell(row = cont, column = 4).value = t_salida.tipo_sal_importe_excel(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = "TOTALES"
					ws.cell(row = cont, column = 4).value = sumas['total']

		else:
			for categorias in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = categorias.nombre.upper()
					cont += 1
					for subcategorias in categorias.subcategoria_set.filter(item__salidaitem__encsalidaitem__fecha__range=(f1, f2), item__salidaitem__encsalidaitem__cerrado=False).distinct():
						ws.cell(row = cont, column = 1).value = subcategorias.nombre.upper()
						cont += 1
						for items in subcategorias.item_set.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct():
							ws.cell(row = cont, column = 1).value = items.codigo
							ws.cell(row = cont, column = 2).value = items.nombre
							ws.cell(row = cont, column = 3).value = items.unidad.simbolo
							ws.cell(row = cont, column = 4).value = items.total_cant_s(f1, f2)
							ws.cell(row = cont, column = 5).value = items.importe_total_s(f1, f2)
							cont += 1
						ws.cell(row = cont, column = 1).value = "TOTAL.- " + subcategorias.nombre.upper()
						ws.cell(row = cont, column = 4).value = subcategorias.total_cant_s(f1, f2)
						ws.cell(row = cont, column = 5).value = subcategorias.importe_total_s(f1, f2)
						cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + categorias.nombre.upper()
					ws.cell(row = cont, column = 4).value = categorias.total_cant_s(f1, f2)
					ws.cell(row = cont, column = 5).value = categorias.importe_total_s(f1, f2)
					cont += 1
					iteration += 1

				else:
					ws.cell(row = cont, column = 1).value = categorias.nombre.upper()
					cont += 1
					for subcategorias in categorias.subcategoria_set.filter(item__salidaitem__encsalidaitem__fecha__range=(f1, f2), item__salidaitem__encsalidaitem__cerrado=False).distinct():
						ws.cell(row = cont, column = 1).value = subcategorias.nombre.upper()
						cont += 1
						for items in subcategorias.item_set.filter(salidaitem__encsalidaitem__fecha__range=(f1, f2), salidaitem__encsalidaitem__cerrado=False).distinct():
							ws.cell(row = cont, column = 1).value = items.codigo
							ws.cell(row = cont, column = 2).value = items.nombre
							ws.cell(row = cont, column = 3).value = items.unidad.simbolo
							ws.cell(row = cont, column = 4).value = items.total_cant_s(f1, f2)
							ws.cell(row = cont, column = 5).value = items.importe_total_s(f1, f2)
							cont += 1
						ws.cell(row = cont, column = 1).value = "TOTAL.- " + subcategorias.nombre.upper()
						ws.cell(row = cont, column = 4).value = subcategorias.total_cant_s(f1, f2)
						ws.cell(row = cont, column = 5).value = subcategorias.importe_total_s(f1, f2)
						cont += 1
					ws.cell(row = cont, column = 1).value = "TOTAL.- " + categorias.nombre.upper()
					ws.cell(row = cont, column = 4).value = categorias.total_cant_s(f1, f2)
					ws.cell(row = cont, column = 5).value = categorias.importe_total_s(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = "TOTALES"
					ws.cell(row = cont, column = 4).value = sumas['total_cant']
					ws.cell(row = cont, column = 5).value = sumas['total']
			
		nombre_archivo = "Reporte Detalle de Salidas.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepDetalleClienteTemplateExcel(FormView):
	form_class = repDetalleClienteForm
	template_name = 'almacen/temp_reporte_detalle_cliente_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepDetalleClienteTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepDetalleClienteExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'REPORTE POR CLIENTES (SALIDAS) DEL ' + str(fecha1) + ' al ' + str(fecha2)
		ws['A2'] = 'Nº'
		ws['B2'] = 'Fecha'
		ws['C2'] = 'Producto'
		ws['D2'] = 'Unidad'
		ws['E2'] = 'Cantidad'
		ws['F2'] = 'Costo_unit.'
		ws['G2'] = 'Importe'

		ws.column_dimensions['A'].width = 5
		ws.column_dimensions['B'].width = 10
		ws.column_dimensions['C'].width = 50
		ws.column_dimensions['D'].width = 8
		ws.column_dimensions['E'].width = 10
		ws.column_dimensions['F'].width = 15
		ws.column_dimensions['G'].width = 10

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:G1')
		
		cont = 3
		iteration = 1
		
		if campo2 == 'T':
			res = cliente.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).order_by('id').distinct()
			num_queries = cliente.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).distinct().count()
			sumas = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))

			for clientes in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = clientes.nombre
					cont += 1
					for encsalidas in clientes.encsalidaitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for salidas in encsalidas.salidaitem_set.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = salidas.encsalidaitem.id
							ws.cell(row = cont, column = 2).value = salidas.encsalidaitem.fecha_f()
							ws.cell(row = cont, column = 3).value = str(salidas.item.codigo) + ' ' + salidas.item.nombre
							ws.cell(row = cont, column = 4).value = salidas.item.unidad.simbolo
							ws.cell(row = cont, column = 5).value = salidas.cantidad
							ws.cell(row = cont, column = 6).value = salidas.costo_unit
							ws.cell(row = cont, column = 7).value = salidas.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + clientes.nombre
					ws.cell(row = cont, column = 5).value = clientes.cantidad_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 6).value = clientes.costo_unit_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = clientes.importe_cliente_s_excel(f1, f2)
					cont += 1
					iteration += 1
				else:
					ws.cell(row = cont, column = 1).value = clientes.nombre
					cont += 1
					for encsalidas in clientes.encsalidaitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for salidas in encsalidas.salidaitem_set.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = salidas.encsalidaitem.id
							ws.cell(row = cont, column = 2).value = salidas.encsalidaitem.fecha_f()
							ws.cell(row = cont, column = 3).value = str(salidas.item.codigo) + ' ' + salidas.item.nombre
							ws.cell(row = cont, column = 4).value = salidas.item.unidad.simbolo
							ws.cell(row = cont, column = 5).value = salidas.cantidad
							ws.cell(row = cont, column = 6).value = salidas.costo_unit
							ws.cell(row = cont, column = 7).value = salidas.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + clientes.nombre
					ws.cell(row = cont, column = 5).value = clientes.cantidad_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 6).value = clientes.costo_unit_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = clientes.importe_cliente_s_excel(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTALES'
					ws.cell(row = cont, column = 5).value = sumas['total_cant']
					ws.cell(row = cont, column = 6).value = sumas['total_costo_u']
					ws.cell(row = cont, column = 7).value = sumas['total']
		else:
			res = cliente.objects.filter(encsalidaitem__fecha__range=(f1, f2), id=campo2, encsalidaitem__cerrado=False).order_by('id').distinct()
			num_queries = cliente.objects.filter(encsalidaitem__fecha__range=(f1, f2), id=campo2, encsalidaitem__cerrado=False).distinct().count()
			sumas = salidaitem.objects.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cliente=campo2, encsalidaitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))

			for clientes in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = clientes.nombre
					cont += 1
					for encsalidas in clientes.encsalidaitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for salidas in encsalidas.salidaitem_set.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = salidas.encsalidaitem.id
							ws.cell(row = cont, column = 2).value = salidas.encsalidaitem.fecha_f()
							ws.cell(row = cont, column = 3).value = str(salidas.item.codigo) + ' ' + salidas.item.nombre
							ws.cell(row = cont, column = 4).value = salidas.item.unidad.simbolo
							ws.cell(row = cont, column = 5).value = salidas.cantidad
							ws.cell(row = cont, column = 6).value = salidas.costo_unit
							ws.cell(row = cont, column = 7).value = salidas.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + clientes.nombre
					ws.cell(row = cont, column = 5).value = clientes.cantidad_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 6).value = clientes.costo_unit_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = clientes.importe_cliente_s_excel(f1, f2)
					cont += 1
					iteration += 1
				else:
					ws.cell(row = cont, column = 1).value = clientes.nombre
					cont += 1
					for encsalidas in clientes.encsalidaitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for salidas in encsalidas.salidaitem_set.filter(encsalidaitem__fecha__range=(f1, f2), encsalidaitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = salidas.encsalidaitem.id
							ws.cell(row = cont, column = 2).value = salidas.encsalidaitem.fecha_f()
							ws.cell(row = cont, column = 3).value = str(salidas.item.codigo) + ' ' + salidas.item.nombre
							ws.cell(row = cont, column = 4).value = salidas.item.unidad.simbolo
							ws.cell(row = cont, column = 5).value = salidas.cantidad
							ws.cell(row = cont, column = 6).value = salidas.costo_unit
							ws.cell(row = cont, column = 7).value = salidas.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + clientes.nombre
					ws.cell(row = cont, column = 5).value = clientes.cantidad_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 6).value = clientes.costo_unit_cliente_s_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = clientes.importe_cliente_s_excel(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTALES'
					ws.cell(row = cont, column = 5).value = sumas['total_cant']
					ws.cell(row = cont, column = 6).value = sumas['total_costo_u']
					ws.cell(row = cont, column = 7).value = sumas['total']

		nombre_archivo = "Reporte Detalle de Clientes.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response

class RepDetalleProovTemplateExcel(FormView):
	form_class = repDetalleProveedorForm
	template_name = 'almacen/temp_reporte_detalle_proov_excel.html'
	def get_context_data(self, **kwargs):
		context = super(RepDetalleProovTemplateExcel, self).get_context_data(**kwargs)
		if 'form'in context:
			context['form'] = self.form_class()
		return context

class RepDetalleProovExcel(TemplateView):
	def get(self, request, *args, **kwargs):
		campo = self.request.GET.get('campo')
		campo1 = self.request.GET.get('campo1')
		campo2 = self.request.GET.get('campo2')
		f1 = datetime.datetime.strptime(campo, "%Y-%m-%d")
		f2 = datetime.datetime.strptime(campo1, "%Y-%m-%d")
		fecha1 = f1.strftime("%d/%m/%Y")
		fecha2 = f2.strftime("%d/%m/%Y")

		wb = Workbook()
		ws = wb.active
		
		ws['A1'] = 'REPORTE POR PROVEEDORES (ENTRADAS) DEL ' + str(fecha1) + ' al ' + str(fecha2)
		ws['A2'] = 'Nº'
		ws['B2'] = 'Fecha'
		ws['C2'] = 'Factura/Recibo'
		ws['D2'] = 'Producto'
		ws['E2'] = 'Unid.'
		ws['F2'] = 'Cant.'
		ws['G2'] = 'Costo Unit.'
		ws['H2'] = 'Importe'

		ws.column_dimensions['A'].width = 5
		ws.column_dimensions['B'].width = 10
		ws.column_dimensions['C'].width = 18
		ws.column_dimensions['D'].width = 50
		ws.column_dimensions['E'].width = 6
		ws.column_dimensions['F'].width = 6
		ws.column_dimensions['G'].width = 15
		ws.column_dimensions['H'].width = 15

		ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
		ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['C2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['D2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['E2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['F2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['G2'].alignment = Alignment(horizontal="center", vertical="center")
		ws['H2'].alignment = Alignment(horizontal="center", vertical="center")

		ws['A1'].font = Font(bold=True)
		ws['A2'].font = Font(bold=True)
		ws['B2'].font = Font(bold=True)
		ws['C2'].font = Font(bold=True)
		ws['D2'].font = Font(bold=True)
		ws['E2'].font = Font(bold=True)
		ws['F2'].font = Font(bold=True)
		ws['G2'].font = Font(bold=True)
		ws['H2'].font = Font(bold=True)

		ws['A1'].fill = PatternFill(start_color='219FB5', end_color='219FB5', fill_type='solid')

		ws.merge_cells('A1:H1')
		
		cont = 3
		iteration = 1
		
		if campo2 == 'T':
			res = proveedor.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).order_by('id').distinct()
			num_queries = proveedor.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).distinct().count()
			sumas = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))

			for proveedores in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = proveedores.nombre
					cont += 1
					for encingresos in proveedores.encingresoitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for ingresos in encingresos.ingresoitem_set.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = ingresos.encingresoitem.id
							ws.cell(row = cont, column = 2).value = ingresos.encingresoitem.fecha_f()
							ws.cell(row = cont, column = 3).value = ingresos.encingresoitem.factura
							ws.cell(row = cont, column = 4).value = str(ingresos.item.codigo) + ' ' + ingresos.item.nombre
							ws.cell(row = cont, column = 5).value = ingresos.item.unidad.simbolo
							ws.cell(row = cont, column = 6).value = ingresos.cantidad
							ws.cell(row = cont, column = 7).value = ingresos.costo_unit
							ws.cell(row = cont, column = 8).value = ingresos.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + proveedores.nombre
					ws.cell(row = cont, column = 6).value = proveedores.cantidad_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = proveedores.costo_unit_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 8).value = proveedores.importe_proveedor_e_excel(f1, f2)
					cont += 1
					iteration += 1
				else:
					ws.cell(row = cont, column = 1).value = proveedores.nombre
					cont += 1
					for encingresos in proveedores.encingresoitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for ingresos in encingresos.ingresoitem_set.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = ingresos.encingresoitem.id
							ws.cell(row = cont, column = 2).value = ingresos.encingresoitem.fecha_f()
							ws.cell(row = cont, column = 3).value = ingresos.encingresoitem.factura
							ws.cell(row = cont, column = 4).value = str(ingresos.item.codigo) + ' ' + ingresos.item.nombre
							ws.cell(row = cont, column = 5).value = ingresos.item.unidad.simbolo
							ws.cell(row = cont, column = 6).value = ingresos.cantidad
							ws.cell(row = cont, column = 7).value = ingresos.costo_unit
							ws.cell(row = cont, column = 8).value = ingresos.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + proveedores.nombre
					ws.cell(row = cont, column = 6).value = proveedores.cantidad_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = proveedores.costo_unit_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 8).value = proveedores.importe_proveedor_e_excel(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTALES'
					ws.cell(row = cont, column = 6).value = sumas['total_cant']
					ws.cell(row = cont, column = 7).value = sumas['total_costo_u']
					ws.cell(row = cont, column = 8).value = sumas['total']
		else:
			res = proveedor.objects.filter(encingresoitem__fecha__range=(f1, f2), id=campo2, encingresoitem__cerrado=False).order_by('id').distinct()
			num_queries = proveedor.objects.filter(encingresoitem__fecha__range=(f1, f2), id=campo2, encingresoitem__cerrado=False).distinct().count()
			sumas = ingresoitem.objects.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__proveedor=campo2, encingresoitem__cerrado=False).aggregate(total_cant=Sum('cantidad'), total_costo_u=Sum('costo_unit'), total=Sum('sub_total'))

			for proveedores in res:
				if iteration < num_queries:
					ws.cell(row = cont, column = 1).value = proveedores.nombre
					cont += 1
					for encingresos in proveedores.encingresoitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for ingresos in encingresos.ingresoitem_set.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = ingresos.encingresoitem.id
							ws.cell(row = cont, column = 2).value = ingresos.encingresoitem.fecha_f()
							ws.cell(row = cont, column = 3).value = ingresos.encingresoitem.factura
							ws.cell(row = cont, column = 4).value = str(ingresos.item.codigo) + ' ' + ingresos.item.nombre
							ws.cell(row = cont, column = 5).value = ingresos.item.unidad.simbolo
							ws.cell(row = cont, column = 6).value = ingresos.cantidad
							ws.cell(row = cont, column = 7).value = ingresos.costo_unit
							ws.cell(row = cont, column = 8).value = ingresos.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + proveedores.nombre
					ws.cell(row = cont, column = 6).value = proveedores.cantidad_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = proveedores.costo_unit_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 8).value = proveedores.importe_proveedor_e_excel(f1, f2)
					cont += 1
					iteration += 1
				else:
					ws.cell(row = cont, column = 1).value = proveedores.nombre
					cont += 1
					for encingresos in proveedores.encingresoitem_set.filter(fecha__range=(f1, f2), cerrado=False):
						for ingresos in encingresos.ingresoitem_set.filter(encingresoitem__fecha__range=(f1, f2), encingresoitem__cerrado=False):
							ws.cell(row = cont, column = 1).value = ingresos.encingresoitem.id
							ws.cell(row = cont, column = 2).value = ingresos.encingresoitem.fecha_f()
							ws.cell(row = cont, column = 3).value = ingresos.encingresoitem.factura
							ws.cell(row = cont, column = 4).value = str(ingresos.item.codigo) + ' ' + ingresos.item.nombre
							ws.cell(row = cont, column = 5).value = ingresos.item.unidad.simbolo
							ws.cell(row = cont, column = 6).value = ingresos.cantidad
							ws.cell(row = cont, column = 7).value = ingresos.costo_unit
							ws.cell(row = cont, column = 8).value = ingresos.sub_total
							cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTAL.- ' + proveedores.nombre
					ws.cell(row = cont, column = 6).value = proveedores.cantidad_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 7).value = proveedores.costo_unit_proveedor_e_excel(f1, f2)
					ws.cell(row = cont, column = 8).value = proveedores.importe_proveedor_e_excel(f1, f2)
					cont += 1
					ws.cell(row = cont, column = 1).value = 'TOTALES'
					ws.cell(row = cont, column = 6).value = sumas['total_cant']
					ws.cell(row = cont, column = 7).value = sumas['total_costo_u']
					ws.cell(row = cont, column = 8).value = sumas['total']

		nombre_archivo = "Reporte Detalle de Proveedores.xlsx"
		response = HttpResponse(content_type = "application/ms-excel")
		content = "attachment; filename = {0}".format(nombre_archivo)
		response['Content-Disposition'] = content
		response.content = save_virtual_workbook(wb)
		return response
